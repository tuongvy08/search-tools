"""Durable import queue, secure upload storage and repo-native worker."""
from __future__ import annotations

import hashlib
import os
from pathlib import Path
import re
import threading
import time
import uuid
from contextlib import contextmanager

from psycopg2.extras import Json, RealDictCursor, execute_values

from db import get_connection
from brand_gateway import acquire_products_import_lock
from import_engine import ImportProblem, limit, workbook_rows, create_stage, build_plan, apply_plan


def upload_dir():
    path = Path(os.environ.get('IMPORT_UPLOAD_DIR', '/var/lib/search-tools/imports'))
    if not path.is_absolute() or path.is_symlink():
        raise ImportProblem('Thư mục upload chưa được cấu hình an toàn.')
    path.mkdir(mode=0o700, parents=True, exist_ok=True)
    return path.resolve()


def upload_path(job_id):
    return upload_dir() / (str(uuid.UUID(str(job_id))) + '.xlsx')


@contextmanager
def connection():
    conn = get_connection(connect_timeout=10)
    try:
        yield conn
    finally:
        conn.close()


def event(cur, job_id, actor, name):
    cur.execute('INSERT INTO product_import_events(job_id,actor,event) VALUES (%s,%s,%s)', (str(job_id), actor, name))


def fetch_job(cur, job_id):
    cur.execute('SELECT * FROM product_import_jobs WHERE id=%s', (str(job_id),))
    return cur.fetchone()


def list_jobs():
    with connection() as conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute('SELECT * FROM product_import_jobs ORDER BY created_at DESC LIMIT 30')
        return cur.fetchall()


def submit(file, mode, actor, user_id, auth_version, submission_key):
    if mode not in ('append','upsert','replace_by_brand'):
        raise ImportProblem('Chế độ nhập không hợp lệ.')
    try:
        submission_key = str(uuid.UUID(submission_key))
    except (ValueError, TypeError, AttributeError):
        raise ImportProblem('Phiên upload không hợp lệ. Hãy tải lại trang.') from None
    filename = (file.filename or '').replace('\\','/').rsplit('/',1)[-1]
    filename = ''.join(c for c in filename if c.isprintable())[:180]
    if not filename.lower().endswith('.xlsx'):
        raise ImportProblem('Chỉ hỗ trợ Excel Workbook (.xlsx).')
    if file.mimetype not in ('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','application/octet-stream','application/zip'):
        raise ImportProblem('Loại nội dung không phù hợp với XLSX.')
    job_id = str(uuid.uuid4())
    path = upload_path(job_id)
    with connection() as conn:
        try:
            with conn, conn.cursor() as cur:
                # Serializes capacity accounting and same-submission admission.
                cur.execute('SELECT pg_advisory_xact_lock(62402401)')
                cur.execute('SELECT id FROM product_import_jobs WHERE actor_user_id=%s AND submission_key=%s', (user_id, submission_key))
                old = cur.fetchone()
                if old:
                    return str(old[0])
                cur.execute("SELECT count(*),COALESCE(sum(file_size),0) FROM product_import_jobs WHERE purged_at IS NULL")
                count, used = cur.fetchone()
                if count >= limit('MAX_RETAINED_JOBS',100) or used >= limit('DISK_BYTES', 8*1024**3):
                    raise ImportProblem('Kho upload đã đầy. Đợi worker dọn tệp hết hạn rồi thử lại.')
                total = 0
                digest = hashlib.sha256()
                fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
                with os.fdopen(fd, 'wb') as dest:
                    while True:
                        block = file.stream.read(1024**2)
                        if not block:
                            break
                        total += len(block)
                        if total > limit('MAX_BYTES',128*1024**2) or used+total > limit('DISK_BYTES',8*1024**3):
                            raise ImportProblem('Tệp vượt giới hạn dung lượng upload hoặc dung lượng kho.')
                        digest.update(block)
                        dest.write(block)
                    dest.flush()
                    os.fsync(dest.fileno())
                # Header validation is cheap; ZIP parsing happens only in worker.
                with path.open('rb') as source:
                    if source.read(4) != b'PK\x03\x04':
                        raise ImportProblem('Nội dung không phải XLSX; không đổi đuôi CSV thành .xlsx.')
                cur.execute("""INSERT INTO product_import_jobs
                    (id,submission_key,actor,actor_user_id,actor_auth_version,filename,file_size,file_sha256,mode,expires_at)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,now()+%s*interval '1 hour')""",
                    (job_id,submission_key,actor,user_id,auth_version,filename,total,digest.hexdigest(),mode,limit('RETENTION_HOURS',48)))
                event(cur,job_id,actor,'uploaded')
        except BaseException:
            path.unlink(missing_ok=True)
            raise
    return job_id


def control(job_id, action, actor, fingerprint='', confirm_delete=''):
    with connection() as conn, conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
        job = fetch_job(cur,job_id)
        if not job:
            raise ImportProblem('Không tìm thấy tác vụ.')
        cur.execute('SELECT * FROM product_import_jobs WHERE id=%s FOR UPDATE', (str(job_id),))
        job = cur.fetchone()
        if action == 'cancel':
            if job['status'] not in ('queued','running') or job['cancel_requested']:
                return
            if job['status'] == 'queued':
                cur.execute("UPDATE product_import_jobs SET status='cancelled',cancel_requested=true,finished_at=now() WHERE id=%s", (str(job_id),))
            else:
                cur.execute('UPDATE product_import_jobs SET cancel_requested=true WHERE id=%s', (str(job_id),))
        elif action in ('apply','retry'):
            cur.execute('SELECT expires_at>now() AND purged_at IS NULL AS valid FROM product_import_jobs WHERE id=%s', (str(job_id),))
            if not cur.fetchone()['valid']:
                raise ImportProblem('Tệp đã hết hạn; hãy upload lại.')
            if action == 'apply':
                if job['status'] == 'queued' and job['phase']=='apply':
                    return  # A double click cannot enqueue another application.
                if job['status'] != 'completed' or job['phase'] != 'preview' or not job['preview_ready']:
                    raise ImportProblem('Cần xem trước thành công trước khi ghi dữ liệu.')
                preview = job['preview'] or {}
                if fingerprint != preview.get('fingerprint'):
                    raise ImportProblem('Xem trước đã thay đổi. Hãy tải lại trang.')
                if job['mode'] == 'replace_by_brand' and confirm_delete != str(preview.get('deleted')):
                    raise ImportProblem('Cần nhập chính xác số dòng sẽ xóa ở bước xác nhận riêng.')
                cur.execute("UPDATE product_import_jobs SET status='queued',phase='apply',preview_ready=false,finished_at=NULL,processed_count=0 WHERE id=%s", (str(job_id),))
            else:
                if job['status'] not in ('failed','cancelled','completed') or (job['status']=='completed' and job['phase']=='apply'):
                    raise ImportProblem('Không thể thử lại tác vụ này.')
                if job['attempts'] >= limit('MAX_ATTEMPTS',6):
                    raise ImportProblem('Đã hết số lần thử. Hãy upload lại workbook.')
                # Always generates a NEW preview; never silently retries apply.
                cur.execute("UPDATE product_import_jobs SET status='queued',phase='preview',preview_ready=false,preview=NULL,cancel_requested=false,finished_at=NULL,errors='[]',error_count=0,processed_count=0 WHERE id=%s", (str(job_id),))
        else:
            raise ImportProblem('Thao tác không hợp lệ.')
        event(cur,job_id,actor,action)


def _job_lock(cur, job_id):
    cur.execute("SELECT pg_try_advisory_lock(hashtextextended(%s,624024))", (str(job_id),))
    return cur.fetchone()[0]


def _job_unlock(cur, job_id):
    cur.execute("SELECT pg_advisory_unlock(hashtextextended(%s,624024))", (str(job_id),))


def claim(conn, only_id=None):
    with conn.cursor() as cur:
        cur.execute("SELECT id FROM product_import_jobs WHERE status IN ('queued','running') AND (%s IS NULL OR id=%s::uuid) ORDER BY created_at LIMIT 100", (only_id,only_id))
        ids = [str(r[0]) for r in cur.fetchall()]
        conn.commit()
        for job_id in ids:
            if not _job_lock(cur,job_id):
                continue
            cur.execute('SELECT status,phase,cancel_requested,attempts,expires_at<now() FROM product_import_jobs WHERE id=%s FOR UPDATE', (job_id,))
            status,phase,cancel,attempts,expired = cur.fetchone()
            if status not in ('queued','running'):
                conn.rollback()
                _job_unlock(cur,job_id)
                continue
            if status=='running':
                # Session lock disappeared: previous worker died. Its product
                # transaction rolled back, or completed state committed with it.
                cur.execute("UPDATE product_import_jobs SET status='failed',preview_ready=false,finished_at=now(),errors=%s,error_count=1 WHERE id=%s", (Json(['Worker bị gián đoạn. Xem trước lại trước khi tiếp tục.']),job_id))
                event(cur,job_id,'worker','crash_recovered')
                conn.commit()
                _job_unlock(cur,job_id)
                continue
            if cancel or expired or attempts>=limit('MAX_ATTEMPTS',6):
                cur.execute("UPDATE product_import_jobs SET status='cancelled',finished_at=now() WHERE id=%s", (job_id,))
                conn.commit()
                _job_unlock(cur,job_id)
                continue
            cur.execute("UPDATE product_import_jobs SET status='running',started_at=now(),heartbeat_at=now(),attempts=attempts+1 WHERE id=%s", (job_id,))
            event(cur,job_id,'worker',phase+'_started')
            conn.commit()
            return job_id
    return None


def run_once(only_id=None):
    with connection() as conn:
        job_id = claim(conn,only_id)
        if not job_id:
            return False
        stop = threading.Event()
        cancelled = threading.Event()
        deadline = time.monotonic()+limit('JOB_SECONDS',7200)
        def watch():
            try:
                with connection() as monitor:
                    monitor.autocommit=True
                    with monitor.cursor() as cur:
                        while not stop.wait(1):
                            cur.execute('UPDATE product_import_jobs SET heartbeat_at=now() WHERE id=%s RETURNING cancel_requested', (job_id,))
                            if cur.fetchone()[0] or time.monotonic()>deadline:
                                cancelled.set()
                                conn.cancel()
                                break
            except Exception:
                cancelled.set()
                conn.cancel()  # fail closed if control channel cannot be checked
        watcher = threading.Thread(target=watch,daemon=True)
        watcher.start()
        def progress(count):
            if cancelled.is_set():
                raise ImportProblem('Tác vụ đã được hủy hoặc vượt thời gian xử lý.')
            with connection() as status_conn, status_conn, status_conn.cursor() as cur:
                cur.execute('UPDATE product_import_jobs SET processed_count=%s,heartbeat_at=now() WHERE id=%s RETURNING cancel_requested', (count,job_id))
                if cur.fetchone()[0]:
                    cancelled.set()
                    raise ImportProblem('Tác vụ đã được hủy.')
        try:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                job = fetch_job(cur,job_id)
            conn.commit()
            if job['phase']=='preview':
                path = upload_path(job_id)
                digest=hashlib.sha256()
                with path.open('rb') as source:
                    for chunk in iter(lambda:source.read(1024**2),b''):
                        digest.update(chunk)
                if digest.hexdigest()!=job['file_sha256']:
                    raise ImportProblem('Tệp lưu trữ đã thay đổi. Hãy upload lại.')
                with conn, conn.cursor() as cur:
                    cur.execute('DELETE FROM product_import_rows WHERE job_id=%s',(job_id,))
                batch=[]
                count=skipped=0
                def blank(_line, number):
                    nonlocal skipped
                    skipped+=number
                for n,row in workbook_rows(path,blank):
                    batch.append((job_id,n,Json(row)))
                    if len(batch)>=limit('CHUNK_ROWS',2000):
                        with conn,conn.cursor() as cur:
                            execute_values(cur,'INSERT INTO product_import_rows VALUES %s',batch)
                        count+=len(batch)
                        progress(count)
                        batch=[]
                if batch:
                    with conn,conn.cursor() as cur:
                        execute_values(cur,'INSERT INTO product_import_rows VALUES %s',batch)
                    count+=len(batch)
                with conn,conn.cursor() as cur:
                    cur.execute('UPDATE product_import_jobs SET row_count=%s,skipped_count=%s WHERE id=%s',(count,skipped,job_id))
                progress(count)
            with conn, conn.cursor() as cur:
                cur.execute("SET LOCAL statement_timeout = %s", (str(limit('SQL_SECONDS',600)*1000),))
                # Serializes preview snapshot and apply with all product writers.
                acquire_products_import_lock(cur)
                if cancelled.is_set():
                    raise ImportProblem('Tác vụ đã được hủy.')
                cur.execute("SELECT 1 FROM app_users WHERE id=%s AND is_admin=true AND account_status='ACTIVE' AND auth_version=%s", (job['actor_user_id'],job['actor_auth_version']))
                if not cur.fetchone():
                    raise ImportProblem('Người tải tệp không còn quyền admin hoặc phiên đã bị thu hồi.')
                create_stage(cur)
                cur.execute('INSERT INTO import_stage SELECT row_number,data FROM product_import_rows WHERE job_id=%s', (job_id,))
                plan = build_plan(cur,job['mode'],apply=job['phase']=='apply')
                if job['phase']=='apply':
                    apply_plan(cur,plan,job['preview'],progress)
                # Stop watcher before final row lock: avoids monitor/job deadlock.
                stop.set()
                watcher.join(timeout=5)
                if cancelled.is_set():
                    raise ImportProblem('Tác vụ đã được hủy hoặc vượt thời gian xử lý.')
                cur.execute('SELECT cancel_requested FROM product_import_jobs WHERE id=%s FOR UPDATE',(job_id,))
                if cur.fetchone()[0]:
                    raise ImportProblem('Tác vụ đã được hủy.')
                cur.execute("""UPDATE product_import_jobs SET status='completed',preview_ready=%s,preview=%s,
                    processed_count=%s,inserted_count=%s,updated_count=%s,deleted_count=%s,finished_at=now(),errors='[]',error_count=0 WHERE id=%s""",
                    (job['phase']=='preview',Json(plan),plan['row_count'],plan['inserted'] if job['phase']=='apply' else 0,
                     plan['updated'] if job['phase']=='apply' else 0,plan['deleted'] if job['phase']=='apply' else 0,job_id))
                event(cur,job_id,'worker',job['phase']+'_completed')
        except Exception as exc:
            conn.rollback()
            # Never store or log exception strings from DB, ZIP, filesystem.
            message = str(exc)[:min(4000,limit('ERROR_CHARS',400))] if isinstance(exc,ImportProblem) else 'Xử lý thất bại. Kiểm tra cấu hình worker hoặc xem trước lại; chưa ghi dữ liệu.'
            with conn,conn.cursor() as cur:
                cur.execute("UPDATE product_import_jobs SET status=CASE WHEN cancel_requested THEN 'cancelled' ELSE 'failed' END,preview_ready=false,errors=%s,error_count=1,finished_at=now() WHERE id=%s", (Json([message]),job_id))
                event(cur,job_id,'worker','stopped')
        finally:
            stop.set()
            watcher.join(timeout=5)
            with conn.cursor() as cur:
                _job_unlock(cur,job_id)
            conn.commit()
        return True


def cleanup():
    """Only UUID files in our private directory; never follow symlinks.

    Session job locks fence cleanup against processing. Audit/summary kept for
    a separate bounded period; sample rows removed with payload expiry.
    """
    with connection() as conn, conn.cursor() as cur:
        cur.execute('SELECT id FROM product_import_jobs WHERE expires_at<now() AND purged_at IS NULL LIMIT 100')
        for (job_id,) in cur.fetchall():
            if not _job_lock(cur,job_id):
                continue
            try:
                upload_path(job_id).unlink(missing_ok=True)
                cur.execute('DELETE FROM product_import_rows WHERE job_id=%s',(job_id,))
                cur.execute("UPDATE product_import_jobs SET purged_at=now(),preview=NULL,preview_ready=false,status=CASE WHEN status IN ('queued','running') THEN 'cancelled' ELSE status END WHERE id=%s", (job_id,))
                conn.commit()
            finally:
                _job_unlock(cur,job_id)
        cur.execute("DELETE FROM product_import_jobs WHERE purged_at IS NOT NULL AND created_at<now()-%s*interval '1 day'", (limit('AUDIT_DAYS',90),))
        cur.execute('DELETE FROM admin_rule_import_previews WHERE expires_at<now()')
        conn.commit()
        for path in upload_dir().iterdir():
            if not re.fullmatch(r'[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\.xlsx',path.name) or path.is_symlink():
                continue
            if time.time()-path.stat().st_mtime < limit('ORPHAN_HOURS',24)*3600:
                continue
            cur.execute('SELECT 1 FROM product_import_jobs WHERE id=%s',(path.stem,))
            if cur.fetchone() is None:
                path.unlink(missing_ok=True)


def read_rules_upload(file):
    """Small legacy rules import; separate 2 MB/10k row admission limit."""
    import csv
    import io
    import tempfile
    from openpyxl import load_workbook
    from import_engine import inspect_workbook
    raw = file.stream.read(2*1024**2+1)
    if len(raw)>2*1024**2:
        raise ImportProblem('File quy tắc tối đa 2 MB.')
    if (file.filename or '').lower().endswith('.csv'):
        iterator=csv.reader(io.StringIO(raw.decode('utf-8-sig')))
        wb=None
        temp=None
    else:
        temp=tempfile.NamedTemporaryFile(suffix='.xlsx')
        temp.write(raw);temp.flush()
        inspect_workbook(temp.name)
        wb=load_workbook(temp.name,read_only=True,data_only=False,keep_links=False)
        if len(wb.worksheets)!=1:
            wb.close();temp.close()
            raise ImportProblem('Chỉ hỗ trợ một trang tính.')
        ws=wb.worksheets[0];ws.reset_dimensions()
        def values():
            for row in ws.iter_rows():
                if any(c.data_type in ('f','e') for c in row):
                    raise ImportProblem('Hãy thay công thức bằng giá trị.')
                yield [c.value for c in row]
        iterator=values()
    try:
        headers=[str(v or '').strip().lower() for v in next(iterator)]
        if len(headers)>16 or len(set(headers))!=len(headers):raise ImportProblem('Tiêu đề không hợp lệ.')
        rows=[]
        for n,values in enumerate(iterator,2):
            if n>10001 or len(values)>len(headers) or any(len(str(v or ''))>4000 for v in values):
                raise ImportProblem('File quy tắc vượt giới hạn.')
            if any(v is not None and str(v).strip() for v in values):rows.append(dict(zip(headers, [str(v).strip() if v is not None else '' for v in values])))
        return rows,set(headers)
    finally:
        if wb:wb.close()
        if temp:temp.close()
