"""Menu "Báo giá → Đơn hàng".

Luồng: upload một hoặc nhiều file báo giá (.xlsx) -> hệ thống tự nhận diện cột
(người dùng chỉnh lại nếu cần) -> sửa/tick chọn hàng trên bảng -> tải về file
"bảng hàng hóa" đúng mẫu để import vào phần mềm Base.

Thiết kế:
- Không lưu file hay dữ liệu phía server. Trình duyệt giữ file; mỗi lần đổi cách
  ghép cột, trình duyệt gửi lại file kèm bảng ghép (`POST /parse`). Khi tải về,
  trình duyệt gửi các dòng đã chọn (`POST /export`) và server kiểm tra lại toàn bộ.
- Quyền dùng: cùng quyền `QUICK_QUOTE` của team (đã kèm quyền xem giá). Module này
  tự kiểm tra, không sửa hệ thống phân quyền (`team_permissions.py`).
- CSRF: dùng token của `session_security` (header `X-CSRF-Token`).
- Chỗ cần chỉnh khi mẫu thay đổi: `QUOTE_HEADERS`, `ORDER_COLUMNS`, `FIELD_META`.
"""
from __future__ import annotations

import io
import json
import logging
import re
import unicodedata
import zipfile
from datetime import datetime
from decimal import ROUND_HALF_UP, ROUND_UP, Decimal
from pathlib import Path

from flask import Blueprint, current_app, jsonify, redirect, render_template, request, send_file, session, url_for
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import session_security
import team_permissions
from quote_workbook_export import (
    MAX_XLSX_BYTES,
    MAX_XLSX_ENTRIES,
    MAX_XLSX_UNCOMPRESSED_BYTES,
    XLSX_MIME,
)

logger = logging.getLogger(__name__)

bp = Blueprint("quote_to_order", __name__, url_prefix="/quote-to-order")

REQUIRED_PERMISSION = "QUICK_QUOTE"
ORDER_TEMPLATE = Path(__file__).resolve().parent / "assets" / "quote_to_order" / "bang-hang-hoa_template.xlsx"
ORDER_SHEET = "Exported file"

MAX_FILES = 10
MAX_ORDER_ROWS = 1000                   # tổng số dòng trong một file đơn hàng
MAX_ROWS_PER_FILE = 1000
MAX_EXPORT_BODY_BYTES = 4 * 1024 * 1024
MAX_SCAN_ROWS = 3000
MAX_SCAN_COLS = 60
MAX_HEADER_SCAN_ROWS = 60
DEFAULT_HEADER_ROW = 16                 # PO: dòng tiêu đề mặc định của form báo giá; lệch thì cảnh báo cho người dùng chỉnh
DEFAULT_VAT_RATE = Decimal("0.08")
DEFAULT_SHEET = "BG"
TYPE_CHOICES = ("Nhập khẩu", "Mua trong nước")

# ---------------------------------------------------------------------------
# CẤU HÌNH MAPPING – chỗ duy nhất cần sửa nếu form thay đổi
# ---------------------------------------------------------------------------
# Trường nội bộ -> các tên cột (đã chuẩn hóa: chữ thường, gộp khoảng trắng) mà hệ
# thống tự nhận diện trong form báo giá. Thứ tự = ưu tiên.
QUOTE_HEADERS = {
    "name": ("tên hàng", "tên sản phẩm", "tên hàng hóa", "product name", "name"),
    "code": ("code", "mã hàng", "mã sản phẩm", "mã sp", "cat no", "catalog number"),
    "cas": ("cas", "cas no", "cas number", "số cas"),
    "brand": ("hãng", "hãng sản xuất", "nhà sản xuất", "brand"),
    "unit": ("đơn vị tính", "đvt", "đơn vị", "unit"),
    "qty": ("số lượng", "số lượng đặt", "sl", "quantity", "qty"),
    # Đơn giá chuyển sang đơn hàng = giá CÓ VAT (PO xác nhận 2026-10-04).
    "price": ("đơn giá có vat (vnđ)", "đơn giá có vat", "đơn giá gồm vat (vnđ)", "đơn giá gồm vat"),
    "cost": ("giá nhập chưa vat", "giá nhập", "đơn giá mua"),
    "type": ("loại hàng", "loại"),
    "note_goods": ("ghi chú hàng hóa", "ghi chú về hàng hóa"),
    "note_other": ("ghi chú khác",),
    # Form báo giá không có cột này: để trống thì khi xuất sẽ bằng Đơn giá (xem build_order_file).
    "min_price": ("giá bán tối thiểu", "giá tối thiểu"),
}
# Cột phụ, chỉ dùng để tự tính lại giá khi file chưa có giá trị tính sẵn.
AUX_HEADERS = {
    "price_ex": ("đơn giá chưa vat (vnđ)", "đơn giá chưa vat", "đơn giá (vnđ)", "đơn giá"),
    "margin": ("%ln", "% ln", "ln%"),
    "seq": ("tt", "stt", "số tt"),
}

# Tiêu đề cột trong file bảng hàng hóa (theo thứ tự trong file mẫu) -> trường nội bộ.
# "Thành tiền" luôn tính = Số lượng × Đơn giá, ghi dạng giá trị (không ghi công thức).
# "Giá bán tối thiểu" trống (không có cột / ô không có giá trị) thì bằng "Đơn giá".
ORDER_COLUMNS = {
    "Tên hàng (*)": "name",
    "Code (*)": "code",
    "Cas": "cas",
    "Hãng (*)": "brand",
    "ĐVT (*)": "unit",
    "Số lượng (*)": "qty",
    "Đơn giá (*)": "price",
    "Thành tiền": "amount",
    "Giá bán tối thiểu": "min_price",
    "Đơn giá mua dự kiến (Sales)": "cost",
    "Loại hàng (*)": "type",
    "Ghi chú về hàng hóa": "note_goods",
    "Ghi chú khác": "note_other",
}

# kind: text | num | choice. Trường mappable=False không cho ghép cột.
FIELD_META = {
    "name": {"label": "Tên hàng", "required": True, "kind": "text", "max_len": 500},
    "code": {"label": "Code", "required": True, "kind": "text", "max_len": 255},
    "cas": {"label": "Cas", "required": False, "kind": "text", "max_len": 255},
    "brand": {"label": "Hãng", "required": True, "kind": "text", "max_len": 255},
    "unit": {"label": "ĐVT", "required": True, "kind": "text", "max_len": 100},
    "qty": {"label": "Số lượng", "required": True, "kind": "num"},
    "price": {"label": "Đơn giá", "required": True, "kind": "num"},
    "cost": {"label": "Đơn giá mua dự kiến", "required": False, "kind": "num"},
    "type": {"label": "Loại hàng", "required": True, "kind": "choice"},
    "note_goods": {"label": "Ghi chú về hàng hóa", "required": False, "kind": "text", "max_len": 2000},
    "note_other": {"label": "Ghi chú khác", "required": False, "kind": "text", "max_len": 2000},
    "min_price": {"label": "Giá bán tối thiểu", "required": False, "kind": "num"},
}
MAPPABLE_FIELDS = tuple(FIELD_META)  # thứ tự ghép tự động

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_TONE_VARIANTS = (("oá", "óa"), ("oà", "òa"), ("oả", "ỏa"), ("oã", "õa"), ("oạ", "ọa"))
_VAT_RE = re.compile(r"thuế\s*vat[^\d%]*(\d+(?:[.,]\d+)?)\s*%", re.IGNORECASE)


class QuoteParseError(ValueError):
    """Lỗi người dùng đọc được (tiếng Việt), không chứa chi tiết nội bộ."""


# ---------------------------------------------------------------------------
# Chuẩn hóa giá trị
# ---------------------------------------------------------------------------
def norm_header(value) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value)).lower()
    text = re.sub(r"\(\*\)", " ", text)
    for old, new in _TONE_VARIANTS:
        text = text.replace(old, new)
    return re.sub(r"\s+", " ", text).strip()


def clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _CONTROL_CHARS.sub("", str(value)).strip()


_THOUSANDS_RE = re.compile(r"^\d{1,3}([.,]\d{3})+$")


def parse_number_text(text: str) -> Decimal | None:
    """'1.250.000' -> 1250000; '0,5' -> 0.5; '1.250.000,5' -> 1250000.5."""
    s = re.sub(r"\s+", "", text.replace(" ", "")).rstrip("đĐ")
    s = re.sub(r"(?i)vn[đd]$", "", s)
    if not s:
        return None
    sign = ""
    if s[0] in "+-":
        sign, s = s[0], s[1:]
    if not s or not re.fullmatch(r"[\d.,]+", s):
        return None
    if "." in s and "," in s:
        decimal_sep = "." if s.rfind(".") > s.rfind(",") else ","
        thousands_sep = "," if decimal_sep == "." else "."
        s = s.replace(thousands_sep, "").replace(decimal_sep, ".")
    elif _THOUSANDS_RE.match(s) and not s.startswith("0"):
        s = re.sub(r"[.,]", "", s)
    else:
        if s.count(",") + s.count(".") > 1:
            return None
        s = s.replace(",", ".")
    try:
        return Decimal(sign + s)
    except Exception:
        return None


def to_decimal(value) -> Decimal | None:
    """Số từ ô Excel hoặc chuỗi nhập tay; không phải số hữu hạn thì trả None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        result = value
    elif isinstance(value, int):
        result = Decimal(value)
    elif isinstance(value, float):
        result = Decimal(repr(value)) if value == value and value not in (float("inf"), float("-inf")) else None
    elif isinstance(value, str):
        result = parse_number_text(value)
    else:
        return None
    return result if result is not None and result.is_finite() else None


def py_number(value: Decimal):
    """Decimal -> int nếu là số nguyên, ngược lại float (để ghi Excel/JSON)."""
    return int(value) if value == value.to_integral_value() else float(value)


def excel_roundup(value: Decimal, digits: int = 0) -> Decimal:
    """ROUNDUP của Excel (làm tròn xa số 0) dùng Decimal, không lỗi float."""
    quantum = Decimal(10) ** -digits
    return (value / quantum).to_integral_value(rounding=ROUND_UP) * quantum


def calc_price_ex_vat(cost: Decimal, margin: Decimal) -> Decimal | None:
    """Công thức form báo giá: Q = ROUNDUP(giá nhập / (1 - %LN), -4)."""
    if margin < 0 or margin >= 1 or cost < 0:
        return None
    return excel_roundup(cost / (Decimal(1) - margin), -4)


def normalize_type(value) -> str:
    key = norm_header(value)
    for choice in TYPE_CHOICES:
        if key == norm_header(choice):
            return choice
    return ""


# ---------------------------------------------------------------------------
# Đọc báo giá
# ---------------------------------------------------------------------------
def check_xlsx_container(raw: bytes) -> None:
    """Chặn file không phải .xlsx, quá lớn, zip bomb hoặc có macro."""
    if not raw or raw[:2] != b"PK":
        raise QuoteParseError("File không đúng định dạng .xlsx.")
    if len(raw) > MAX_XLSX_BYTES:
        raise QuoteParseError(f"File quá lớn, tối đa {MAX_XLSX_BYTES // (1024 * 1024)}MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            infos = zf.infolist()
            names = {info.filename.lower() for info in infos}
            if len(infos) > MAX_XLSX_ENTRIES or sum(i.file_size for i in infos) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise QuoteParseError("File quá phức tạp hoặc quá lớn khi giải nén, hệ thống từ chối đọc.")
            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise QuoteParseError("File không đúng định dạng .xlsx.")
            if "xl/vbaproject.bin" in names:
                raise QuoteParseError("Không hỗ trợ file có macro. Hãy lưu lại thành .xlsx thường.")
    except zipfile.BadZipFile as exc:
        raise QuoteParseError("File bị hỏng hoặc không phải .xlsx.") from exc


def _load_rows(raw: bytes, sheet: str | None):
    """Trả (danh sách sheet, sheet đang đọc, các dòng dưới dạng tuple giá trị)."""
    check_xlsx_container(raw)
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        logger.warning("quote_to_order: cannot open workbook: %s", type(exc).__name__)
        raise QuoteParseError("Không đọc được file. File có thể bị hỏng hoặc không phải .xlsx.") from exc
    try:
        names = list(wb.sheetnames)
        if not names:
            raise QuoteParseError("File không có sheet nào.")
        chosen = sheet if sheet in names else (DEFAULT_SHEET if DEFAULT_SHEET in names else names[0])
        rows = []
        for row in wb[chosen].iter_rows(max_col=MAX_SCAN_COLS, values_only=True):
            rows.append(row)
            if len(rows) >= MAX_SCAN_ROWS:
                break
        return names, chosen, rows
    except QuoteParseError:
        raise
    except Exception as exc:
        logger.warning("quote_to_order: cannot read sheet: %s", type(exc).__name__)
        raise QuoteParseError("Không đọc được nội dung sheet. File có thể bị hỏng.") from exc
    finally:
        wb.close()


def _row_titles(row) -> dict[int, str]:
    return {i: norm_header(v) for i, v in enumerate(row) if clean_text(v)}


def _aliases(aliases) -> tuple[str, ...]:
    """Chuẩn hóa tên gọi giống tiêu đề cột để so khớp đối xứng (hóa/hoá, ạ/ä...)."""
    return tuple(norm_header(a) for a in aliases)


def _is_header_row(row) -> bool:
    titles = set(_row_titles(row).values())
    return bool(titles & set(_aliases(QUOTE_HEADERS["name"])) and titles & set(_aliases(QUOTE_HEADERS["code"])))


def _find_header_row(rows) -> int | None:
    """Số dòng (1-based) của dòng tiêu đề. Ưu tiên dòng mặc định (16); không có thì tìm theo nội dung."""
    if len(rows) >= DEFAULT_HEADER_ROW and _is_header_row(rows[DEFAULT_HEADER_ROW - 1]):
        return DEFAULT_HEADER_ROW
    for number, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS], start=1):
        if _is_header_row(row):
            return number
    return None


def _first_column(titles: dict[int, str], aliases) -> int | None:
    for alias in _aliases(aliases):
        for index, title in titles.items():
            if title == alias:
                return index
    return None


def auto_mapping(titles: dict[int, str]) -> dict[str, int | None]:
    result: dict[str, int | None] = {}
    used: set[int] = set()
    for key in MAPPABLE_FIELDS:
        taken = {i: t for i, t in titles.items() if i not in used}
        index = _first_column(taken, QUOTE_HEADERS[key])
        result[key] = index
        if index is not None:
            used.add(index)
    return result


def _sanitize_mapping(mapping, column_indexes: set[int], auto: dict) -> dict[str, int | None]:
    """Ghép cột do client gửi: trường vắng -> tự nhận diện; chỉ số lạ -> bỏ trống."""
    if not isinstance(mapping, dict):
        return dict(auto)
    result: dict[str, int | None] = {}
    for key in MAPPABLE_FIELDS:
        if key not in mapping:
            result[key] = auto[key]
            continue
        value = mapping[key]
        ok = isinstance(value, int) and not isinstance(value, bool) and value in column_indexes
        result[key] = value if ok else None
    return result


def _detect_vat_rate(rows, start: int) -> Decimal:
    for row in rows[start:]:
        for cell in row:
            if isinstance(cell, str):
                match = _VAT_RE.search(cell)
                if match:
                    rate = parse_number_text(match.group(1))
                    if rate is not None and 0 < rate < 100:
                        return rate / Decimal(100)
    return DEFAULT_VAT_RATE


def parse_quote(raw: bytes, *, sheet: str | None = None, header_row: int | None = None, mapping=None) -> dict:
    names, sheet_name, rows = _load_rows(raw, sheet)
    manual_header = isinstance(header_row, int) and not isinstance(header_row, bool) and 1 <= header_row <= len(rows)
    found = header_row if manual_header else _find_header_row(rows)
    # Cảnh báo chỉ khi hệ thống tự tìm thấy ở dòng khác 16; người dùng đã chủ động nhập dòng thì coi là đã xác nhận.
    base = {"sheet": sheet_name, "sheets": names, "header_row": found, "default_header_row": DEFAULT_HEADER_ROW,
            "header_mismatch": found is not None and not manual_header and found != DEFAULT_HEADER_ROW}
    if found is None:
        return {**base, "needs_header": True, "columns": [], "mapping": {}, "auto_mapping": {}, "items": [],
                "vat_rate": float(DEFAULT_VAT_RATE),
                "message": f"Không tìm thấy dòng tiêu đề (mặc định ở dòng {DEFAULT_HEADER_ROW}; cần có cột "
                           "'Tên hàng' và 'Code'). Hãy nhập đúng số dòng tiêu đề của file này."}

    titles = _row_titles(rows[found - 1])
    columns = [{"index": i, "letter": get_column_letter(i + 1), "title": clean_text(rows[found - 1][i])}
               for i in sorted(titles)]
    auto = auto_mapping(titles)
    chosen = _sanitize_mapping(mapping, set(titles), auto)
    aux = {key: _first_column(titles, aliases) for key, aliases in AUX_HEADERS.items()}
    price_vat_col = _first_column(titles, QUOTE_HEADERS["price"])
    vat_rate = _detect_vat_rate(rows, found)

    def cell(row, index):
        return row[index] if index is not None and index < len(row) else None

    def ex_vat_price(row):
        cached = to_decimal(cell(row, aux["price_ex"]))
        if cached is not None and cached > 0:
            return cached
        cost, margin = to_decimal(cell(row, chosen["cost"])), to_decimal(cell(row, aux["margin"]))
        if cost is None or margin is None:
            return None
        return calc_price_ex_vat(cost, margin)

    def price_of(row):
        index = chosen["price"]
        value = to_decimal(cell(row, index))
        if value is not None and value > 0:
            return value
        if index is None:
            return None
        # Chưa có giá trị tính sẵn (file chưa mở/lưu bằng Excel) -> tự tính lại theo công thức form.
        if index == price_vat_col:
            ex = ex_vat_price(row)
            return None if ex is None else (ex * (Decimal(1) + vat_rate)).quantize(Decimal(1), ROUND_HALF_UP)
        if index == aux["price_ex"]:
            return ex_vat_price(row)
        return None

    seq_col = aux["seq"] if aux["seq"] is not None else 0
    items: list[dict] = []
    for offset, row in enumerate(rows[found:], start=found + 1):
        first = cell(row, seq_col)
        if isinstance(first, str):
            first = first.strip() or None
        if isinstance(first, str) and not first.isdigit():
            break                                   # gặp phần tổng ("Tổng giá gồm VAT"...)
        text = {key: clean_text(cell(row, chosen[key])) for key in MAPPABLE_FIELDS if FIELD_META[key]["kind"] == "text"}
        if not text["name"] and not text["code"]:
            continue                                # dòng trống/placeholder chưa nhập hàng
        item = dict(text)
        qty = to_decimal(cell(row, chosen["qty"]))
        item["qty"] = py_number(qty) if qty is not None else None
        price = price_of(row)
        item["price"] = py_number(price) if price is not None else None
        for key in ("cost", "min_price"):
            value = to_decimal(cell(row, chosen[key]))
            if key == "min_price" and value is not None and value <= 0:
                value = None                        # ô có nhưng chưa có giá trị -> coi như trống
            item[key] = py_number(value) if value is not None else None
        item["type"] = normalize_type(cell(row, chosen["type"]))
        item["source_row"] = offset
        items.append(item)
        if len(items) > MAX_ROWS_PER_FILE:
            raise QuoteParseError(f"File có quá {MAX_ROWS_PER_FILE} dòng hàng, hệ thống từ chối đọc.")

    return {**base, "needs_header": False, "columns": columns, "mapping": chosen, "auto_mapping": auto,
            "items": items, "vat_rate": float(vat_rate),
            "unmapped_required": [k for k in MAPPABLE_FIELDS if FIELD_META[k]["required"] and chosen[k] is None]}


# ---------------------------------------------------------------------------
# Kiểm tra dòng gửi lên khi tải về + tạo file đơn hàng
# ---------------------------------------------------------------------------
def clean_order_rows(rows) -> tuple[list[dict], list[dict]]:
    """Kiểm tra lại dữ liệu từ client. Trả (dòng sạch, danh sách lỗi theo chỉ số dòng)."""
    clean: list[dict] = []
    errors: list[dict] = []
    for index, raw in enumerate(rows):
        if not isinstance(raw, dict):
            errors.append({"index": index, "missing": [], "invalid": ["Dữ liệu dòng"]})
            continue
        row: dict = {}
        missing: list[str] = []
        invalid: list[str] = []
        for key, meta in FIELD_META.items():
            value, label = raw.get(key), meta["label"]
            kind = meta["kind"]
            if kind == "text":
                if value is not None and not isinstance(value, (str, int, float)) or isinstance(value, bool):
                    invalid.append(label)
                    continue
                row[key] = clean_text(value)
                if len(row[key]) > meta["max_len"]:
                    invalid.append(label)
            elif kind == "num":
                number = to_decimal(value)
                if value not in (None, "") and number is None:
                    invalid.append(label)
                elif number is not None and (number < 0 or number >= Decimal(10) ** 15):
                    invalid.append(label)
                else:
                    row[key] = number
            else:
                row[key] = clean_text(value)
                if row[key] and row[key] not in TYPE_CHOICES:
                    invalid.append(label)
        for key, meta in FIELD_META.items():
            if not meta["required"] or meta["label"] in invalid:
                continue
            value = row.get(key)
            if value in ("", None):
                missing.append(meta["label"])
        if row.get("qty") is not None and row["qty"] <= 0 and "Số lượng" not in invalid:
            invalid.append("Số lượng")
        if missing or invalid:
            errors.append({"index": index, "missing": missing, "invalid": invalid})
        else:
            clean.append(row)
    return clean, errors


def build_order_file(rows: list[dict]) -> bytes:
    """Ghi các dòng đã kiểm tra vào bản sao file mẫu (giữ nguyên 2 sheet của mẫu)."""
    wb = load_workbook(ORDER_TEMPLATE)
    ws = wb[ORDER_SHEET]
    col_of = {clean_text(c.value): c.column for c in ws[1] if c.value}
    missing_headers = [h for h in ORDER_COLUMNS if h not in col_of]
    if missing_headers:
        raise RuntimeError(f"File mẫu đơn hàng không còn các cột: {missing_headers}")
    for row_number, item in enumerate(rows, start=2):
        values = dict(item)
        values["amount"] = _amount(item.get("qty"), item.get("price"))
        if values.get("min_price") is None:
            values["min_price"] = item.get("price")     # PO: Giá bán tối thiểu trống thì = Đơn giá (đã gồm VAT)
        for header, key in ORDER_COLUMNS.items():
            value = values.get(key)
            if value is None or value == "":
                continue
            cell = ws.cell(row=row_number, column=col_of[header])
            if isinstance(value, Decimal):
                cell.value = py_number(value)
            elif isinstance(value, (int, float)):
                cell.value = value
            else:
                cell.value = value
                cell.data_type = "s"            # chuỗi bắt đầu bằng "=" không bị hiểu là công thức
                if key in ("code", "cas"):
                    cell.number_format = "@"    # Code/CAS giữ dạng chữ, không bị đổi thành số/ngày
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _amount(qty, price):
    if qty is None or price is None:
        return None
    return py_number((qty * price).quantize(Decimal("0.01"), ROUND_HALF_UP))


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
def _json(payload: dict, status: int = 200):
    response = jsonify(payload)
    response.status_code = status
    response.headers["Cache-Control"] = "no-store, private"
    return response


def _guard(api: bool):
    """Trả None nếu được phép; ngược lại trả response từ chối."""
    if not session.get("authenticated"):
        if api:
            return _json({"ok": False, "error": "Phiên đăng nhập đã hết hạn. Vui lòng đăng nhập lại."}, 401)
        return redirect(url_for("login"))
    if not team_permissions.can(REQUIRED_PERMISSION):
        message = "Team của bạn chưa có quyền dùng chức năng này."
        return _json({"ok": False, "error": message}, 403) if api else (message, 403)
    if api:
        token = request.headers.get("X-CSRF-Token", "") or request.form.get("csrf_token", "")
        if not session_security.verify_csrf_token(token):
            return _json({"ok": False, "error": "Phiên làm việc không hợp lệ hoặc đã hết hạn. Hãy tải lại trang."}, 400)
    return None


@bp.get("/")
def index():
    denied = _guard(api=False)
    if denied is not None:
        return denied
    config = {
        "fields": [{"key": key, **{k: meta[k] for k in ("label", "required", "kind")}}
                   for key, meta in FIELD_META.items()],
        "typeChoices": list(TYPE_CHOICES),
        "limits": {"maxFiles": MAX_FILES, "maxRows": MAX_ORDER_ROWS, "maxFileBytes": MAX_XLSX_BYTES},
        "urls": {"parse": url_for("quote_to_order.parse"), "export": url_for("quote_to_order.export")},
    }
    response = current_app.make_response(render_template("quote_to_order.html", config=config))
    response.headers["Cache-Control"] = "no-store, private"
    return response


@bp.post("/parse")
def parse():
    denied = _guard(api=True)
    if denied is not None:
        return denied
    if request.content_length and request.content_length > MAX_XLSX_BYTES + 512 * 1024:
        return _json({"ok": False, "error": f"File quá lớn, tối đa {MAX_XLSX_BYTES // (1024 * 1024)}MB."}, 413)
    upload = request.files.get("quote")
    filename = (upload.filename or "").strip() if upload else ""
    if not upload or not filename:
        return _json({"ok": False, "error": "Vui lòng chọn file báo giá định dạng .xlsx."}, 400)
    if not filename.lower().endswith(".xlsx"):
        return _json({"ok": False, "error": f"'{filename}' không phải file .xlsx."}, 400)
    raw = upload.stream.read(MAX_XLSX_BYTES + 1)
    mapping = _json_field("mapping")
    header_row = _int_field("header_row")
    try:
        result = parse_quote(raw, sheet=request.form.get("sheet") or None, header_row=header_row, mapping=mapping)
    except QuoteParseError as exc:
        return _json({"ok": False, "error": f"{filename}: {exc}"}, 400)
    except Exception:
        logger.exception("quote_to_order: unexpected parse failure")
        return _json({"ok": False, "error": f"{filename}: không đọc được file báo giá."}, 400)
    return _json({"ok": True, "filename": filename, **result})


@bp.post("/export")
def export():
    denied = _guard(api=True)
    if denied is not None:
        return denied
    if request.content_length and request.content_length > MAX_EXPORT_BODY_BYTES:
        return _json({"ok": False, "error": "Dữ liệu gửi lên quá lớn."}, 413)
    body = request.get_json(silent=True)
    rows = body.get("rows") if isinstance(body, dict) else None
    if not isinstance(rows, list) or not rows:
        return _json({"ok": False, "error": "Chưa chọn dòng hàng nào."}, 400)
    if len(rows) > MAX_ORDER_ROWS:
        return _json({"ok": False, "error": f"Tối đa {MAX_ORDER_ROWS} dòng mỗi file đơn hàng."}, 400)
    clean, errors = clean_order_rows(rows)
    if errors:
        return _json({"ok": False, "error": "Có dòng chưa hợp lệ, hãy sửa rồi tải lại.", "rows": errors}, 400)
    try:
        data = build_order_file(clean)
    except Exception:
        logger.exception("quote_to_order: cannot build order workbook")
        return _json({"ok": False, "error": "Không tạo được file đơn hàng. Vui lòng thử lại hoặc báo quản trị."}, 500)
    name = f"bang-hang-hoa_{datetime.now():%Y%m%d_%H%M}.xlsx"
    response = send_file(io.BytesIO(data), as_attachment=True, download_name=name, mimetype=XLSX_MIME)
    response.headers["Cache-Control"] = "no-store, private"
    return response


def _json_field(name: str):
    raw = request.form.get(name)
    if not raw:
        return None
    try:
        return json.loads(raw)
    except ValueError:
        return None


def _int_field(name: str) -> int | None:
    raw = (request.form.get(name) or "").strip()
    return int(raw) if raw.isdigit() and len(raw) <= 6 else None
