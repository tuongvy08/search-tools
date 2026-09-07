from __future__ import annotations

import copy
import posixpath
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from xml.etree import ElementTree as ET

from openpyxl.formula.tokenizer import Token, Tokenizer, TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ENTRIES = 1000

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_MARKUP_COMPAT = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"

CALC_CHAIN_PATH = "xl/calcChain.xml"
CALC_CHAIN_REL_TYPE = NS_REL + "/calcChain"
CALC_CHAIN_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"

XML_DECLARATION = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
IGNORABLE_ATTR = f"{{{NS_MARKUP_COMPAT}}}Ignorable"
XML_NS = "http://www.w3.org/XML/1998/namespace"
GENERATED_PREFIX_RE = re.compile(r"^ns\d+$")
ROOT_START_TAG_RE = re.compile(rb"<[A-Za-z_][^>]*>")
XMLNS_DECL_RE = re.compile(rb'xmlns:([^\s=]+)\s*=')

ET.register_namespace("", NS_MAIN)
ET.register_namespace("r", NS_REL)
ET.register_namespace("mc", NS_MARKUP_COMPAT)

BG_HEADERS = {
    "B": "Tên hàng",
    "C": "Code",
    "D": "Cas",
    "E": "Hãng",
    "F": "Đơn vị tính",
    "M": "Ghi chú hàng hóa",
    "N": "Ghi chú khác",
    "P": "Giá nhập chưa VAT",
}
TEXT_COLUMNS = {
    "B": "Name",
    "C": "Code",
    "D": "Cas",
    "E": "Brand",
    "F": "Size",
    "M": "Note",
    "N": "Compliance_Combined",
}
MAPPED_COLUMNS = ("A", "B", "C", "D", "E", "F", "M", "N", "P")
PRODUCT_START_ROW = 17


class WorkbookExportError(ValueError):
    pass


@dataclass
class TemplateInfo:
    sheet_path: str
    total_row: int
    capacity: int


def export_quick_quote_workbook(raw: bytes, products: list[dict], mapping: dict | None = None) -> bytes:
    legacy_mapping = mapping is None
    mapping = _normalized_mapping(mapping)
    entries = _read_valid_xlsx_entries(raw)
    shared_strings = _read_shared_strings(entries)
    workbook_path, sheet_path = _find_worksheet_path(entries, mapping["sheet"])
    sheet_root, sheet_namespaces = _parse_xml_part(entries[sheet_path])
    total_row = _find_total_row(sheet_root, shared_strings, mapping["total_label"], mapping["product_start_row"])
    if legacy_mapping:
        _validate_bg_template(sheet_root, shared_strings, total_row)
    else:
        _validate_mapped_template(sheet_root, shared_strings, total_row, mapping)

    product_count = len(products)
    capacity = total_row - mapping["product_start_row"]
    if capacity < 1:
        raise WorkbookExportError("Template BG không có vùng dòng sản phẩm hợp lệ trước dòng Tổng giá.")
    if product_count > capacity:
        inserted = product_count - capacity
        _insert_product_rows(entries, sheet_path, mapping["sheet"], sheet_root, total_row, inserted)
        _shift_workbook_defined_names(entries, workbook_path, mapping["sheet"], total_row, inserted)
        _shift_cross_part_formula_references(
            entries, sheet_path, mapping["sheet"], total_row, inserted
        )
        total_row += inserted

    _write_products(sheet_root, products, total_row, mapping)
    _update_footer_formulas(sheet_root, total_row, product_count, mapping)
    _strip_formula_cached_values(sheet_root)
    _update_dimension(sheet_root)
    entries[sheet_path] = _serialize_part(sheet_root, sheet_namespaces)
    _strip_all_formula_caches(entries, exclude_sheet_path=sheet_path)
    _drop_calc_chain(entries, workbook_path)
    _force_full_recalculation(entries, workbook_path)
    return _write_xlsx_entries(entries)


def inspect_bg_template(raw: bytes, mapping: dict | None = None) -> TemplateInfo:
    legacy_mapping = mapping is None
    mapping = _normalized_mapping(mapping)
    entries = _read_valid_xlsx_entries(raw)
    shared_strings = _read_shared_strings(entries)
    _workbook_path, sheet_path = _find_worksheet_path(entries, mapping["sheet"])
    sheet_root = _parse_xml(entries[sheet_path])
    total_row = _find_total_row(sheet_root, shared_strings, mapping["total_label"], mapping["product_start_row"])
    if legacy_mapping:
        _validate_bg_template(sheet_root, shared_strings, total_row)
    else:
        _validate_mapped_template(sheet_root, shared_strings, total_row, mapping)
    return TemplateInfo(sheet_path=sheet_path, total_row=total_row, capacity=total_row - mapping["product_start_row"])


def _normalized_mapping(mapping: dict | None) -> dict:
    if mapping is None:
        return {
            "sheet": "BG", "header_row": 16, "product_start_row": 17,
            "total_label": "Tổng giá", "total_formula_column": "J",
            "mapping": {
                "sequence": "A", "Name": "B", "Code": "C", "Cas": "D",
                "Brand": "E", "Size": "F", "Note": "M",
                "Compliance_Combined": "N", "Unit_Price_Value": "P",
            },
        }
    return mapping


def _read_valid_xlsx_entries(raw: bytes) -> dict[str, bytes]:
    if not isinstance(raw, (bytes, bytearray)) or len(raw) < 64 or raw[:2] != b"PK":
        raise WorkbookExportError("File không phải .xlsx OOXML hợp lệ.")
    if len(raw) > MAX_XLSX_BYTES:
        raise WorkbookExportError(f"File .xlsx quá lớn, tối đa {MAX_XLSX_BYTES // (1024 * 1024)}MB.")
    try:
        with zipfile.ZipFile(BytesIO(raw), "r") as zf:
            infos = zf.infolist()
            if len(infos) > MAX_XLSX_ENTRIES:
                raise WorkbookExportError("File .xlsx có quá nhiều ZIP entry.")
            total_uncompressed = sum(info.file_size for info in infos)
            if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise WorkbookExportError("File .xlsx giải nén quá lớn, từ chối để tránh zip bomb.")
            names = [info.filename for info in infos]
            _validate_zip_names(names)
            if "[Content_Types].xml" not in names:
                raise WorkbookExportError("File thiếu [Content_Types].xml, không phải .xlsx OOXML chuẩn.")
            lower_names = {name.lower() for name in names}
            if "xl/vbaproject.bin" in lower_names or any(name.endswith(".bin") and "vba" in name for name in lower_names):
                raise WorkbookExportError("Không hỗ trợ workbook có macro/VBA.")
            content_types = zf.read("[Content_Types].xml")
            if b"macroEnabled" in content_types or b"vbaProject" in content_types:
                raise WorkbookExportError("Không hỗ trợ workbook có macro/VBA.")
            entries = {info.filename: zf.read(info.filename) for info in infos}
            _reject_active_workbook_content(entries)
    except zipfile.BadZipFile as exc:
        raise WorkbookExportError("File .xlsx bị hỏng hoặc không phải ZIP hợp lệ.") from exc
    return entries


def _reject_active_workbook_content(entries: dict[str, bytes]) -> None:
    unsafe_path_fragments = (
        "xl/externallinks/", "xl/connections", "xl/activex/",
        "xl/embeddings/", "xl/oleobjects/", "xl/ctrlprops/",
    )
    if any(any(fragment in path.lower() for fragment in unsafe_path_fragments) for path in entries):
        raise WorkbookExportError("Template chứa external link, connection hoặc active/embedded content không được hỗ trợ.")
    unsafe_rel_suffixes = ("/externalLink", "/oleObject", "/package", "/control", "/attachedTemplate")
    for path, raw in entries.items():
        if not path.endswith(".rels"):
            continue
        root = _parse_xml(raw)
        for rel in root.findall(f"{{{NS_PKG_REL}}}Relationship"):
            rel_type = rel.attrib.get("Type", "")
            if rel_type.endswith(unsafe_rel_suffixes):
                raise WorkbookExportError("Template chứa relationship chủ động/nhúng không được hỗ trợ.")
            if rel.attrib.get("TargetMode") == "External" and not rel_type.endswith("/hyperlink"):
                raise WorkbookExportError("Template chứa external relationship không an toàn.")
    for path, raw in entries.items():
        if not path.startswith("xl/worksheets/") or not path.endswith(".xml"):
            continue
        root = _parse_xml(raw)
        if any(_local_name(element.tag) in {"oleObjects", "controls"} for element in root.iter()):
            raise WorkbookExportError("Template chứa OLE/control không được hỗ trợ.")


def _validate_zip_names(names: list[str]) -> None:
    for name in names:
        if not name or name.startswith("/") or "\\" in name:
            raise WorkbookExportError("File .xlsx có ZIP entry không hợp lệ.")
        parts = name.split("/")
        if any(part in {"", ".", ".."} for part in parts):
            raise WorkbookExportError("File .xlsx có ZIP path không an toàn.")


def _parse_xml(raw: bytes) -> ET.Element:
    _reject_unsafe_xml(raw)
    try:
        return ET.fromstring(raw)
    except ET.ParseError as exc:
        raise WorkbookExportError("XML trong workbook bị lỗi.") from exc


def _reject_unsafe_xml(raw: bytes) -> None:
    upper = raw.upper()
    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
        raise WorkbookExportError("XML trong workbook không hợp lệ.")


def _parse_xml_part(raw: bytes) -> tuple[ET.Element, dict[str, str]]:
    """Parse an OOXML part and keep its original namespace prefix declarations."""
    _reject_unsafe_xml(raw)
    declarations: dict[str, str] = {}
    root = None
    try:
        for event, payload in ET.iterparse(BytesIO(raw), events=("start-ns", "start")):
            if event == "start-ns":
                prefix, uri = payload
                declarations.setdefault(prefix, uri)
            elif root is None:
                root = payload
    except ET.ParseError as exc:
        raise WorkbookExportError("XML trong workbook bị lỗi.") from exc
    if root is None:
        raise WorkbookExportError("XML trong workbook bị lỗi.")
    return root, declarations


def _serialize_part(root: ET.Element, declarations: dict[str, str]) -> bytes:
    """Serialize a part so Excel still sees the original prefixes and mc:Ignorable stays declared."""
    _sanitize_ignorable(root, declarations)
    _restore_namespace_declarations(root, declarations)
    data = XML_DECLARATION + ET.tostring(root, encoding="utf-8", xml_declaration=False)
    _validate_markup_compatibility(data)
    return data


def _sanitize_ignorable(root: ET.Element, declarations: dict[str, str]) -> None:
    """Drop mc:Ignorable tokens the source part never declared, so output stays valid."""
    value = root.get(IGNORABLE_ATTR)
    if not value:
        return
    tokens = value.split()
    kept = [token for token in tokens if token in declarations]
    if kept == tokens:
        return
    if kept:
        root.set(IGNORABLE_ATTR, " ".join(kept))
    else:
        del root.attrib[IGNORABLE_ATTR]


def _restore_namespace_declarations(root: ET.Element, declarations: dict[str, str]) -> None:
    serialized_prefixes: dict[str, str] = {}
    for prefix, uri in declarations.items():
        if prefix == "xml" or GENERATED_PREFIX_RE.match(prefix):
            continue
        ET.register_namespace(prefix, uri)
        serialized_prefixes[uri] = prefix

    used = _used_namespace_uris(root)
    for prefix, uri in declarations.items():
        if not prefix or prefix == "xml" or GENERATED_PREFIX_RE.match(prefix):
            continue
        if uri in used and serialized_prefixes.get(uri) == prefix:
            continue
        root.set(f"xmlns:{prefix}", uri)


def _used_namespace_uris(root: ET.Element) -> set[str]:
    used: set[str] = set()
    for element in root.iter():
        if isinstance(element.tag, str) and element.tag.startswith("{"):
            used.add(element.tag[1:].partition("}")[0])
        for name in element.attrib:
            if name.startswith("{"):
                used.add(name[1:].partition("}")[0])
    used.discard(XML_NS)
    return used


def _validate_markup_compatibility(data: bytes) -> None:
    root = _parse_xml(data)
    ignorable = root.get(IGNORABLE_ATTR, "").split()
    if not ignorable:
        return
    declared = _root_declared_prefixes(data)
    missing = [token for token in ignorable if token not in declared]
    if missing:
        raise WorkbookExportError(
            "XML xuất ra thiếu khai báo namespace cho mc:Ignorable: " + " ".join(missing)
        )


def _root_declared_prefixes(data: bytes) -> set[str]:
    match = ROOT_START_TAG_RE.search(data)
    if not match:
        raise WorkbookExportError("XML xuất ra không có root element hợp lệ.")
    return {prefix.decode("utf-8") for prefix in XMLNS_DECL_RE.findall(match.group(0))}


def _find_worksheet_path(entries: dict[str, bytes], sheet_name: str) -> tuple[str, str]:
    rels_root = _parse_xml(entries.get("_rels/.rels", b""))
    office_target = None
    for rel in rels_root.findall(f"{{{NS_PKG_REL}}}Relationship"):
        rel_type = rel.attrib.get("Type", "")
        if rel_type.endswith("/officeDocument"):
            office_target = rel.attrib.get("Target")
            break
    workbook_path = _normalize_target("", office_target or "xl/workbook.xml")
    if workbook_path not in entries:
        raise WorkbookExportError("Workbook thiếu xl/workbook.xml.")

    workbook_root = _parse_xml(entries[workbook_path])
    rels = _relationship_map(entries.get(_rels_path_for(workbook_path), b""))
    for sheet in workbook_root.findall(f".//{{{NS_MAIN}}}sheet"):
        if sheet.attrib.get("name") != sheet_name:
            continue
        rel_id = sheet.attrib.get(f"{{{NS_REL}}}id")
        target = rels.get(rel_id)
        if not target:
            raise WorkbookExportError(f"Sheet {sheet_name} không có relationship hợp lệ.")
        sheet_path = _normalize_target(posixpath.dirname(workbook_path), target)
        if sheet_path not in entries:
            raise WorkbookExportError(f"Không tìm thấy worksheet XML cho sheet {sheet_name}.")
        return workbook_path, sheet_path
    raise WorkbookExportError("Template không có sheet BG.")


def _rels_path_for(part_path: str) -> str:
    return posixpath.join(posixpath.dirname(part_path), "_rels", posixpath.basename(part_path) + ".rels")


def _drop_calc_chain(entries: dict[str, bytes], workbook_path: str) -> None:
    """Stale calcChain entries make Excel drop formulas; remove the part, its rel and content type."""
    calc_chain_paths = {CALC_CHAIN_PATH}
    rels_path = _rels_path_for(workbook_path)
    rels_raw = entries.get(rels_path)
    if rels_raw:
        rels_root, rels_namespaces = _parse_xml_part(rels_raw)
        removed_rels = False
        for rel in list(rels_root.findall(f"{{{NS_PKG_REL}}}Relationship")):
            if rel.attrib.get("Type") != CALC_CHAIN_REL_TYPE:
                continue
            target = _normalize_target(posixpath.dirname(workbook_path), rel.attrib.get("Target", ""))
            if target:
                calc_chain_paths.add(target)
            rels_root.remove(rel)
            removed_rels = True
        if removed_rels:
            entries[rels_path] = _serialize_part(rels_root, rels_namespaces)

    for path in calc_chain_paths:
        entries.pop(path, None)
    _drop_calc_chain_content_types(entries, calc_chain_paths)


def _drop_calc_chain_content_types(entries: dict[str, bytes], calc_chain_paths: set[str]) -> None:
    raw = entries.get("[Content_Types].xml")
    if not raw:
        return
    root, namespaces = _parse_xml_part(raw)
    part_names = {"/" + path for path in calc_chain_paths}
    removed = False
    for override in list(root.findall(f"{{{NS_CONTENT_TYPES}}}Override")):
        if override.attrib.get("PartName") in part_names or override.attrib.get("ContentType") == CALC_CHAIN_CONTENT_TYPE:
            root.remove(override)
            removed = True
    if removed:
        entries["[Content_Types].xml"] = _serialize_part(root, namespaces)


def _force_full_recalculation(entries: dict[str, bytes], workbook_path: str) -> None:
    raw = entries.get(workbook_path)
    if not raw:
        return
    root, namespaces = _parse_xml_part(raw)
    calc_pr = root.find(f"{{{NS_MAIN}}}calcPr")
    if calc_pr is None:
        calc_pr = ET.Element(f"{{{NS_MAIN}}}calcPr")
        ext_lst = root.find(f"{{{NS_MAIN}}}extLst")
        children = list(root)
        root.insert(children.index(ext_lst) if ext_lst is not None else len(children), calc_pr)
    calc_pr.attrib["calcMode"] = "auto"
    calc_pr.attrib["fullCalcOnLoad"] = "1"
    calc_pr.attrib["forceFullCalc"] = "1"
    entries[workbook_path] = _serialize_part(root, namespaces)


def _relationship_map(raw: bytes) -> dict[str, str]:
    root = _parse_xml(raw)
    out = {}
    for rel in root.findall(f"{{{NS_PKG_REL}}}Relationship"):
        rel_id = rel.attrib.get("Id")
        target = rel.attrib.get("Target")
        if rel_id and target:
            out[rel_id] = target
    return out


def _normalize_target(base_dir: str, target: str) -> str:
    if not target:
        return ""
    path = target.lstrip("/") if target.startswith("/") else posixpath.join(base_dir, target)
    return posixpath.normpath(path)


def _read_shared_strings(entries: dict[str, bytes]) -> list[str]:
    raw = entries.get("xl/sharedStrings.xml")
    if not raw:
        return []
    root = _parse_xml(raw)
    out = []
    for si in root.findall(f"{{{NS_MAIN}}}si"):
        texts = [t.text or "" for t in si.findall(f".//{{{NS_MAIN}}}t")]
        out.append("".join(texts))
    return out


def _cell_text(cell: ET.Element | None, shared_strings: list[str]) -> str:
    if cell is None:
        return ""
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(t.text or "" for t in cell.findall(f".//{{{NS_MAIN}}}t")).strip()
    value = cell.find(f"{{{NS_MAIN}}}v")
    if value is None or value.text is None:
        return ""
    if cell_type == "s":
        try:
            return shared_strings[int(value.text)].strip()
        except (ValueError, IndexError):
            return ""
    return str(value.text).strip()


def _find_total_row(sheet_root: ET.Element, shared_strings: list[str], label: str = "Tổng giá", start_row: int = 1) -> int:
    for row in _sheet_data(sheet_root).findall(f"{{{NS_MAIN}}}row"):
        for cell in row.findall(f"{{{NS_MAIN}}}c"):
            if int(row.attrib.get("r", "0")) >= start_row and _cell_text(cell, shared_strings).casefold() == label.casefold():
                return int(row.attrib.get("r", "0"))
    raise WorkbookExportError("Template BG thiếu dòng Tổng giá.")


def _cell_formula(cell: ET.Element | None) -> str:
    if cell is None:
        return ""
    formula = cell.find(f"{{{NS_MAIN}}}f")
    return "" if formula is None else (formula.text or "")


def _validate_mapped_template(sheet_root: ET.Element, shared_strings: list[str], total_row: int, mapping: dict) -> None:
    missing = []
    for field in ("Name", "Code"):
        col = mapping["mapping"].get(field)
        actual = _cell_text(_get_cell(sheet_root, f"{col}{mapping['header_row']}") if col else None, shared_strings)
        if not col or not actual:
            missing.append(field)
    if missing:
        raise WorkbookExportError("Mapping thiếu header bắt buộc: " + ", ".join(missing))
    if total_row <= mapping["product_start_row"]:
        raise WorkbookExportError("Dòng Tổng giá nằm trước vùng sản phẩm.")
    _validate_footer_formulas(sheet_root, total_row, mapping.get("total_formula_column", "J"))


def _validate_bg_template(sheet_root: ET.Element, shared_strings: list[str], total_row: int) -> None:
    missing = []
    for col, expected in BG_HEADERS.items():
        actual = _cell_text(_get_cell(sheet_root, f"{col}16"), shared_strings)
        if actual != expected:
            missing.append(f"{col}16 cần '{expected}', hiện là '{actual or '<trống>'}'")
    if missing:
        raise WorkbookExportError("Header sheet BG không đúng: " + "; ".join(missing))
    if total_row <= PRODUCT_START_ROW:
        raise WorkbookExportError("Dòng Tổng giá nằm trước vùng sản phẩm.")
    _validate_footer_formulas(sheet_root, total_row)


def _validate_footer_formulas(sheet_root: ET.Element, total_row: int, formula_col: str = "J") -> None:
    footer_refs = {
        "Tổng giá": f"{formula_col}{total_row}",
        "VAT": f"{formula_col}{total_row + 1}",
        "Tổng gồm VAT": f"{formula_col}{total_row + 2}",
    }
    errors = []
    for label, ref in footer_refs.items():
        cell = _get_cell(sheet_root, ref)
        if cell is None:
            errors.append(f"{ref} thiếu công thức {label}")
            continue
        formula = _cell_formula(cell)
        if not formula:
            errors.append(f"{ref} phải là công thức {label}, không phải giá trị tĩnh")
    if errors:
        raise WorkbookExportError("Footer sheet BG không hợp lệ: " + "; ".join(errors))


def _sheet_data(sheet_root: ET.Element) -> ET.Element:
    sheet_data = sheet_root.find(f"{{{NS_MAIN}}}sheetData")
    if sheet_data is None:
        raise WorkbookExportError("Worksheet thiếu sheetData.")
    return sheet_data


def _get_cell(sheet_root: ET.Element, ref: str) -> ET.Element | None:
    row_number = _ref_row(ref)
    row = _find_row(sheet_root, row_number)
    if row is None:
        return None
    for cell in row.findall(f"{{{NS_MAIN}}}c"):
        if cell.attrib.get("r") == ref:
            return cell
    return None


def _find_row(sheet_root: ET.Element, row_number: int) -> ET.Element | None:
    for row in _sheet_data(sheet_root).findall(f"{{{NS_MAIN}}}row"):
        if int(row.attrib.get("r", "0")) == row_number:
            return row
    return None


def _ensure_row(sheet_root: ET.Element, row_number: int) -> ET.Element:
    sheet_data = _sheet_data(sheet_root)
    row = _find_row(sheet_root, row_number)
    if row is not None:
        return row
    row = ET.Element(f"{{{NS_MAIN}}}row", {"r": str(row_number)})
    rows = sheet_data.findall(f"{{{NS_MAIN}}}row")
    for idx, existing in enumerate(rows):
        if int(existing.attrib.get("r", "0")) > row_number:
            sheet_data.insert(idx, row)
            return row
    sheet_data.append(row)
    return row


def _ensure_cell(sheet_root: ET.Element, col: str, row_number: int) -> ET.Element:
    row = _ensure_row(sheet_root, row_number)
    ref = f"{col}{row_number}"
    for cell in row.findall(f"{{{NS_MAIN}}}c"):
        if cell.attrib.get("r") == ref:
            return cell
    cell = ET.Element(f"{{{NS_MAIN}}}c", {"r": ref})
    cells = row.findall(f"{{{NS_MAIN}}}c")
    target_col_index = _col_to_index(col)
    for idx, existing in enumerate(cells):
        existing_col = _ref_col(existing.attrib.get("r", "A1"))
        if _col_to_index(existing_col) > target_col_index:
            row.insert(idx, cell)
            return cell
    row.append(cell)
    return cell


def _insert_product_rows(entries: dict[str, bytes], sheet_path: str, sheet_name: str,
                         sheet_root: ET.Element, total_row: int, count: int) -> None:
    sheet_data = _sheet_data(sheet_root)
    template_row = _find_row(sheet_root, total_row - 1)
    if template_row is None:
        raise WorkbookExportError("Template BG thiếu dòng sản phẩm mẫu cuối.")
    template_row_source = copy.deepcopy(template_row)
    if any(formula.attrib.get("t") in {"shared", "array", "dataTable"}
           for formula in sheet_root.findall(f".//{{{NS_MAIN}}}f")):
        raise WorkbookExportError("Không thể mở rộng template có shared/array/data-table formula.")

    for row in sheet_data.findall(f"{{{NS_MAIN}}}row"):
        row_num = int(row.attrib.get("r", "0"))
        if row_num >= total_row:
            _shift_row(row, count, total_row)
    _shift_existing_sheet_formula_references(
        sheet_root, sheet_name, total_row, count, skip_row=total_row - 1
    )
    for formula in template_row_source.findall(f".//{{{NS_MAIN}}}f"):
        if formula.text:
            formula.text = _shift_formula_rows_for_sheet(
                formula.text, sheet_name, total_row, count,
                include_unqualified=True, expand_boundary=False,
            )
        if formula.attrib.get("ref"):
            formula.attrib["ref"] = _shift_range_rows(formula.attrib["ref"], total_row, count)

    insert_at = 0
    for idx, row in enumerate(sheet_data.findall(f"{{{NS_MAIN}}}row")):
        if int(row.attrib.get("r", "0")) > total_row - 1:
            insert_at = idx
            break
    else:
        insert_at = len(sheet_data)

    for offset in range(count):
        new_row_number = total_row + offset
        new_row = copy.deepcopy(template_row_source)
        _retarget_row(new_row, new_row_number)
        _translate_copied_row_formulas(new_row, new_row_number - (total_row - 1))
        sheet_data.insert(insert_at + offset, new_row)

    _shift_merged_ranges(sheet_root, total_row, count)
    _shift_dimension_refs(sheet_root, total_row, count)
    _shift_sheet_row_structures(sheet_root, total_row, count)
    _shift_related_row_structures(entries, sheet_path, total_row, count)


def _shift_row(row: ET.Element, amount: int, threshold_row: int) -> None:
    old_num = int(row.attrib.get("r", "0"))
    new_num = old_num + amount
    _retarget_row(row, new_num)


def _retarget_row(
    row: ET.Element,
    new_row_number: int,
) -> None:
    row.attrib["r"] = str(new_row_number)
    for cell in row.findall(f"{{{NS_MAIN}}}c"):
        ref = cell.attrib.get("r")
        if not ref:
            continue
        col = _ref_col(ref)
        cell.attrib["r"] = f"{col}{new_row_number}"


def _shift_formula_rows(formula: str, threshold_row: int, amount: int) -> str:
    return _rewrite_formula_range_tokens(
        formula, lambda row: row + amount if row >= threshold_row else row
    )


def _normalized_formula_sheet_name(worksheet_prefix: str) -> str:
    value = worksheet_prefix[:-1] if worksheet_prefix.endswith("!") else worksheet_prefix
    if len(value) >= 2 and value.startswith("'") and value.endswith("'"):
        value = value[1:-1].replace("''", "'")
    return value


def _shift_formula_rows_for_sheet(formula: str, sheet_name: str, threshold_row: int,
                                  amount: int, *, include_unqualified: bool,
                                  expand_boundary: bool = True) -> str:
    try:
        tokens = Tokenizer("=" + formula).items
    except TokenizerError as exc:
        raise WorkbookExportError("Công thức trong template không thể phân tích an toàn.") from exc
    out = []
    for token in tokens:
        if token.type != Token.OPERAND or token.subtype != Token.RANGE:
            out.append(token.value)
            continue
        worksheet, _coordinate = Translator.strip_ws_name(token.value)
        targets_sheet = (
            (not worksheet and include_unqualified)
            or (worksheet and _normalized_formula_sheet_name(worksheet) == sheet_name)
        )
        if targets_sheet:
            coordinate = token.value[len(worksheet):]
            rewritten = (_expand_range_at_insertion(coordinate, threshold_row, amount)
                         if expand_boundary else _shift_range_rows(coordinate, threshold_row, amount))
            out.append(worksheet + rewritten)
        else:
            out.append(token.value)
    return "".join(out)


def _shift_existing_sheet_formula_references(sheet_root: ET.Element, sheet_name: str,
                                             threshold_row: int, amount: int,
                                             *, skip_row: int | None = None) -> None:
    skipped = set()
    if skip_row is not None:
        row = _find_row(sheet_root, skip_row)
        if row is not None:
            skipped = {id(formula) for formula in row.findall(f".//{{{NS_MAIN}}}f")}
    for formula in sheet_root.findall(f".//{{{NS_MAIN}}}f"):
        if formula.text:
            formula.text = _shift_formula_rows_for_sheet(
                formula.text, sheet_name, threshold_row, amount, include_unqualified=True,
                expand_boundary=id(formula) not in skipped,
            )
        if formula.attrib.get("ref"):
            formula.attrib["ref"] = _shift_range_rows(formula.attrib["ref"], threshold_row, amount)


def _translate_copied_row_formulas(row: ET.Element, row_delta: int) -> None:
    """Apply Excel copy-down semantics to formulas in a cloned product row.

    Relative row references move with the cloned row; absolute row references
    stay fixed. Mapped input cells are overwritten later by _write_products.
    """
    for formula in row.findall(f".//{{{NS_MAIN}}}f"):
        if formula.text:
            try:
                translated = Translator("=" + formula.text, origin="A1").translate_formula(
                    row_delta=row_delta, col_delta=0
                )
            except (TokenizerError, TranslatorError, ValueError) as exc:
                raise WorkbookExportError("Không thể dịch công thức của dòng sản phẩm mẫu.") from exc
            formula.text = translated[1:] if translated.startswith("=") else translated


def _rewrite_formula_range_tokens(formula: str, row_transform) -> str:
    """Rewrite rows in A1 RANGE operands without touching quoted literals or names."""
    try:
        tokens = Tokenizer("=" + formula).items
    except TokenizerError as exc:
        raise WorkbookExportError("Công thức trong template không thể phân tích an toàn.") from exc
    out = []
    for token in tokens:
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            out.append(_rewrite_range_operand_rows(token.value, row_transform))
        else:
            out.append(token.value)
    return "".join(out)


def _rewrite_range_operand_rows(value: str, row_transform) -> str:
    worksheet, coordinate = Translator.strip_ws_name(value)
    parts = coordinate.split(":")
    rewritten = []
    cell_re = re.compile(r"^(\$?[A-Za-z]{1,3})(\$?)([1-9][0-9]{0,6})$")
    row_re = re.compile(r"^(\$?)([1-9][0-9]{0,6})$")
    for part in parts:
        match = cell_re.fullmatch(part)
        if match:
            rewritten.append(f"{match.group(1)}{match.group(2)}{row_transform(int(match.group(3)))}")
            continue
        match = row_re.fullmatch(part)
        if match:
            rewritten.append(f"{match.group(1)}{row_transform(int(match.group(2)))}")
            continue
        # Structured references and defined names are not A1 row references.
        rewritten.append(part)
    return worksheet + ":".join(rewritten)


def _shift_merged_ranges(sheet_root: ET.Element, threshold_row: int, amount: int) -> None:
    merge_cells = sheet_root.find(f"{{{NS_MAIN}}}mergeCells")
    if merge_cells is None:
        return
    existing = merge_cells.findall(f"{{{NS_MAIN}}}mergeCell")
    template_merges = []
    for merge_cell in existing:
        ref = merge_cell.attrib.get("ref", "")
        parts = ref.split(":")
        if len(parts) == 2 and all(_ref_row(part) == threshold_row - 1 for part in parts):
            template_merges.append(parts)
        merge_cell.attrib["ref"] = _shift_range_rows(ref, threshold_row, amount)
    for offset in range(1, amount + 1):
        for first, last in template_merges:
            cloned_ref = ":".join(
                f"{_ref_col(part)}{_ref_row(part) + offset}" for part in (first, last)
            )
            merge_cells.append(ET.Element(f"{{{NS_MAIN}}}mergeCell", {"ref": cloned_ref}))
    if "count" in merge_cells.attrib:
        merge_cells.attrib["count"] = str(len(merge_cells.findall(f"{{{NS_MAIN}}}mergeCell")))


def _shift_dimension_refs(sheet_root: ET.Element, threshold_row: int, amount: int) -> None:
    dimension = sheet_root.find(f"{{{NS_MAIN}}}dimension")
    if dimension is not None and dimension.attrib.get("ref"):
        dimension.attrib["ref"] = _shift_range_rows(dimension.attrib["ref"], threshold_row, amount)


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _shift_sqref(value: str, threshold_row: int, amount: int) -> str:
    return " ".join(_shift_range_rows(ref, threshold_row, amount) for ref in value.split())


def _expand_sqref_at_insertion(value: str, threshold_row: int, amount: int) -> str:
    return " ".join(_expand_range_at_insertion(ref, threshold_row, amount) for ref in value.split())


def _shift_sheet_row_structures(sheet_root: ET.Element, threshold_row: int, amount: int) -> None:
    expanding_ref_tags = {"autoFilter", "sortState"}
    shifting_ref_tags = {"hyperlink"}
    expanding_sqref_tags = {"conditionalFormatting", "dataValidation", "ignoredError", "protectedRange"}
    for element in sheet_root.iter():
        local = _local_name(element.tag)
        if local in expanding_ref_tags and element.attrib.get("ref"):
            element.attrib["ref"] = _expand_range_at_insertion(element.attrib["ref"], threshold_row, amount)
        if local in shifting_ref_tags and element.attrib.get("ref"):
            element.attrib["ref"] = _shift_range_rows(element.attrib["ref"], threshold_row, amount)
        if local in expanding_sqref_tags and element.attrib.get("sqref"):
            element.attrib["sqref"] = _expand_sqref_at_insertion(element.attrib["sqref"], threshold_row, amount)
        if local == "selection" and element.attrib.get("sqref"):
            element.attrib["sqref"] = _shift_sqref(element.attrib["sqref"], threshold_row, amount)
        if local == "pane" and element.attrib.get("topLeftCell"):
            element.attrib["topLeftCell"] = _shift_single_ref(
                element.attrib["topLeftCell"], threshold_row, amount
            )
        if local == "selection" and element.attrib.get("activeCell"):
            element.attrib["activeCell"] = _shift_single_ref(
                element.attrib["activeCell"], threshold_row, amount
            )
    for container in sheet_root.iter():
        if _local_name(container.tag) not in {"rowBreaks", "colBreaks"}:
            continue
        if _local_name(container.tag) == "colBreaks":
            continue
        for brk in container:
            if _local_name(brk.tag) == "brk" and brk.attrib.get("id", "").isdigit():
                row_number = int(brk.attrib["id"])
                if row_number >= threshold_row:
                    brk.attrib["id"] = str(row_number + amount)


def _shift_related_row_structures(entries: dict[str, bytes], sheet_path: str,
                                  threshold_row: int, amount: int) -> None:
    rels_raw = entries.get(_rels_path_for(sheet_path))
    if not rels_raw:
        return
    rels_root = _parse_xml(rels_raw)
    for rel in rels_root.findall(f"{{{NS_PKG_REL}}}Relationship"):
        rel_type = rel.attrib.get("Type", "")
        target = _normalize_target(posixpath.dirname(sheet_path), rel.attrib.get("Target", ""))
        if not target or target not in entries:
            continue
        if rel_type.endswith("/drawing"):
            _shift_drawing_anchors(entries, target, threshold_row, amount)
        elif rel_type.endswith("/vmlDrawing"):
            _shift_vml_anchors(entries, target, threshold_row, amount)
        elif rel_type.endswith("/table"):
            _shift_table_refs(entries, target, threshold_row, amount)
        elif rel_type.endswith("/comments"):
            _shift_comment_refs(entries, target, threshold_row, amount)


def _shift_drawing_anchors(entries: dict[str, bytes], path: str,
                           threshold_row: int, amount: int) -> None:
    root, namespaces = _parse_xml_part(entries[path])
    changed = False
    for marker in root.iter():
        if _local_name(marker.tag) not in {"from", "to"}:
            continue
        for child in marker:
            if _local_name(child.tag) != "row" or not str(child.text or "").isdigit():
                continue
            zero_based = int(child.text)
            if zero_based + 1 >= threshold_row:
                child.text = str(zero_based + amount)
                changed = True
    if changed:
        entries[path] = _serialize_part(root, namespaces)


def _shift_vml_anchors(entries: dict[str, bytes], path: str,
                       threshold_row: int, amount: int) -> None:
    root, namespaces = _parse_xml_part(entries[path])
    changed = False
    for anchor in root.iter():
        if _local_name(anchor.tag) != "Anchor" or not anchor.text:
            continue
        parts = [part.strip() for part in anchor.text.split(",")]
        if len(parts) != 8 or any(not re.fullmatch(r"-?\d+", part) for part in parts):
            continue
        for index in (2, 6):
            if int(parts[index]) + 1 >= threshold_row:
                parts[index] = str(int(parts[index]) + amount)
                changed = True
        anchor.text = ", ".join(parts)
    if changed:
        entries[path] = _serialize_part(root, namespaces)


def _shift_table_refs(entries: dict[str, bytes], path: str,
                      threshold_row: int, amount: int) -> None:
    root, namespaces = _parse_xml_part(entries[path])
    changed = False
    for element in root.iter():
        if _local_name(element.tag) in {"table", "autoFilter", "sortState"} and element.attrib.get("ref"):
            element.attrib["ref"] = _expand_range_at_insertion(element.attrib["ref"], threshold_row, amount)
            changed = True
    if changed:
        entries[path] = _serialize_part(root, namespaces)


def _shift_comment_refs(entries: dict[str, bytes], path: str,
                        threshold_row: int, amount: int) -> None:
    root, namespaces = _parse_xml_part(entries[path])
    changed = False
    for comment in root.iter():
        if _local_name(comment.tag) == "comment" and comment.attrib.get("ref"):
            comment.attrib["ref"] = _shift_single_ref(comment.attrib["ref"], threshold_row, amount)
            changed = True
    if changed:
        entries[path] = _serialize_part(root, namespaces)


def _shift_workbook_defined_names(entries: dict[str, bytes], workbook_path: str,
                                  sheet_name: str, threshold_row: int, amount: int) -> None:
    raw = entries.get(workbook_path)
    if not raw:
        return
    root, namespaces = _parse_xml_part(raw)
    changed = False
    quoted = "'" + sheet_name.replace("'", "''") + "'!"
    plain = sheet_name + "!"
    for defined_name in root.findall(f".//{{{NS_MAIN}}}definedName"):
        text = defined_name.text or ""
        if quoted not in text and plain not in text:
            continue
        defined_name.text = _shift_formula_rows_for_sheet(
            text, sheet_name, threshold_row, amount, include_unqualified=False
        )
        changed = changed or defined_name.text != text
    if changed:
        entries[workbook_path] = _serialize_part(root, namespaces)


def _shift_cross_part_formula_references(entries: dict[str, bytes], current_sheet_path: str,
                                         sheet_name: str, threshold_row: int, amount: int) -> None:
    for path in list(entries):
        if path == current_sheet_path or not path.startswith("xl/worksheets/") or not path.endswith(".xml"):
            continue
        root, namespaces = _parse_xml_part(entries[path])
        changed = False
        for formula in root.findall(f".//{{{NS_MAIN}}}f"):
            if not formula.text:
                continue
            updated = _shift_formula_rows_for_sheet(
                formula.text, sheet_name, threshold_row, amount, include_unqualified=False
            )
            if updated != formula.text:
                formula.text = updated
                changed = True
        if changed:
            _strip_formula_cached_values(root)
            entries[path] = _serialize_part(root, namespaces)

    for path in list(entries):
        if not path.startswith("xl/charts/") or not path.endswith(".xml"):
            continue
        root, namespaces = _parse_xml_part(entries[path])
        changed = False
        for parent in root.iter():
            direct_formulas = [child for child in parent if _local_name(child.tag) == "f" and child.text]
            parent_changed = False
            for formula in direct_formulas:
                updated = _shift_formula_rows_for_sheet(
                    formula.text, sheet_name, threshold_row, amount, include_unqualified=False
                )
                if updated != formula.text:
                    formula.text = updated
                    changed = parent_changed = True
            if parent_changed:
                for child in list(parent):
                    if _local_name(child.tag) in {"numCache", "strCache", "multiLvlStrCache"}:
                        parent.remove(child)
        if changed:
            entries[path] = _serialize_part(root, namespaces)


def _shift_range_rows(ref: str, threshold_row: int, amount: int) -> str:
    parts = ref.split(":")
    return ":".join(_shift_single_ref(part, threshold_row, amount) for part in parts)


def _expand_range_at_insertion(ref: str, threshold_row: int, amount: int) -> str:
    parts = ref.split(":")
    if len(parts) != 2:
        return _shift_range_rows(ref, threshold_row, amount)
    row_match = re.compile(r"^\$?[A-Z]{1,3}\$?(\d+)$", re.IGNORECASE)
    row_only_match = re.compile(r"^\$?(\d+)$")
    first = row_match.fullmatch(parts[0])
    last = row_match.fullmatch(parts[1])
    if not first or not last:
        first = row_only_match.fullmatch(parts[0])
        last = row_only_match.fullmatch(parts[1])
    if first and last and int(first.group(1)) < threshold_row and int(last.group(1)) == threshold_row - 1:
        return parts[0] + ":" + _shift_single_ref(parts[1], threshold_row - 1, amount)
    return _shift_range_rows(ref, threshold_row, amount)


def _shift_single_ref(ref: str, threshold_row: int, amount: int) -> str:
    match = re.fullmatch(r"(\$?[A-Z]{1,3})(\$?)(\d+)", ref)
    if match:
        row_num = int(match.group(3))
        if row_num >= threshold_row:
            row_num += amount
        return f"{match.group(1)}{match.group(2)}{row_num}"
    row_only = re.fullmatch(r"(\$?)(\d+)", ref)
    if row_only:
        row_num = int(row_only.group(2))
        if row_num >= threshold_row:
            row_num += amount
        return f"{row_only.group(1)}{row_num}"
    return ref


def _write_products(sheet_root: ET.Element, products: list[dict], total_row: int, mapping: dict) -> None:
    start_row = mapping["product_start_row"]
    field_columns = mapping["mapping"]
    product_end_row = max(start_row + len(products) - 1, start_row - 1)
    for index, product in enumerate(products, start=1):
        row_num = start_row + index - 1
        stt = product.get("STT")
        if stt is None:
            _write_number_cell(sheet_root, field_columns["sequence"], row_num, index)
        else:
            _write_stt_cell(sheet_root, row_num, stt, field_columns["sequence"])
        for key, col in field_columns.items():
            if key in {"sequence", "Unit_Price_Value"}:
                continue
            _write_text_cell(sheet_root, col, row_num, str(product.get(key) or ""))
        price = product.get("Unit_Price_Value")
        if price is None:
            # Placeholder rows have no price: clear the cell instead of writing
            # 0 so it stays visually blank while still contributing 0 to the
            # per-row total formula and the footer SUM.
            if field_columns.get("Unit_Price_Value"):
                _clear_cell(_ensure_cell(sheet_root, field_columns["Unit_Price_Value"], row_num))
        else:
            if field_columns.get("Unit_Price_Value"):
                _write_number_cell(sheet_root, field_columns["Unit_Price_Value"], row_num, price)

    for row_num in range(product_end_row + 1, total_row):
        for col in field_columns.values():
            cell = _get_cell(sheet_root, f"{col}{row_num}")
            if cell is not None:
                _clear_cell(cell)


def _write_stt_cell(sheet_root: ET.Element, row_number: int, stt, col: str = "A") -> None:
    """STT may be a plain integer (5) or a sub-ordinal string (5.1); write as
    number when possible, otherwise as inline text to preserve the label."""
    if isinstance(stt, (int, float)) and float(stt).is_integer():
        _write_number_cell(sheet_root, col, row_number, int(stt))
        return
    text = str(stt)
    if re.fullmatch(r"\d+", text):
        _write_number_cell(sheet_root, col, row_number, int(text))
    else:
        _write_text_cell(sheet_root, col, row_number, text)


def _write_text_cell(sheet_root: ET.Element, col: str, row_number: int, text: str) -> None:
    cell = _ensure_cell(sheet_root, col, row_number)
    _clear_cell(cell)
    cell.attrib["t"] = "inlineStr"
    inline = ET.SubElement(cell, f"{{{NS_MAIN}}}is")
    t = ET.SubElement(inline, f"{{{NS_MAIN}}}t")
    if text.strip() != text or "\n" in text:
        t.attrib["{http://www.w3.org/XML/1998/namespace}space"] = "preserve"
    t.text = text


def _write_number_cell(sheet_root: ET.Element, col: str, row_number: int, value) -> None:
    cell = _ensure_cell(sheet_root, col, row_number)
    _clear_cell(cell)
    cell.attrib.pop("t", None)
    v = ET.SubElement(cell, f"{{{NS_MAIN}}}v")
    if isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
        v.text = str(int(value))
    else:
        v.text = str(float(value))


def _write_formula_cell(sheet_root: ET.Element, col: str, row_number: int, formula: str) -> None:
    cell = _ensure_cell(sheet_root, col, row_number)
    _clear_cell(cell)
    cell.attrib.pop("t", None)
    f = ET.SubElement(cell, f"{{{NS_MAIN}}}f")
    f.text = formula


def _clear_cell(cell: ET.Element) -> None:
    style = cell.attrib.get("s")
    ref = cell.attrib.get("r")
    cell.attrib.clear()
    if ref:
        cell.attrib["r"] = ref
    if style is not None:
        cell.attrib["s"] = style
    for child in list(cell):
        cell.remove(child)


def _update_footer_formulas(sheet_root: ET.Element, total_row: int, product_count: int, mapping: dict) -> None:
    start_row = mapping["product_start_row"]
    formula_col = mapping.get("total_formula_column", "J")
    last_product_row = max(start_row + product_count - 1, start_row)
    _write_formula_cell(sheet_root, formula_col, total_row, f"SUM({formula_col}{start_row}:{formula_col}{last_product_row})")


def _strip_formula_cached_values(sheet_root: ET.Element) -> bool:
    """Never carry template-era calculated values into a newly priced quote."""
    changed = False
    for cell in sheet_root.findall(f".//{{{NS_MAIN}}}c"):
        if cell.find(f"{{{NS_MAIN}}}f") is None:
            continue
        for value in list(cell.findall(f"{{{NS_MAIN}}}v")):
            cell.remove(value)
            changed = True
    return changed


def _strip_all_formula_caches(entries: dict[str, bytes], *, exclude_sheet_path: str) -> None:
    for path in list(entries):
        if path == exclude_sheet_path or not path.startswith("xl/worksheets/") or not path.endswith(".xml"):
            continue
        root, namespaces = _parse_xml_part(entries[path])
        if _strip_formula_cached_values(root):
            entries[path] = _serialize_part(root, namespaces)
    for path in list(entries):
        if not path.startswith("xl/charts/") or not path.endswith(".xml"):
            continue
        root, namespaces = _parse_xml_part(entries[path])
        changed = False
        for parent in root.iter():
            for child in list(parent):
                if _local_name(child.tag) in {"numCache", "strCache", "multiLvlStrCache"}:
                    parent.remove(child)
                    changed = True
        if changed:
            entries[path] = _serialize_part(root, namespaces)


def _update_dimension(sheet_root: ET.Element) -> None:
    dimension = sheet_root.find(f"{{{NS_MAIN}}}dimension")
    if dimension is None:
        return
    max_row = 1
    max_col = 1
    for row in _sheet_data(sheet_root).findall(f"{{{NS_MAIN}}}row"):
        row_num = int(row.attrib.get("r", "0"))
        max_row = max(max_row, row_num)
        for cell in row.findall(f"{{{NS_MAIN}}}c"):
            ref = cell.attrib.get("r")
            if ref:
                max_col = max(max_col, _col_to_index(_ref_col(ref)))
    dimension.attrib["ref"] = f"A1:{_index_to_col(max_col)}{max_row}"


def _write_xlsx_entries(entries: dict[str, bytes]) -> bytes:
    out = BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return out.getvalue()


def _ref_col(ref: str) -> str:
    match = re.match(r"\$?([A-Z]{1,3})", ref)
    return match.group(1) if match else "A"


def _ref_row(ref: str) -> int:
    match = re.search(r"(\d+)$", ref)
    return int(match.group(1)) if match else 0


def _col_to_index(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch.upper()) - ord("A") + 1)
    return n


def _index_to_col(index: int) -> str:
    out = ""
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(ord("A") + rem) + out
    return out or "A"
