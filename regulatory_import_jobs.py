"""Durable background preview/apply jobs for Phase 6D1 regulatory rules."""

from __future__ import annotations

import hashlib
import os
from pathlib import Path
import threading
import time
import uuid

from openpyxl import load_workbook
from psycopg2.extras import Json, RealDictCursor, execute_values

import import_jobs
from import_engine import ImportProblem, inspect_workbook, limit
from regulatory import (
    EXPORT_ALLOW,
    MATCH_FIELD_LABELS,
    REGULATORY_LOCK_KEY,
    acquire_regulatory_lock,
    catalog_fingerprint,
    clean_text,
    normalize_match_value,
    normalized_identity,
    stable_key_for_label,
)


REGULATORY_JOB_LOCK_NAMESPACE = 624026
_HEADER_STATUS = "tình trạng quản lý"
_HEADER_NOTE = "ghi chú quản lý"
_FIELD_HEADERS = {"cas": "cas", "code": "code", "tên": "name"}


def connection():
    return import_jobs.connection()


def upload_path(job_id) -> Path:
    return import_jobs.upload_dir() / ("regulatory-" + str(uuid.UUID(str(job_id))) + ".xlsx")


def _event(cur, job_id, actor, name, detail=None):
    cur.execute(
        "INSERT INTO regulatory_import_events(job_id,actor,event,detail) VALUES (%s,%s,%s,%s)",
        (str(job_id), actor, name, Json(detail) if detail is not None else None),
    )


def fetch_job(cur, job_id):
    cur.execute("SELECT * FROM regulatory_import_jobs WHERE id=%s", (str(job_id),))
    return cur.fetchone()


def list_jobs():
    with connection() as conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM regulatory_import_jobs ORDER BY created_at DESC LIMIT 30")
        return cur.fetchall()


def submit(file, mode, actor, user_id, auth_version, submission_key):
    if mode not in ("upsert", "replace_scoped"):
        raise ImportProblem("Chế độ nhập quy tắc không hợp lệ.")
    try:
        submission_key = str(uuid.UUID(str(submission_key)))
    except (ValueError, TypeError, AttributeError):
        raise ImportProblem("Phiên upload không hợp lệ. Hãy tải lại trang.") from None
    filename = (file.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
    filename = "".join(char for char in filename if char.isprintable())[:180]
    if not filename.lower().endswith(".xlsx"):
        raise ImportProblem("Chỉ hỗ trợ Excel Workbook (.xlsx).")

    job_id = str(uuid.uuid4())
    path = upload_path(job_id)
    with connection() as conn:
        try:
            with conn, conn.cursor() as cur:
                acquire_regulatory_lock(cur)
                cur.execute(
                    "SELECT id FROM regulatory_import_jobs WHERE actor_user_id=%s AND submission_key=%s",
                    (user_id, submission_key),
                )
                old = cur.fetchone()
                if old:
                    return str(old[0])
                total = 0
                digest = hashlib.sha256()
                fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
                with os.fdopen(fd, "wb") as dest:
                    while True:
                        block = file.stream.read(1024 * 1024)
                        if not block:
                            break
                        total += len(block)
                        if total > limit("REGULATORY_MAX_BYTES", 16 * 1024**2):
                            raise ImportProblem("File quy tắc tối đa 16 MB.")
                        digest.update(block)
                        dest.write(block)
                    dest.flush()
                    os.fsync(dest.fileno())
                with path.open("rb") as source:
                    if source.read(4) != b"PK\x03\x04":
                        raise ImportProblem("Nội dung không phải XLSX.")
                cur.execute(
                    """INSERT INTO regulatory_import_jobs
                       (id,submission_key,actor,actor_user_id,actor_auth_version,filename,file_size,
                        file_sha256,mode,expires_at)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,now()+%s*interval '1 hour')""",
                    (job_id, submission_key, actor, user_id, auth_version, filename, total,
                     digest.hexdigest(), mode, limit("RETENTION_HOURS", 48)),
                )
                _event(cur, job_id, actor, "uploaded")
        except BaseException:
            path.unlink(missing_ok=True)
            raise
    return job_id


def control(job_id, action, actor, fingerprint="", confirm_delete=""):
    with connection() as conn, conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM regulatory_import_jobs WHERE id=%s FOR UPDATE", (str(job_id),))
        job = cur.fetchone()
        if not job:
            raise ImportProblem("Không tìm thấy tác vụ quy tắc.")
        if action == "cancel":
            if job["status"] not in ("queued", "running") or job["cancel_requested"]:
                return
            if job["status"] == "queued":
                cur.execute(
                    "UPDATE regulatory_import_jobs SET status='cancelled',cancel_requested=true,finished_at=now() WHERE id=%s",
                    (str(job_id),),
                )
            elif job["status"] == "running":
                cur.execute("UPDATE regulatory_import_jobs SET cancel_requested=true WHERE id=%s", (str(job_id),))
        elif action == "apply":
            cur.execute(
                "SELECT expires_at>now() AND purged_at IS NULL AS valid "
                "FROM regulatory_import_jobs WHERE id=%s",
                (str(job_id),),
            )
            if not cur.fetchone()["valid"]:
                raise ImportProblem("Tệp đã hết hạn; hãy upload lại.")
            preview = job["preview"] or {}
            if job["status"] != "completed" or job["phase"] != "preview" or not job["preview_ready"]:
                raise ImportProblem("Cần xem trước thành công trước khi ghi dữ liệu.")
            if fingerprint != preview.get("fingerprint"):
                raise ImportProblem("Xem trước đã thay đổi. Hãy tải lại trang.")
            if job["mode"] == "replace_scoped" and confirm_delete != str(preview.get("deleted", 0)):
                raise ImportProblem("Cần nhập chính xác số quy tắc sẽ xóa.")
            cur.execute(
                """UPDATE regulatory_import_jobs SET status='queued',phase='apply',preview_ready=false,
                   finished_at=NULL,processed_count=0 WHERE id=%s""",
                (str(job_id),),
            )
        elif action == "retry":
            if job["status"] not in ("failed", "cancelled"):
                raise ImportProblem("Không thể thử lại tác vụ này.")
            if job["attempts"] >= limit("MAX_ATTEMPTS", 6):
                raise ImportProblem("Đã hết số lần thử. Hãy upload lại workbook.")
            cur.execute(
                """UPDATE regulatory_import_jobs SET status='queued',phase='preview',preview_ready=false,
                   preview=NULL,cancel_requested=false,finished_at=NULL,errors='[]',error_count=0,
                   processed_count=0 WHERE id=%s""",
                (str(job_id),),
            )
        else:
            raise ImportProblem("Thao tác không hợp lệ.")
        _event(cur, job_id, actor, action)


def _parse_workbook(path: Path):
    inspect_workbook(path)
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if len(wb.worksheets) != 1:
            raise ImportProblem("Workbook quy tắc chỉ được có một trang tính.")
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        iterator = ws.iter_rows()
        try:
            header_cells = next(iterator)
        except StopIteration:
            raise ImportProblem("Workbook đang trống; không có phạm vi để xem trước.") from None
        if any(cell.data_type in ("f", "e") for cell in header_cells):
            raise ImportProblem("Tiêu đề không được chứa công thức.")
        headers = [clean_text(cell.value, max_chars=80).casefold() for cell in header_cells]
        if len(headers) != 3 or len(set(headers)) != 3:
            raise ImportProblem("Workbook phải có đúng ba cột không trùng nhau.")
        fields = [header for header in headers if header in _FIELD_HEADERS]
        if len(fields) != 1 or set(headers) != {fields[0], _HEADER_STATUS, _HEADER_NOTE}:
            raise ImportProblem(
                "Header phải là CAS|Code|Tên + Tình trạng quản lý + Ghi chú quản lý."
            )
        field_header = fields[0]
        field = _FIELD_HEADERS[field_header]
        indexes = {header: index for index, header in enumerate(headers)}
        rows = []
        for row_number, cells in enumerate(iterator, start=2):
            if row_number > limit("REGULATORY_MAX_ROWS", 100000) + 1:
                raise ImportProblem("File quy tắc vượt giới hạn số dòng.")
            if any(cell.data_type in ("f", "e") for cell in cells):
                raise ImportProblem(f"Dòng {row_number}: hãy thay công thức bằng giá trị.")
            values = [cell.value for cell in cells]
            if not any(clean_text(value) for value in values):
                continue
            try:
                value = normalize_match_value(field, values[indexes[field_header]])
                status_label = clean_text(values[indexes[_HEADER_STATUS]], max_chars=120)
                note = clean_text(values[indexes[_HEADER_NOTE]], max_chars=4000)
                if not status_label:
                    raise ImportProblem("Tình trạng quản lý không được để trống.")
            except (ValueError, ImportProblem) as exc:
                raise ImportProblem(f"Dòng {row_number}: {exc}") from exc
            rows.append({
                "row_number": row_number,
                "match_field": field,
                "match_value": value,
                "status_label": status_label,
                "note": note,
            })
        if not rows:
            raise ImportProblem("Workbook không có dòng quy tắc; không thể dùng file rỗng để xóa dữ liệu.")
        return rows
    finally:
        wb.close()


def _load_existing(cur):
    cur.execute("SELECT id,stable_key,label,priority,export_policy FROM regulatory_statuses ORDER BY priority,id")
    statuses = [dict(zip(("id", "stable_key", "label", "priority", "export_policy"), row)) for row in cur.fetchall()]
    cur.execute(
        """SELECT r.id,s.label,r.match_field,r.match_value,COALESCE(r.note,''),r.status_id
           FROM regulatory_rules r JOIN regulatory_statuses s ON s.id=r.status_id
           WHERE r.is_active=true"""
    )
    rules = [dict(zip(("id", "status_label", "match_field", "match_value", "note", "status_id"), row))
             for row in cur.fetchall()]
    return statuses, rules


def build_plan(cur, rows, mode):
    statuses, existing_rules = _load_existing(cur)
    status_by_norm = {normalized_identity(item["label"]): item for item in statuses}
    new_statuses = []
    seen_status = set(status_by_norm)
    seen_rows = {}
    prepared = []
    for row in rows:
        status_norm = normalized_identity(row["status_label"])
        if not status_norm:
            raise ImportProblem(f"Dòng {row['row_number']}: tình trạng không hợp lệ.")
        if status_norm not in seen_status:
            seen_status.add(status_norm)
            new_statuses.append(row["status_label"])
        value_norm = normalized_identity(row["match_value"])
        key = (status_norm, row["match_field"], value_norm)
        if key in seen_rows:
            raise ImportProblem(
                f"Dòng {row['row_number']}: trùng {MATCH_FIELD_LABELS[row['match_field']]} trong cùng tình trạng "
                f"với dòng {seen_rows[key]}; hãy giữ đúng một dòng."
            )
        seen_rows[key] = row["row_number"]
        prepared.append(dict(row, status_norm=status_norm, value_norm=value_norm))

    existing = {
        (normalized_identity(item["status_label"]), item["match_field"], normalized_identity(item["match_value"])): item
        for item in existing_rules
    }
    inserted = updated = 0
    changes = []
    for row in prepared:
        key = (row["status_norm"], row["match_field"], row["value_norm"])
        old = existing.get(key)
        if old is None:
            inserted += 1
            kind = "thêm"
        elif clean_text(old["note"]) != row["note"]:
            updated += 1
            kind = "cập nhật"
        else:
            kind = "giữ nguyên"
        if len(changes) < 20:
            changes.append({"action": kind, "field": row["match_field"], "value": row["match_value"],
                            "status": row["status_label"]})

    scopes = sorted({(row["match_field"], row["status_norm"]) for row in prepared})
    uploaded_keys = {(row["status_norm"], row["match_field"], row["value_norm"]) for row in prepared}
    delete_ids = []
    if mode == "replace_scoped":
        scope_set = set(scopes)
        delete_ids = [item["id"] for key, item in existing.items()
                      if (key[1], key[0]) in scope_set and key not in uploaded_keys]
    fingerprint = catalog_fingerprint(cur)
    plan_core = {
        "row_count": len(prepared), "inserted": inserted, "updated": updated,
        "deleted": len(delete_ids), "unchanged": len(prepared) - inserted - updated,
        "new_statuses": new_statuses,
        "scopes": [{"match_field": field,
                    "status": next(row["status_label"] for row in prepared if row["status_norm"] == status_norm)}
                   for field, status_norm in scopes],
        "sample": changes,
    }
    plan_digest = hashlib.sha256(repr((prepared, mode, plan_core)).encode("utf-8")).hexdigest()
    return dict(plan_core, fingerprint=fingerprint, plan_digest=plan_digest, _delete_ids=delete_ids,
                _prepared=prepared)


def apply_plan(cur, rows, mode, expected):
    current_fingerprint = catalog_fingerprint(cur)
    if current_fingerprint != expected.get("fingerprint"):
        raise ImportProblem("Danh mục hoặc quy tắc đã đổi từ lúc xem trước. Hãy xem trước lại.")
    plan = build_plan(cur, rows, mode)
    if plan["plan_digest"] != expected.get("plan_digest"):
        raise ImportProblem("Kế hoạch áp dụng không còn giống bản xem trước. Hãy xem trước lại.")

    if plan["new_statuses"]:
        cur.execute("SELECT COALESCE(max(priority),0) FROM regulatory_statuses")
        priority = int(cur.fetchone()[0])
        for label in plan["new_statuses"]:
            priority += 10
            cur.execute(
                """INSERT INTO regulatory_statuses(stable_key,label,priority,export_policy)
                   VALUES (%s,%s,%s,%s)""",
                (stable_key_for_label(label), label, priority, EXPORT_ALLOW),
            )

    if plan["_delete_ids"]:
        cur.execute("DELETE FROM regulatory_rules WHERE id=ANY(%s)", (plan["_delete_ids"],))
    for row in plan["_prepared"]:
        cur.execute(
            "SELECT id,stable_key,label,priority FROM regulatory_statuses WHERE upper(btrim(label))=upper(btrim(%s))",
            (row["status_label"],),
        )
        status_id, stable_key, label, priority = cur.fetchone()
        cur.execute(
            """INSERT INTO regulatory_rules
               (rule_type,rule_label,match_field,match_value,priority,is_active,note,status_id)
               VALUES (%s,%s,%s,%s,%s,true,%s,%s)
               ON CONFLICT (status_id,match_field,upper(btrim(match_value)))
               DO UPDATE SET note=EXCLUDED.note,match_value=EXCLUDED.match_value,
                             rule_label=EXCLUDED.rule_label,rule_type=EXCLUDED.rule_type,
                             priority=EXCLUDED.priority,is_active=true,updated_at=now()""",
            (stable_key, label, row["match_field"], row["match_value"], priority,
             row["note"] or None, status_id),
        )
    return plan


def _job_lock(cur, job_id):
    cur.execute(
        "SELECT pg_try_advisory_lock(hashtextextended(%s,%s))",
        (str(job_id), REGULATORY_JOB_LOCK_NAMESPACE),
    )
    return cur.fetchone()[0]


def _job_unlock(cur, job_id):
    cur.execute(
        "SELECT pg_advisory_unlock(hashtextextended(%s,%s))",
        (str(job_id), REGULATORY_JOB_LOCK_NAMESPACE),
    )


def _claim(conn, only_id=None):
    with conn.cursor() as cur:
        cur.execute(
            """SELECT id FROM regulatory_import_jobs
               WHERE status IN ('queued','running') AND (%s IS NULL OR id=%s::uuid)
               ORDER BY created_at LIMIT 100""",
            (only_id, only_id),
        )
        ids = [str(row[0]) for row in cur.fetchall()]
        conn.commit()
        for job_id in ids:
            if not _job_lock(cur, job_id):
                continue
            cur.execute(
                """SELECT status,phase,cancel_requested,attempts,expires_at<now()
                   FROM regulatory_import_jobs WHERE id=%s FOR UPDATE""",
                (job_id,),
            )
            row = cur.fetchone()
            if not row or row[0] not in ("queued", "running"):
                conn.rollback()
                _job_unlock(cur, job_id)
                continue
            status, phase, cancelled, attempts, expired = row
            if status == "running":
                cur.execute(
                    """UPDATE regulatory_import_jobs SET status='failed',preview_ready=false,
                       finished_at=now(),errors=%s,error_count=1 WHERE id=%s""",
                    (Json(["Worker bị gián đoạn. Hãy xem trước lại trước khi tiếp tục."]), job_id),
                )
                _event(cur, job_id, "worker", "crash_recovered")
                conn.commit()
                _job_unlock(cur, job_id)
                continue
            if cancelled or expired or attempts >= limit("MAX_ATTEMPTS", 6):
                cur.execute(
                    "UPDATE regulatory_import_jobs SET status='cancelled',finished_at=now() WHERE id=%s",
                    (job_id,),
                )
                conn.commit()
                _job_unlock(cur, job_id)
                continue
            cur.execute(
                """UPDATE regulatory_import_jobs SET status='running',started_at=now(),heartbeat_at=now(),
                   attempts=attempts+1 WHERE id=%s""",
                (job_id,),
            )
            _event(cur, job_id, "worker", phase + "_started")
            conn.commit()
            return job_id
    return None


def run_once(only_id=None):
    with connection() as conn:
        job_id = _claim(conn, only_id)
        if not job_id:
            return False
        stop = threading.Event()
        cancelled = threading.Event()
        deadline = time.monotonic() + limit("JOB_SECONDS", 7200)

        def watch():
            try:
                with connection() as monitor:
                    monitor.autocommit = True
                    with monitor.cursor() as cur:
                        while not stop.wait(1):
                            cur.execute(
                                """UPDATE regulatory_import_jobs SET heartbeat_at=now()
                                   WHERE id=%s AND purged_at IS NULL
                                   RETURNING cancel_requested,expires_at<=now()""",
                                (job_id,),
                            )
                            row = cur.fetchone()
                            if not row or row[0] or row[1] or time.monotonic() > deadline:
                                cancelled.set()
                                conn.cancel()
                                break
            except Exception:
                cancelled.set()
                conn.cancel()

        watcher = threading.Thread(target=watch, daemon=True)
        watcher.start()
        try:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                job = fetch_job(cur, job_id)
            conn.commit()
            path = upload_path(job_id)
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            if digest != job["file_sha256"]:
                raise ImportProblem("Tệp lưu trữ đã thay đổi. Hãy upload lại.")
            rows = _parse_workbook(path)
            if cancelled.is_set():
                raise ImportProblem("Tác vụ đã được hủy hoặc vượt thời gian xử lý.")
            with conn, conn.cursor() as cur:
                cur.execute("SET LOCAL statement_timeout = %s", (str(limit("SQL_SECONDS", 600) * 1000),))
                acquire_regulatory_lock(cur)
                if cancelled.is_set():
                    raise ImportProblem("Tác vụ đã được hủy hoặc vượt thời gian xử lý.")
                cur.execute(
                    """SELECT 1 FROM app_users WHERE id=%s AND is_admin=true
                       AND account_status='ACTIVE' AND auth_version=%s""",
                    (job["actor_user_id"], job["actor_auth_version"]),
                )
                if not cur.fetchone():
                    raise ImportProblem("Người tải tệp không còn quyền admin hoặc phiên đã bị thu hồi.")
                cur.execute("DELETE FROM regulatory_import_rows WHERE job_id=%s", (job_id,))
                execute_values(
                    cur,
                    "INSERT INTO regulatory_import_rows(job_id,row_number,data) VALUES %s",
                    [(job_id, row["row_number"], Json(row)) for row in rows],
                )
                plan = build_plan(cur, rows, job["mode"])
                if job["phase"] == "apply":
                    plan = apply_plan(cur, rows, job["mode"], job["preview"] or {})
                stop.set()
                watcher.join(timeout=5)
                if cancelled.is_set():
                    raise ImportProblem("Tác vụ đã được hủy hoặc vượt thời gian xử lý.")
                cur.execute(
                    """SELECT cancel_requested,expires_at<=now(),purged_at IS NOT NULL
                       FROM regulatory_import_jobs WHERE id=%s FOR UPDATE""",
                    (job_id,),
                )
                final_state = cur.fetchone()
                if not final_state or any(final_state):
                    raise ImportProblem("Tác vụ đã được hủy hoặc hết hạn.")
                public_plan = {key: value for key, value in plan.items() if not key.startswith("_")}
                cur.execute(
                    """UPDATE regulatory_import_jobs SET status='completed',preview_ready=%s,preview=%s,
                       row_count=%s,processed_count=%s,inserted_count=%s,updated_count=%s,
                       deleted_count=%s,finished_at=now(),errors='[]',error_count=0 WHERE id=%s""",
                    (job["phase"] == "preview", Json(public_plan), len(rows), len(rows),
                     plan["inserted"] if job["phase"] == "apply" else 0,
                     plan["updated"] if job["phase"] == "apply" else 0,
                     plan["deleted"] if job["phase"] == "apply" else 0, job_id),
                )
                _event(cur, job_id, "worker", job["phase"] + "_completed", public_plan)
        except Exception as exc:
            conn.rollback()
            message = str(exc)[:400] if isinstance(exc, (ImportProblem, ValueError)) else (
                "Xử lý quy tắc thất bại; chưa ghi dữ liệu."
            )
            with conn, conn.cursor() as cur:
                cur.execute(
                    """UPDATE regulatory_import_jobs
                       SET status=CASE WHEN cancel_requested THEN 'cancelled' ELSE 'failed' END,
                           preview_ready=false,errors=%s,error_count=1,finished_at=now() WHERE id=%s""",
                    (Json([message]), job_id),
                )
                _event(cur, job_id, "worker", "stopped")
        finally:
            stop.set()
            watcher.join(timeout=5)
            with conn.cursor() as cur:
                _job_unlock(cur, job_id)
            conn.commit()
        return True


def cleanup():
    with connection() as conn, conn.cursor() as cur:
        cur.execute(
            """SELECT id FROM regulatory_import_jobs
               WHERE expires_at<now() AND purged_at IS NULL ORDER BY expires_at LIMIT 100"""
        )
        ids = [str(row[0]) for row in cur.fetchall()]
        conn.commit()
        for job_id in ids:
            # Expiry requests cancellation, but cleanup never removes a live
            # worker's rows or upload: that worker owns the session job lock.
            cur.execute(
                """UPDATE regulatory_import_jobs SET cancel_requested=true
                   WHERE id=%s AND status='running' AND expires_at<now() AND purged_at IS NULL""",
                (job_id,),
            )
            conn.commit()
            if not _job_lock(cur, job_id):
                continue
            try:
                cur.execute(
                    """SELECT status,expires_at<now(),purged_at IS NOT NULL
                       FROM regulatory_import_jobs WHERE id=%s FOR UPDATE""",
                    (job_id,),
                )
                row = cur.fetchone()
                if not row or not row[1] or row[2]:
                    conn.rollback()
                    continue
                status = row[0]
                if status in ("queued", "running"):
                    status = "cancelled"
                    cur.execute(
                        """UPDATE regulatory_import_jobs SET status='cancelled',cancel_requested=true,
                           preview_ready=false,finished_at=COALESCE(finished_at,now()) WHERE id=%s""",
                        (job_id,),
                    )
                if status not in ("completed", "failed", "cancelled"):
                    conn.rollback()
                    continue
                cur.execute("DELETE FROM regulatory_import_rows WHERE job_id=%s", (job_id,))
                cur.execute(
                    """UPDATE regulatory_import_jobs SET purged_at=now(),preview=NULL,
                       preview_ready=false WHERE id=%s""",
                    (job_id,),
                )
                upload_path(job_id).unlink(missing_ok=True)
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                _job_unlock(cur, job_id)
                conn.commit()
