"""Disposable migration + 200k-row rehearsal; optional loopback browser server.

Requires DATABASE_URL pointing at a dedicated localhost PostgreSQL test server.
Only creates/deletes helper-prefixed temporary databases, never app databases.
"""
import argparse
import json
import os
from pathlib import Path
import resource
import signal
import sys
import tempfile
import threading
import time
import uuid
from urllib.parse import urlparse

import psycopg2
from psycopg2.extras import RealDictCursor
from openpyxl import Workbook
from werkzeug.datastructures import FileStorage
from werkzeug.security import generate_password_hash

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT),str(ROOT/'tests')]
from pg_temp_db import create_full_schema_temp_db,drop_temp_db,apply_brand_master_and_currency_migrations,apply_dynamic_brand_currency_migration,run_migration_via_psql


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--rows',type=int,default=200000);parser.add_argument('--serve',action='store_true');args=parser.parse_args()
    original=os.environ.get('DATABASE_URL','')
    if urlparse(original).hostname not in ('127.0.0.1','localhost') or urlparse(original).path!='/postgres':
        raise RuntimeError('Use a dedicated loopback maintenance database only')
    name,dsn=create_full_schema_temp_db();report={};stop=threading.Event();worker=None
    with tempfile.TemporaryDirectory(prefix='phase6c2-uploads-') as directory:
        try:
            conn=psycopg2.connect(dsn);conn.autocommit=True
            with conn.cursor() as cur:
                apply_brand_master_and_currency_migrations(cur);apply_dynamic_brand_currency_migration(cur)
                cur.execute((ROOT/'sql/migration_004_import_jobs.sql').read_text())
                cur.execute("INSERT INTO app_users(username,password_hash,is_admin,account_status) VALUES ('qa-admin',%s,true,'ACTIVE') RETURNING id",(generate_password_hash('qa-only'),));uid=cur.fetchone()[0]
            for _ in range(2):
                code,output=run_migration_via_psql(dsn,ROOT/'sql/migration_024_admin_import_center.sql')
                if code:raise RuntimeError(output)
            report['migration_twice_via_psql']=True
            os.environ.update(DATABASE_URL=dsn,IMPORT_UPLOAD_DIR=directory,FLASK_SECRET_KEY='disposable-phase6c2-browser',DISABLE_IP_ALLOWLIST='1',ENABLE_LEGACY_PASSWORD_LOGIN='0')
            import import_jobs as jobs
            def state(jid):
                with conn.cursor(cursor_factory=RealDictCursor) as cur:return jobs.fetch_job(cur,jid)
            path=Path(directory)/'benchmark.xlsx';wb=Workbook(write_only=True);ws=wb.create_sheet();ws.append(['brand','code','name','size','price'])
            for n in range(args.rows):ws.append(['TRC',f'QA-{n}',f'Sản phẩm tham chiếu {n}','1g',str(100+n)])
            wb.save(path);wb.close()
            report.update(rows=args.rows,workbook_bytes=path.stat().st_size)
            with path.open('rb') as source:
                jid=jobs.submit(FileStorage(source,filename='Danh-muc-TRC.xlsx',content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),'replace_by_brand','user:'+str(uid),uid,1,str(uuid.uuid4()))
            start=time.monotonic();jobs.run_once(jid);report['preview_seconds']=round(time.monotonic()-start,2)
            item=state(jid)
            if item['status']!='completed':raise RuntimeError(item['errors'])
            if not args.serve:
                p=item['preview'];jobs.control(jid,'apply','qa',p['fingerprint'],str(p['deleted']))
                start=time.monotonic();jobs.run_once(jid);report['apply_seconds']=round(time.monotonic()-start,2)
                done=state(jid)
                if done['status']!='completed':raise RuntimeError(done['errors'])
                report['inserted']=done['inserted_count']
                # Second pass measures large populated-catalog upsert and its fingerprint.
                with path.open('rb') as source:
                    second=jobs.submit(FileStorage(source,filename='Upsert-TRC.xlsx',content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),'upsert','user:'+str(uid),uid,1,str(uuid.uuid4()))
                start=time.monotonic();jobs.run_once(second);report['upsert_preview_seconds']=round(time.monotonic()-start,2)
                item=state(second)
                if item['status']!='completed':raise RuntimeError(item['errors'])
                p=item['preview'];jobs.control(second,'apply','qa',p['fingerprint'])
                start=time.monotonic();jobs.run_once(second);report['upsert_apply_seconds']=round(time.monotonic()-start,2)
                done=state(second)
                if done['status']!='completed':raise RuntimeError(done['errors'])
                report['updated']=done['updated_count']
                with conn.cursor() as cur:
                    cur.execute('SELECT count(*) FROM products');report['final_product_count']=cur.fetchone()[0]
                    cur.execute("UPDATE product_import_jobs SET expires_at=now()-interval '1 second'")
                jobs.cleanup()
                with conn.cursor() as cur:
                    cur.execute('SELECT count(*) FROM product_import_rows');report['remaining_staged_rows']=cur.fetchone()[0]
                report['remaining_uuid_uploads']=len([p for p in Path(directory).glob('*.xlsx') if p.name!='benchmark.xlsx'])
                report['python_peak_rss_mb']=round(resource.getrusage(resource.RUSAGE_SELF).ru_maxrss/(1024**2 if sys.platform=='darwin' else 1024),1)
            else:
                import search
                # Stable synthetic file for the browser's actual upload flow.
                sample=Path('/private/tmp/p6c2-browser.xlsx');wb=Workbook();ws=wb.active;ws.append(['brand','code','name','size','price']);ws.append(['Brand Mới QA','NEW-001','Mẫu chuẩn kiểm thử giao diện','5mg','250']);wb.save(sample);wb.close()
                def work():
                    while not stop.wait(1):jobs.run_once()
                worker=threading.Thread(target=work);worker.start()
                def finish(*_):raise SystemExit(0)
                signal.signal(signal.SIGTERM,finish);signal.signal(signal.SIGINT,finish)
                print('PHASE6C2_BROWSER http://127.0.0.1:5062 qa-admin / qa-only',flush=True)
                search.app.run(host='127.0.0.1',port=5062,use_reloader=False,debug=False)
        finally:
            stop.set()
            if worker:worker.join(15)
            if 'conn' in locals():conn.close()
            os.environ['DATABASE_URL']=original
            drop_temp_db(name)
            with psycopg2.connect(original) as check,check.cursor() as cur:
                cur.execute('SELECT count(*) FROM pg_database WHERE datname=%s',(name,));report['temporary_db_removed']=cur.fetchone()[0]==0
            print(json.dumps(report,ensure_ascii=False,indent=2),flush=True)


if __name__=='__main__':main()
