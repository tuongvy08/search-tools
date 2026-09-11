"""Durable background full-snapshot inventory imports for Phase 6D2."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
import os
from pathlib import Path
import threading
import time
import unicodedata
import uuid

from openpyxl import load_workbook
from psycopg2.extras import Json, RealDictCursor, execute_values

from brand_gateway import (
    acquire_products_import_lock,
    load_brand_gateway,
    preview_import_rows_brands,
    register_and_resolve_import_rows,
)
import import_jobs
from import_engine import ImportProblem, inspect_workbook, limit
from regulatory import normalize_cas
from stock import STOCK_LOCK_KEY, acquire_stock_lock, active_snapshot, normalized_text, snapshot_fingerprint


STOCK_JOB_LOCK_NAMESPACE = 624029
HEADERS = (
    "name", "code", "cas", "brand", "size", "giá tồn kho", "số lượng tồn", "hạn sử dụng",
)

HEADER_ALIASES = {
    "name": "name",
    "code": "code",
    "cas": "cas",
    "brand": "brand",
    "size": "size",
    "giá tồn kho": "giá tồn kho",
    "số lượng tồn": "số lượng tồn",
    "hạn sử dụng": "hạn sử dụng",
    "stock price": "giá tồn kho",
    "qty": "số lượng tồn",
    "expiry": "hạn sử dụng",
}

_HEADER_REQUIRES = "Name | Code | Cas | Brand | Size | Giá tồn kho | Số lượng tồn | Hạn sử dụng."


def connection():
    return import_jobs.connection()


def upload_path(job_id) -> Path:
    return import_jobs.upload_dir() / ("stock-" + str(uuid.UUID(str(job_id))) + ".xlsx")


def _event(cur, job_id, actor, event, detail=None):
    cur.execute(
        "INSERT INTO stock_import_events(job_id,actor,event,detail) VALUES (%s,%s,%s,%s)",
        (str(job_id), actor, event, Json(detail) if detail is not None else None),
    )


def fetch_job(cur, job_id):
    cur.execute("SELECT * FROM stock_import_jobs WHERE id=%s", (str(job_id),))
    return cur.fetchone()


def list_jobs():
    with connection() as conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM stock_import_jobs ORDER BY created_at DESC LIMIT 30")
        return cur.fetchall()


def list_snapshots():
    with connection() as conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute(
            """
            SELECT s.*,state.active_snapshot_id=s.id AS is_active
            FROM stock_snapshots s CROSS JOIN stock_state state
            WHERE state.singleton=TRUE
            ORDER BY s.created_at DESC LIMIT 30
            """
        )
        return cur.fetchall()


def submit(file, actor, user_id, auth_version, submission_key):
    try:
        submission_key = str(uuid.UUID(str(submission_key)))
    except (ValueError, TypeError, AttributeError):
        raise ImportProblem("Phiên upload không hợp lệ. Hãy tải lại trang.") from None
    filename = (file.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
    filename = "".join(char for char in filename if char.isprintable())[:180]
    if not filename.lower().endswith(".xlsx"):
        raise ImportProblem("Chỉ hỗ trợ Excel Workbook (.xlsx).")
    job_id = str(uuid.uuid4())
    path = upload_path(job_id)
    with connection() as conn:
        try:
            with conn, conn.cursor() as cur:
                acquire_stock_lock(cur)
                cur.execute(
                    "SELECT id FROM stock_import_jobs WHERE actor_user_id=%s AND submission_key=%s",
                    (user_id, submission_key),
                )
                old = cur.fetchone()
                if old:
                    return str(old[0])
                total = 0
                digest = hashlib.sha256()
                fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
                with os.fdopen(fd, "wb") as dest:
                    while True:
                        block = file.stream.read(1024 * 1024)
                        if not block:
                            break
                        total += len(block)
                        if total > limit("STOCK_MAX_BYTES", 32 * 1024**2):
                            raise ImportProblem("File tồn kho tối đa 32 MB.")
                        digest.update(block)
                        dest.write(block)
                    dest.flush()
                    os.fsync(dest.fileno())
                with path.open("rb") as source:
                    if source.read(4) != b"PK\x03\x04":
                        raise ImportProblem("Nội dung không phải XLSX.")
                cur.execute(
                    """INSERT INTO stock_import_jobs
                       (id,submission_key,actor,actor_user_id,actor_auth_version,filename,file_size,file_sha256)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (job_id, submission_key, actor, user_id, auth_version, filename, total, digest.hexdigest()),
                )
                _event(cur, job_id, actor, "uploaded")
        except BaseException:
            path.unlink(missing_ok=True)
            raise
    return job_id


def control(job_id, action, actor, fingerprint="", confirm_replace=""):
    with connection() as conn, conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute("SELECT * FROM stock_import_jobs WHERE id=%s FOR UPDATE", (str(job_id),))
        job = cur.fetchone()
        if not job:
            raise ImportProblem("Không tìm thấy tác vụ tồn kho.")
        if action == "cancel":
            if job["status"] not in ("queued", "running") or job["cancel_requested"]:
                return
            if job["status"] == "queued":
                cur.execute(
                    "UPDATE stock_import_jobs SET status='cancelled',cancel_requested=true,finished_at=now() WHERE id=%s",
                    (str(job_id),),
                )
            else:
                cur.execute("UPDATE stock_import_jobs SET cancel_requested=true WHERE id=%s", (str(job_id),))
        elif action == "apply":
            preview = job["preview"] or {}
            if job["status"] != "completed" or job["phase"] != "preview" or not job["preview_ready"]:
                raise ImportProblem("Cần xem trước thành công trước khi thay tồn kho.")
            if fingerprint != preview.get("fingerprint"):
                raise ImportProblem("Xem trước đã thay đổi. Hãy tải lại trang.")
            if confirm_replace != str(preview.get("current_rows", 0)):
                raise ImportProblem("Cần nhập chính xác số dòng tồn hiện tại sẽ được thay thế.")
            cur.execute(
                """UPDATE stock_import_jobs SET status='queued',phase='apply',preview_ready=false,
                   finished_at=NULL,processed_count=0 WHERE id=%s""",
                (str(job_id),),
            )
        elif action == "retry":
            if job["status"] not in ("failed", "cancelled"):
                raise ImportProblem("Không thể xem trước lại tác vụ này.")
            if job["attempts"] >= limit("MAX_ATTEMPTS", 6):
                raise ImportProblem("Đã hết số lần thử. Hãy upload lại workbook.")
            cur.execute(
                """UPDATE stock_import_jobs SET status='queued',phase='preview',preview_ready=false,
                   preview=NULL,cancel_requested=false,finished_at=NULL,errors='[]',error_count=0,
                   processed_count=0 WHERE id=%s""",
                (str(job_id),),
            )
        else:
            raise ImportProblem("Thao tác không hợp lệ.")
        _event(cur, job_id, actor, action)


def _clean(value, field, maximum, *, required=True):
    text = unicodedata.normalize("NFC", "" if value is None else str(value)).strip()
    if required and not text:
        raise ImportProblem(f"{field} không được để trống.")
    if len(text) > maximum:
        raise ImportProblem(f"{field} vượt quá {maximum} ký tự.")
    return text


def _quantity(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ImportProblem("Số lượng tồn không được để trống.")
    if isinstance(value, bool):
        raise ImportProblem("Số lượng tồn phải là số nguyên không âm.")
    try:
        number = Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, ValueError):
        raise ImportProblem("Số lượng tồn phải là số nguyên không âm.") from None
    if not number.is_finite() or number < 0 or number != number.to_integral_value():
        raise ImportProblem("Số lượng tồn phải là số nguyên không âm.")
    return int(number)


def _price(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ImportProblem("Giá tồn kho phải là số VND không âm hoặc để trống.")
    try:
        number = Decimal(str(value).strip().replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        raise ImportProblem("Giá tồn kho phải là số VND không âm hoặc để trống.") from None
    if not number.is_finite() or number < 0 or number.as_tuple().exponent < -2:
        raise ImportProblem("Giá tồn kho phải là số VND không âm, tối đa 2 chữ số thập phân.")
    return number


def _expiry(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ImportProblem("Hạn sử dụng phải là ngày Excel, YYYY-MM-DD hoặc DD/MM/YYYY.")


def parse_workbook(path: Path, progress=lambda *_: None):
    inspect_workbook(path)
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if len(wb.worksheets) != 1:
            raise ImportProblem("Workbook tồn kho chỉ được có một trang tính.")
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        iterator = ws.iter_rows()
        try:
            header_cells = next(iterator)
        except StopIteration:
            raise ImportProblem("Workbook đang trống; tồn kho hiện tại không thay đổi.") from None
        if any(cell.data_type in ("f", "e") for cell in header_cells):
            raise ImportProblem("Tiêu đề không được chứa công thức.")
        headers = [
            _clean(cell.value, "Header", 80).casefold() for cell in header_cells
        ]
        if len(headers) != len(HEADERS):
            raise ImportProblem(f"Header phải đúng thứ tự: {_HEADER_REQUIRES}")
        normalized = []
        for header in headers:
            mapped = HEADER_ALIASES.get(header)
            if not mapped:
                raise ImportProblem(f"Header phải đúng thứ tự: {_HEADER_REQUIRES}")
            normalized.append(mapped)
        if len(set(normalized)) != len(normalized):
            raise ImportProblem(f"Header phải đúng thứ tự: {_HEADER_REQUIRES}")
        if tuple(normalized) != tuple(HEADER_ALIASES.get(h, h) for h in HEADERS):
            raise ImportProblem(
                f"Header phải đúng thứ tự: {_HEADER_REQUIRES}"
            )
        rows = []
        errors = []
        for row_number, cells in enumerate(iterator, start=2):
            if row_number > limit("STOCK_MAX_ROWS", 100000) + 1:
                raise ImportProblem("File tồn kho vượt giới hạn số dòng.")
            values = [cell.value for cell in cells[:len(HEADERS)]]
            values.extend([None] * (len(HEADERS) - len(values)))
            if not any(_clean(value, "Giá trị", 4000, required=False) for value in values):
                continue
            if any(cell.data_type in ("f", "e") for cell in cells):
                errors.append(f"Dòng {row_number}: hãy thay công thức bằng giá trị.")
                continue
            try:
                cas_text = _clean(values[2], "CAS", 32, required=False)
                rows.append({
                    "row_number": row_number,
                    "name": _clean(values[0], "Name", 500),
                    "code": _clean(values[1], "Code", 500),
                    "cas": normalize_cas(cas_text) if cas_text else None,
                    "brand": _clean(values[3], "Brand", 180),
                    "size": _clean(values[4], "Size", 500),
                    "stock_price_vnd": _price(values[5]),
                    "quantity": _quantity(values[6]),
                    "expiry_date": _expiry(values[7]),
                })
            except (ImportProblem, ValueError) as exc:
                errors.append(f"Dòng {row_number}: {exc}")
            if len(errors) >= 100:
                break
            if row_number % 100 == 0:
                progress(row_number - 1)
        if errors:
            raise ImportProblem("\n".join(errors))
        if not rows:
            raise ImportProblem("Workbook không có dòng tồn kho; file rỗng không được dùng để xóa toàn bộ tồn.")
        progress(len(rows))
        return rows
    finally:
        wb.close()


def _row_identity(row):
    return (
        normalized_text(row["brand"]), normalized_text(row["code"]),
        normalized_text(row["size"]), row["expiry_date"].isoformat() if row["expiry_date"] else None,
    )


def _row_payload(row):
    return (
        _row_identity(row), row["name"], row["code"], row["brand"], row["size"], row["cas"],
        str(row["stock_price_vnd"]) if row["stock_price_vnd"] is not None else None,
        row["quantity"],
    )


def _prepare_rows(cur, rows, *, register=False):
    if register:
        resolved, created = register_and_resolve_import_rows(cur, rows)
        new_brands = [item["name"] for item in created]
        errors = []
    else:
        resolved, errors, proposed = preview_import_rows_brands(rows, load_brand_gateway(cur))
        new_brands = [item["name"] for item in proposed]
    if errors:
        raise ImportProblem("\n".join(errors[:100]))
    seen = {}
    for row in resolved:
        identity = _row_identity(row)
        if identity in seen:
            expiry_label = row["expiry_date"].isoformat() if row["expiry_date"] else "trống"
            raise ImportProblem(
                f"Dòng {row['row_number']}: trùng Brand + Code + Size + Hạn sử dụng "
                f"với dòng {seen[identity]} (hạn {expiry_label}); không tự cộng số lượng."
            )
        seen[identity] = row["row_number"]
    return resolved, new_brands


def build_plan(cur, rows):
    prepared, new_brands = _prepare_rows(cur, rows, register=False)
    state = active_snapshot(cur)
    cur.execute(
        """SELECT i.name,i.code,i.cas,i.brand,i.size,i.stock_price_vnd,i.quantity,i.expiry_date
           FROM stock_items i WHERE i.snapshot_id=%s""",
        (state["id"],),
    )
    current = []
    for name, code, cas, brand, size, price, quantity, expiry in cur.fetchall():
        current.append({"name": name, "code": code, "cas": cas, "brand": brand, "size": size,
                        "stock_price_vnd": price, "quantity": quantity, "expiry_date": expiry})
    old_by_id = {_row_identity(row): _row_payload(row) for row in current}
    new_by_id = {_row_identity(row): _row_payload(row) for row in prepared}
    added = sum(1 for key in new_by_id if key not in old_by_id)
    changed = sum(1 for key in new_by_id if key in old_by_id and new_by_id[key] != old_by_id[key])
    unchanged = sum(1 for key in new_by_id if key in old_by_id and new_by_id[key] == old_by_id[key])
    removed = sum(1 for key in old_by_id if key not in new_by_id)
    digest_rows = sorted(repr(_row_payload(row)) for row in prepared)
    plan_digest = hashlib.sha256("\n".join(digest_rows).encode("utf-8")).hexdigest()
    sample = [{
        "name": row["name"], "code": row["code"], "cas": row["cas"] or "",
        "brand": row["brand"], "size": row["size"], "quantity": row["quantity"],
        "stock_price": str(row["stock_price_vnd"]) if row["stock_price_vnd"] is not None else "",
        "expiry": row["expiry_date"].isoformat() if row["expiry_date"] else "",
    } for row in prepared[:20]]
    return {
        "total_rows": len(prepared), "brands": sorted({row["brand"] for row in prepared}, key=str.casefold),
        "brand_count": len({normalized_text(row["brand"]) for row in prepared}),
        "new_brands": new_brands, "added": added, "changed": changed,
        "unchanged": unchanged, "removed": removed, "current_rows": len(current),
        "active_snapshot_id": state["id"], "active_revision": state["revision"],
        "fingerprint": snapshot_fingerprint(cur), "plan_digest": plan_digest, "sample": sample,
        "_prepared": prepared,
    }


def _apply(cur, rows, expected, job):
    if snapshot_fingerprint(cur) != expected.get("fingerprint"):
        raise ImportProblem("Snapshot tồn kho đã thay đổi từ lúc xem trước. Hãy xem trước lại.")
    plan = build_plan(cur, rows)
    if plan["plan_digest"] != expected.get("plan_digest"):
        raise ImportProblem("Nội dung hoặc kế hoạch áp dụng không còn giống bản xem trước.")
    cur.execute(
        """SELECT 1 FROM app_users WHERE id=%s AND is_admin AND account_status='ACTIVE'
           AND auth_version=%s FOR UPDATE""",
        (job["actor_user_id"], job["actor_auth_version"]),
    )
    if not cur.fetchone():
        raise ImportProblem("Quyền quản trị hoặc phiên đăng nhập đã thay đổi; không áp dụng tồn kho.")
    prepared, new_brands = _prepare_rows(cur, rows, register=True)
    if sorted(new_brands, key=str.casefold) != sorted(plan["new_brands"], key=str.casefold):
        raise ImportProblem("Danh mục brand đã đổi từ lúc xem trước. Hãy xem trước lại.")
    state = active_snapshot(cur, for_update=True)
    snapshot_id = str(uuid.uuid4())
    cur.execute(
        """INSERT INTO stock_snapshots
           (id,source_kind,source_job_id,replaced_snapshot_id,actor,row_count,content_sha256)
           VALUES (%s,'IMPORT',%s,%s,%s,%s,%s)""",
        (snapshot_id, job["id"], state["id"], job["actor"], len(prepared), plan["plan_digest"]),
    )
    values = [(
        snapshot_id, row["name"], row["code"], row["cas"], row["brand"], row["size"],
        row["stock_price_vnd"], row["quantity"], row["expiry_date"],
        normalized_text(row["brand"]), normalized_text(row["code"]), normalized_text(row["size"]),
        normalized_text(row["cas"]) or None,
    ) for row in prepared]
    execute_values(
        cur,
        """INSERT INTO stock_items
           (snapshot_id,name,code,cas,brand,size,stock_price_vnd,quantity,expiry_date,
            brand_norm,code_norm,size_norm,cas_norm) VALUES %s""",
        values,
        page_size=1000,
    )
    cur.execute(
        """UPDATE stock_state SET active_snapshot_id=%s,revision=revision+1,updated_at=now()
           WHERE singleton=TRUE""",
        (snapshot_id,),
    )
    cur.execute(
        "INSERT INTO stock_snapshot_events(snapshot_id,actor,event,detail) VALUES (%s,%s,'activated',%s)",
        (snapshot_id, job["actor"], Json({"replaced_snapshot_id": state["id"], "row_count": len(prepared)})),
    )
    return plan, snapshot_id


def _snapshot_state_fingerprint(snapshot_id, revision):
    payload = json.dumps(
        {"id": str(snapshot_id) if snapshot_id else None, "revision": int(revision)},
        sort_keys=True,
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def restore_snapshot(snapshot_id, actor, actor_user_id, actor_auth_version, expected_revision, expected_fingerprint):
    source_id = str(uuid.UUID(str(snapshot_id)))
    if expected_revision is None or not str(expected_revision).strip() or not str(expected_fingerprint or "").strip():
        raise ImportProblem("Phiên giao diện đã cũ; hãy làm mới trang rồi thử lại.")
    with connection() as conn, conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
        acquire_stock_lock(cur)
        cur.execute(
            """SELECT 1 FROM app_users WHERE id=%s AND is_admin AND account_status='ACTIVE'
               AND auth_version=%s FOR UPDATE""",
            (actor_user_id, actor_auth_version),
        )
        if not cur.fetchone():
            raise ImportProblem("Quyền quản trị hoặc phiên đăng nhập đã thay đổi; không khôi phục snapshot.")
        state = active_snapshot(cur, for_update=True)
        if expected_revision is not None:
            try:
                if int(expected_revision) != state["revision"]:
                    raise ImportProblem("Hệ thống đã thay đổi. Hãy làm mới lại trước khi khôi phục.")
            except (TypeError, ValueError):
                raise ImportProblem("Phiên giao diện đã cũ; hãy làm mới trang rồi thử lại.") from None
        if expected_fingerprint and expected_fingerprint != _snapshot_state_fingerprint(state["id"], state["revision"]):
            raise ImportProblem("Hệ thống đã thay đổi. Hãy làm mới lại trước khi khôi phục.")
        cur.execute("SELECT * FROM stock_snapshots WHERE id=%s", (source_id,))
        source = cur.fetchone()
        if not source:
            raise ImportProblem("Không tìm thấy snapshot cần khôi phục.")
        if state["id"] == source_id:
            raise ImportProblem("Snapshot này đang được sử dụng.")
        new_id = str(uuid.uuid4())
        cur.execute(
            """INSERT INTO stock_snapshots
               (id,source_kind,restored_from_snapshot_id,replaced_snapshot_id,actor,row_count,content_sha256)
               VALUES (%s,'RESTORE',%s,%s,%s,%s,%s)""",
            (new_id, source_id, state["id"], actor, source["row_count"], source["content_sha256"]),
        )
        cur.execute(
            """INSERT INTO stock_items
               (snapshot_id,name,code,cas,brand,size,stock_price_vnd,quantity,expiry_date,
                brand_norm,code_norm,size_norm,cas_norm)
               SELECT %s,name,code,cas,brand,size,stock_price_vnd,quantity,expiry_date,
                      brand_norm,code_norm,size_norm,cas_norm
               FROM stock_items WHERE snapshot_id=%s""",
            (new_id, source_id),
        )
        cur.execute(
            "UPDATE stock_state SET active_snapshot_id=%s,revision=revision+1,updated_at=now() WHERE singleton=TRUE",
            (new_id,),
        )
        cur.execute(
            "INSERT INTO stock_snapshot_events(snapshot_id,actor,event,detail) VALUES (%s,%s,'restored',%s)",
            (new_id, actor, Json({"restored_from_snapshot_id": source_id, "replaced_snapshot_id": state["id"]})),
        )
        return new_id


def _job_lock(cur, job_id):
    cur.execute(
        "SELECT pg_try_advisory_lock(hashtextextended(%s,%s)) AS acquired",
        (str(job_id), STOCK_JOB_LOCK_NAMESPACE),
    )
    row = cur.fetchone()
    return row["acquired"] if isinstance(row, dict) else row[0]


def _job_unlock(cur, job_id):
    cur.execute("SELECT pg_advisory_unlock(hashtextextended(%s,%s))", (str(job_id), STOCK_JOB_LOCK_NAMESPACE))


def _claim(conn, only_id=None):
    with conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute(
            """SELECT id FROM stock_import_jobs WHERE status IN ('queued','running')
               AND (%s IS NULL OR id=%s::uuid) ORDER BY created_at LIMIT 100""",
            (only_id, only_id),
        )
        ids = [str(row["id"]) for row in cur.fetchall()]
        conn.commit()
        for job_id in ids:
            if not _job_lock(cur, job_id):
                continue
            cur.execute("SELECT * FROM stock_import_jobs WHERE id=%s FOR UPDATE", (job_id,))
            job = cur.fetchone()
            if not job or job["status"] not in ("queued", "running"):
                conn.rollback(); _job_unlock(cur, job_id); continue
            if job["status"] == "running":
                cur.execute(
                    """UPDATE stock_import_jobs SET status='failed',preview_ready=false,finished_at=now(),
                       errors=%s,error_count=1 WHERE id=%s""",
                    (Json(["Worker bị gián đoạn. Hãy xem trước lại trước khi tiếp tục."]), job_id),
                )
                _event(cur, job_id, "worker", "crash_recovered")
                conn.commit(); _job_unlock(cur, job_id); continue
            if job["cancel_requested"] or job["attempts"] >= limit("MAX_ATTEMPTS", 6):
                cur.execute("UPDATE stock_import_jobs SET status='cancelled',finished_at=now() WHERE id=%s", (job_id,))
                conn.commit(); _job_unlock(cur, job_id); continue
            cur.execute(
                """UPDATE stock_import_jobs SET status='running',started_at=now(),heartbeat_at=now(),
                   attempts=attempts+1 WHERE id=%s RETURNING *""",
                (job_id,),
            )
            job = cur.fetchone()
            _event(cur, job_id, "worker", job["phase"] + "_started")
            conn.commit()
            return job
    return None


def run_once(only_id=None):
    with connection() as conn:
        job = _claim(conn, only_id)
        if not job:
            return False
        job_id = str(job["id"])
        stop = threading.Event()
        cancelled = threading.Event()
        deadline = time.monotonic() + limit("JOB_SECONDS", 7200)

        def watch():
            try:
                with connection() as monitor:
                    monitor.autocommit = True
                    with monitor.cursor() as cur:
                        while not stop.wait(1):
                            cur.execute(
                                """UPDATE stock_import_jobs SET heartbeat_at=now() WHERE id=%s
                                   RETURNING cancel_requested""",
                                (job_id,),
                            )
                            row = cur.fetchone()
                            if not row or row[0] or time.monotonic() > deadline:
                                cancelled.set(); conn.cancel(); break
            except Exception:
                cancelled.set(); conn.cancel()

        thread = threading.Thread(target=watch, daemon=True)
        thread.start()

        def progress(count):
            if cancelled.is_set():
                raise ImportProblem("Tác vụ đã bị hủy; snapshot hiện tại được giữ nguyên.")
            with connection() as progress_conn, progress_conn, progress_conn.cursor() as cur:
                cur.execute(
                    "UPDATE stock_import_jobs SET processed_count=%s,heartbeat_at=now() "
                    "WHERE id=%s RETURNING cancel_requested",
                    (count, job_id),
                )
                row = cur.fetchone()
                if row and row[0]:
                    cancelled.set()
                    raise ImportProblem("Tác vụ đã bị hủy; snapshot hiện tại được giữ nguyên.")

        try:
            rows = parse_workbook(upload_path(job_id), progress)
            if cancelled.is_set():
                raise ImportProblem("Tác vụ đã bị hủy; snapshot hiện tại được giữ nguyên.")
            with conn:
                with conn.cursor() as lock_cur:
                    lock_cur.execute(
                        "SELECT set_config('statement_timeout', %s, true)",
                        (str(limit("SQL_SECONDS", 600) * 1000),),
                    )
                    acquire_products_import_lock(lock_cur)
                    acquire_stock_lock(lock_cur)
                with conn.cursor(cursor_factory=RealDictCursor) as job_cur:
                    job_cur.execute("SELECT * FROM stock_import_jobs WHERE id=%s", (job_id,))
                    live_job = job_cur.fetchone()
                if not live_job:
                    raise ImportProblem("Không tìm thấy tác vụ tồn kho.")
                if live_job["status"] != "running" or live_job["cancel_requested"]:
                    raise ImportProblem("Tác vụ đã bị hủy; snapshot hiện tại được giữ nguyên.")
                if live_job["phase"] == "preview":
                    with conn.cursor() as cur:
                        plan = build_plan(cur, rows)
                        public_plan = {key: value for key, value in plan.items() if not key.startswith("_")}
                        # Stop the monitor before taking the final job-row lock;
                        # this avoids a monitor/job-lock cycle while preserving
                        # the final cancellation check as the commit gate.
                        stop.set()
                        thread.join(timeout=5)
                        if cancelled.is_set() or time.monotonic() > deadline:
                            raise ImportProblem("Tác vụ đã bị hủy; snapshot hiện tại được giữ nguyên.")
                        cur.execute("SELECT cancel_requested FROM stock_import_jobs WHERE id=%s FOR UPDATE", (job_id,))
                        if cur.fetchone()[0]:
                            raise ImportProblem("Tác vụ đã bị hủy; snapshot hiện tại được giữ nguyên.")
                        cur.execute(
                            """UPDATE stock_import_jobs SET status='completed',preview_ready=true,preview=%s,
                               row_count=%s,processed_count=%s,error_count=0,errors='[]',finished_at=now()
                               WHERE id=%s""",
                            (Json(public_plan), len(rows), len(rows), job_id),
                        )
                        _event(cur, job_id, "worker", "preview_completed")
                else:
                    with conn.cursor() as cur:
                        plan, snapshot_id = _apply(cur, rows, live_job["preview"] or {}, live_job)
                        stop.set()
                        thread.join(timeout=5)
                        if cancelled.is_set() or time.monotonic() > deadline:
                            raise ImportProblem("Tác vụ đã bị hủy; snapshot hiện tại được giữ nguyên.")
                        cur.execute("SELECT cancel_requested FROM stock_import_jobs WHERE id=%s FOR UPDATE", (job_id,))
                        if cur.fetchone()[0]:
                            raise ImportProblem("Tác vụ đã bị hủy; snapshot hiện tại được giữ nguyên.")
                        cur.execute(
                            """UPDATE stock_import_jobs SET status='completed',preview_ready=false,
                               inserted_count=%s,deleted_count=%s,processed_count=%s,error_count=0,
                               errors='[]',finished_at=now() WHERE id=%s""",
                            (plan["total_rows"], plan["current_rows"], len(rows), job_id),
                        )
                        _event(cur, job_id, "worker", "apply_completed", {"snapshot_id": snapshot_id})
        except BaseException as exc:
            conn.rollback()
            message = str(exc).strip() or "Không xử lý được workbook tồn kho."
            with conn, conn.cursor() as cur:
                cur.execute(
                    """UPDATE stock_import_jobs
                       SET status=CASE WHEN cancel_requested THEN 'cancelled' ELSE 'failed' END,
                           preview_ready=false,errors=%s,error_count=1,finished_at=now()
                       WHERE id=%s""",
                    (Json(message.splitlines()[:100]), job_id),
                )
                _event(cur, job_id, "worker", "failed")
        finally:
            stop.set(); thread.join(timeout=2)
            try:
                with conn.cursor() as cur:
                    _job_unlock(cur, job_id)
                conn.commit()
            except Exception:
                conn.rollback()
        return True
