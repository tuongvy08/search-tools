import hashlib
import io
import json
import os
import unittest
from pathlib import Path
from unittest import mock

import psycopg2
from openpyxl import load_workbook

import search
import quote_workbook_export as qwe
from auth_test_helpers import start_auth_db_patch
from pg_temp_db import create_full_schema_temp_db, drop_temp_db, probe_postgres_reachable
from test_quote_workbook_export import make_workbook, product


ROOT = Path(__file__).resolve().parents[1]


def changed_mapping_workbook():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "BG"
    ws["Q16"] = "Tên hàng"; ws["R16"] = "Code"; ws["P16"] = "Đơn giá"
    for row in range(17, 26): ws[f"J{row}"] = f"=P{row}*1"
    ws["B26"] = "Tổng giá"; ws["J26"] = "=SUM(J17:J25)"
    ws["J27"] = "=J26*10%"; ws["J28"] = "=J26+J27"
    output = io.BytesIO(); wb.save(output); wb.close()
    mapping = search._quote_template_mapping_snapshot()
    mapping["mapping"] = {"sequence": "A", "Name": "Q", "Code": "R", "Unit_Price_Value": "P"}
    mapping["total_formula_column"] = "J"
    return output.getvalue(), mapping


class Phase6C1MappingTests(unittest.TestCase):
    def test_migration_is_additive_and_assignment_is_one_per_team(self):
        sql = (ROOT / "sql" / "migration_022_team_quote_templates.sql").read_text()
        self.assertIn("CREATE TABLE IF NOT EXISTS team_quote_templates", sql)
        self.assertIn("team_id INTEGER PRIMARY KEY", sql)
        self.assertIn("template_id BIGINT NOT NULL", sql)
        self.assertIn("BEGIN;", sql)
        self.assertIn("mapping_v2_json", sql)
        self.assertNotIn("DROP CONSTRAINT", sql)
        self.assertNotIn("products ", sql.lower())

    def test_changed_columns_are_validated_and_used_by_same_engine(self):
        raw, mapping = changed_mapping_workbook()
        normalized = search._validate_quote_template_mapping(raw, mapping)
        exported = search.export_quick_quote_workbook(raw, [product(1)], normalized)
        check = load_workbook(io.BytesIO(exported), data_only=False)
        self.assertEqual(check["BG"]["Q17"].value, product(1)["Name"])
        self.assertEqual(check["BG"]["R17"].value, product(1)["Code"])
        check.close()

    def test_duplicate_missing_and_malformed_mapping_fail_closed(self):
        raw = make_workbook()
        base = search._quote_template_mapping_snapshot()
        for mapping in (
            {**base, "mapping": {"sequence": "A", "Name": "B"}},
            {**base, "mapping": {"sequence": "A", "Name": "B", "Code": "B"}},
            {**base, "mapping": {"sequence": "A", "Name": "B", "Code": "XFE"}},
        ):
            with self.assertRaises(search.QuoteTemplateError):
                search._validate_quote_template_mapping(raw, mapping)


@unittest.skipUnless(probe_postgres_reachable(), "local PostgreSQL required")
class Phase6C1ResolverPgTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dbname, cls.dsn = create_full_schema_temp_db()
        cls.env = mock.patch.dict(os.environ, DATABASE_URL=cls.dsn, DISABLE_IP_ALLOWLIST="1")
        cls.env.start()
        conn = psycopg2.connect(cls.dsn)
        try:
            with conn, conn.cursor() as cur:
                cur.execute("INSERT INTO teams(name) VALUES ('A'),('B') RETURNING id")
                cls.team_a, cls.team_b = [row[0] for row in cur.fetchall()]
                cls.templates = []
                for filename, active in (("global.xlsx", True), ("team.xlsx", False)):
                    raw, mapping = (changed_mapping_workbook() if filename == "team.xlsx"
                                    else (make_workbook(), search._quote_template_mapping_snapshot()))
                    cur.execute("""INSERT INTO quote_templates(filename,content,content_sha256,content_size,profile_version,mapping_json,mapping_v2_json,is_active)
                        VALUES (%s,%s,%s,%s,'BG_V1',%s::jsonb,%s::jsonb,%s) RETURNING id""",
                        (filename, psycopg2.Binary(raw), hashlib.sha256(raw).hexdigest(), len(raw),
                         json.dumps(search._quote_template_mapping_snapshot()), json.dumps(mapping), active))
                    cls.templates.append(cur.fetchone()[0])
                cur.execute("INSERT INTO team_quote_templates(team_id,template_id,assigned_by) VALUES (%s,%s,'test')",
                            (cls.team_a, cls.templates[1]))
        finally:
            conn.close()

    @classmethod
    def tearDownClass(cls):
        cls.env.stop(); drop_temp_db(cls.dbname)

    def test_team_global_and_validated_admin_override(self):
        conn = psycopg2.connect(self.dsn)
        try:
            self.assertEqual(search._get_active_quote_template(conn, team_id=self.team_a)["source"], "team")
            self.assertEqual(search._get_active_quote_template(conn, team_id=self.team_b)["source"], "global")
            chosen = search._get_active_quote_template(conn, template_id=self.templates[1], admin_override=True)
            self.assertEqual(chosen["source"], "admin")
            with self.assertRaises(search.QuoteTemplateError):
                search._get_active_quote_template(conn, template_id=999999, admin_override=True)

            team_template = search._get_active_quote_template(conn, team_id=self.team_a, include_content=True)
            exported = search.export_quick_quote_workbook(
                team_template["content"], [product(1)], team_template["mapping"]
            )
            check = load_workbook(io.BytesIO(exported), data_only=False)
            self.assertEqual(check["BG"]["Q17"].value, product(1)["Name"])
            self.assertEqual(check["BG"]["R17"].value, product(1)["Code"])
            self.assertIsNone(check["BG"]["B17"].value)
            check.close()
        finally:
            conn.close()

    def test_migration_rehearses_twice(self):
        sql = (ROOT / "sql" / "migration_022_team_quote_templates.sql").read_text()
        conn = psycopg2.connect(self.dsn)
        try:
            with conn, conn.cursor() as cur:
                cur.execute(sql); cur.execute(sql)
        finally:
            conn.close()


class Phase6C1SharedBoundaryTests(unittest.TestCase):
    def setUp(self):
        self._ip_patch = mock.patch.dict(os.environ, {"DISABLE_IP_ALLOWLIST": "1"})
        self._ip_patch.start()
        self.addCleanup(self._ip_patch.stop)
        start_auth_db_patch(self)
        search.app.testing = True

    def _client(self):
        client = search.app.test_client()
        with client.session_transaction() as sess:
            sess.update(authenticated=True, user_id=1, auth_version=1, is_admin=True,
                        team_id=None, csrf_token="csrf-test")
        return client

    def test_search_and_quick_quote_call_one_builder(self):
        client = self._client()
        result = (b"PK-export", "quote.xlsx", {"id": 1})
        with mock.patch.object(search, "get_connection", return_value=mock.MagicMock()), \
             mock.patch.object(search, "_quote_context", return_value={"is_admin": True, "team_id": None, "template_id": None, "grants": frozenset(search.team_permissions.REGISTRY)}), \
             mock.patch.object(search, "_create_quote_workbook", return_value=result) as builder:
            a = client.post("/api/results/quote-export", data={"source": "SEARCH", "selections": json.dumps([
                {"product_id": 9}, {"product_id": 3}
            ])},
                            headers={"X-CSRF-Token": "csrf-test"})
            b = client.post("/api/quote-assistant/workbook/export",
                            data={"selections": json.dumps([{"product_id": 1}])},
                            headers={"X-CSRF-Token": "csrf-test"})
        self.assertEqual((a.status_code, b.status_code), (200, 200))
        self.assertEqual(builder.call_count, 2)
        self.assertEqual(builder.call_args_list[0].kwargs["selections"], [
            {"ord": 1, "product_id": 9}, {"ord": 2, "product_id": 3}
        ])

    def test_search_export_requires_the_claimed_source_capability(self):
        client = self._client()
        context = {"is_admin": False, "team_id": 7, "template_id": None,
                   "grants": frozenset({"EXPORT", "VIEW_COMPLIANCE"})}
        with mock.patch.object(search, "get_connection", return_value=mock.MagicMock()), \
             mock.patch.object(search, "_quote_context", return_value=context), \
             mock.patch.object(search, "_create_quote_workbook") as builder:
            response = client.post("/api/results/quote-export", data={
                "source": "SEARCH", "selections": json.dumps([{"product_id": 9}])
            }, headers={"X-CSRF-Token": "csrf-test"})
        self.assertEqual(response.status_code, 403)
        builder.assert_not_called()

    def test_search_export_rejects_missing_or_unknown_source_before_database_access(self):
        client = self._client()
        with mock.patch.object(search, "get_connection") as get_connection:
            for source in (None, "", "QUICK_QUOTE", "NOT_A_SOURCE"):
                data = {"selections": json.dumps([{"product_id": 9}])}
                if source is not None:
                    data["source"] = source
                with self.subTest(source=source):
                    response = client.post(
                        "/api/results/quote-export", data=data,
                        headers={"X-CSRF-Token": "csrf-test"},
                    )
                    self.assertEqual(response.status_code, 400)
                    self.assertIn("Nguồn kết quả", response.get_json()["error"])
        get_connection.assert_not_called()

    def test_export_requires_visible_compliance_to_avoid_status_oracle(self):
        self.assertFalse(search.team_permissions.allows("EXPORT", {"EXPORT"}))
        self.assertTrue(search.team_permissions.allows(
            "EXPORT", {"EXPORT", "VIEW_COMPLIANCE"}
        ))

    def test_restricted_export_uses_clean_workbook_without_template_stale_data(self):
        template = {
            "id": 1, "filename": "quote.xlsx", "content": b"PK-STALE-PRICE-SECRET",
            "mapping": search._quote_template_mapping_snapshot(),
        }
        products = [{
            "Name": "HIDDEN-NAME-SECRET", "Code": "ITEM-X", "Unit_Price": "999,000",
            "Unit_Price_Value": 999000, "Compliance": "Được bán",
        }]
        context = {"is_admin": False, "team_id": 7, "template_id": None,
                   "grants": frozenset({"EXPORT", "VIEW_CODE", "VIEW_COMPLIANCE"})}
        with mock.patch.object(search, "_get_active_quote_template", return_value=template), \
             mock.patch.object(search, "_quote_export_products", return_value=products):
            exported, _name, _template = search._create_quote_workbook(
                mock.MagicMock(), selections=[{"ord": 1, "product_id": 9}], context=context
            )
        check = load_workbook(io.BytesIO(exported), data_only=False)
        self.assertEqual(check.sheetnames, ["Báo giá"])
        self.assertEqual(
            [cell.value for cell in check["Báo giá"][1]],
            ["Code", "Tình trạng quản lý"],
        )
        self.assertEqual(
            [cell.value for cell in check["Báo giá"][2]], ["ITEM-X", "Được bán"]
        )
        check.close()
        package_text = b"\n".join(qwe._read_valid_xlsx_entries(exported).values())
        self.assertNotIn(b"STALE-PRICE-SECRET", package_text)
        self.assertNotIn(b"HIDDEN-NAME-SECRET", package_text)

    def test_fully_authorized_export_keeps_the_approved_template(self):
        template = {
            "id": 1, "filename": "approved.xlsx", "content": b"PK-APPROVED-TEMPLATE",
            "mapping": search._quote_template_mapping_snapshot(),
        }
        context = {"is_admin": True, "team_id": None, "template_id": None,
                   "grants": frozenset(search.team_permissions.REGISTRY)}
        products = [{"Code": "ITEM-X", "Compliance": "Được bán"}]
        with mock.patch.object(search, "_get_active_quote_template", return_value=template), \
             mock.patch.object(search, "_quote_export_products", return_value=products), \
             mock.patch.object(search, "export_quick_quote_workbook", return_value=b"PK-OUT") as engine:
            exported, _name, _template = search._create_quote_workbook(
                mock.MagicMock(), selections=[{"ord": 1, "product_id": 9}], context=context
            )
        self.assertEqual(exported, b"PK-OUT")
        engine.assert_called_once_with(template["content"], products, template["mapping"])

    def test_staff_context_spoofing_is_rejected(self):
        with search.app.test_request_context("/?team_id=999&template_id=999"):
            search.session.update(authenticated=True, is_admin=False, team_id=1)
            with self.assertRaises(search.QuoteTemplateError):
                search._quote_context(mock.MagicMock())

    def test_admin_team_context_requires_team_export_capability(self):
        conn = mock.MagicMock()
        conn.cursor.return_value.__enter__.return_value.fetchone.return_value = (["SEARCH"],)
        with search.app.test_request_context("/?team_id=7"):
            search.session.update(authenticated=True, is_admin=True)
            with self.assertRaisesRegex(search.QuoteTemplateError, "quyền xuất báo giá"):
                search._quote_context(conn)

    def test_quick_quote_export_requires_valid_csrf(self):
        client = self._client()
        payload = {"selections": json.dumps([{"product_id": 1}])}
        missing = client.post("/api/quote-assistant/workbook/export", data=payload)
        wrong = client.post("/api/quote-assistant/workbook/export", data=payload,
                            headers={"X-CSRF-Token": "wrong"})
        self.assertEqual((missing.status_code, wrong.status_code), (400, 400))
        self.assertIn("CSRF", missing.get_json()["error"])


if __name__ == "__main__":
    unittest.main()
