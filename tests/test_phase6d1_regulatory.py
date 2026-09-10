"""Phase 6D1 regulatory contract on isolated real PostgreSQL only."""

from io import BytesIO
import os
from pathlib import Path
import secrets
import shutil
import tempfile
import threading
import time
import unittest
from unittest import mock

import psycopg2
from psycopg2.extras import RealDictCursor
from openpyxl import Workbook, load_workbook
from werkzeug.datastructures import FileStorage

os.environ.setdefault("FLASK_SECRET_KEY", "phase6d1-test-only")

import regulatory_import_jobs
import admin_regulatory
from import_engine import ImportProblem
from regulatory import normalize_cas, normalized_identity, product_resolver_lateral
import search
from auth_test_helpers import start_auth_db_patch
import pg_temp_db


ROOT = Path(__file__).resolve().parents[1]
MIGRATION_026 = (ROOT / "sql" / "migration_026_regulatory_management.sql").read_text(encoding="utf-8")
MIGRATION_027 = (ROOT / "sql" / "migration_027_regulatory_status_colors.sql").read_text(encoding="utf-8")


def _workbook_bytes(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


@unittest.skipUnless(pg_temp_db.probe_postgres_reachable(), "isolated local PostgreSQL required")
class Phase6D1RegulatoryPgTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.db_name, cls.dsn = pg_temp_db.create_full_schema_temp_db()
        cls.upload_dir = tempfile.mkdtemp(prefix="phase6d1-uploads-")
        cls.env = mock.patch.dict(os.environ, {
            "DATABASE_URL": cls.dsn,
            "DISABLE_IP_ALLOWLIST": "1",
            "IMPORT_UPLOAD_DIR": cls.upload_dir,
        })
        cls.env.start()
        cls.conn = psycopg2.connect(cls.dsn)
        cls.conn.autocommit = True
        with cls.conn.cursor() as cur:
            cur.execute("INSERT INTO app_users(username,password_hash,is_admin,account_status,auth_version) VALUES ('phase6d1_admin','x',true,'ACTIVE',1) RETURNING id")
            cls.admin_id = cur.fetchone()[0]

    @classmethod
    def tearDownClass(cls):
        cls.conn.close()
        cls.env.stop()
        try:
            pg_temp_db.drop_temp_db(cls.db_name)
        finally:
            shutil.rmtree(cls.upload_dir, ignore_errors=True)

    def setUp(self):
        with self.conn.cursor() as cur:
            cur.execute("""UPDATE app_users SET is_admin=true,account_status='ACTIVE',auth_version=1
                           WHERE id=%s""", (self.admin_id,))
            cur.execute("DELETE FROM regulatory_import_jobs")
            cur.execute("DELETE FROM regulatory_status_events")
            cur.execute("DELETE FROM products")
            cur.execute("DELETE FROM brand_compliance_settings")
            cur.execute("DELETE FROM regulatory_rules")
            cur.execute("DELETE FROM regulatory_statuses WHERE stable_key LIKE 'CUSTOM_%'")
            cur.execute("""UPDATE regulatory_statuses SET
                           label=CASE stable_key WHEN 'CAM_NHAP' THEN 'CẤM NHẬP' WHEN 'PHU_LUC_II' THEN 'Phụ lục II' WHEN 'PHU_LUC_III' THEN 'Phụ lục III' WHEN 'DUOC_BAN' THEN 'Được bán' WHEN 'CAN_GIAY_PHEP' THEN 'Cần giấy phép' ELSE 'Chưa xác định' END,
                           priority=CASE stable_key WHEN 'CAM_NHAP' THEN 10 WHEN 'PHU_LUC_II' THEN 20 WHEN 'PHU_LUC_III' THEN 30 WHEN 'DUOC_BAN' THEN 40 WHEN 'CAN_GIAY_PHEP' THEN 50 ELSE 60 END,
                           color_key=CASE stable_key WHEN 'CAM_NHAP' THEN 'red' WHEN 'PHU_LUC_II' THEN 'amber' WHEN 'PHU_LUC_III' THEN 'teal' WHEN 'DUOC_BAN' THEN 'green' WHEN 'CAN_GIAY_PHEP' THEN 'blue' ELSE 'gray' END""")

    def _status(self, key):
        with self.conn.cursor() as cur:
            cur.execute("SELECT id,label,priority,export_policy FROM regulatory_statuses WHERE stable_key=%s", (key,))
            return cur.fetchone()

    def _rule(self, key, field, value, note=""):
        status_id, label, priority, _ = self._status(key)
        with self.conn.cursor() as cur:
            cur.execute("""INSERT INTO regulatory_rules(rule_type,rule_label,match_field,match_value,priority,is_active,note,status_id)
                           VALUES (%s,%s,%s,%s,%s,true,%s,%s)""",
                        (key,label,field,value,priority,note,status_id))

    def _admin_client(self):
        start_auth_db_patch(self)
        search.app.testing = True
        client = search.app.test_client()
        with client.session_transaction() as sess:
            sess.update(authenticated=True,user_id=self.admin_id,auth_version=1,is_admin=True,
                        username="phase6d1_admin",csrf_token="phase6d1-csrf")
        return client

    def _real_admin_client(self):
        search.app.testing = True
        client = search.app.test_client()
        with client.session_transaction() as sess:
            sess.update(authenticated=True,user_id=self.admin_id,auth_version=1,is_admin=True,
                        username="phase6d1_admin",csrf_token="phase6d1-csrf")
        return client

    def _team_client(self, grants):
        marker=secrets.token_hex(4)
        with self.conn.cursor() as cur:
            cur.execute("""INSERT INTO teams(name,lifecycle_status,permission_keys)
                           VALUES (%s,'ACTIVE',%s) RETURNING id""",(f"Regulatory {marker}",list(grants)))
            team_id=cur.fetchone()[0]
            cur.execute("""INSERT INTO app_users(username,password_hash,team_id,is_admin,account_status,auth_version)
                           VALUES (%s,'x',%s,false,'ACTIVE',1) RETURNING id""",(f"reg_{marker}",team_id))
            user_id=cur.fetchone()[0]
        client=search.app.test_client()
        with client.session_transaction() as sess:
            sess.update(authenticated=True,user_id=user_id,auth_version=1,is_admin=False,
                        team_id=team_id,username=f"reg_{marker}")
        return client

    def test_cross_field_priority_substring_notes_and_blank_no_match(self):
        self._rule("PHU_LUC_II", "cas", "50-00-0", "ghi chú CAS")
        self._rule("CAM_NHAP", "code", "CODE-X", "ghi chú Z")
        self._rule("CAM_NHAP", "name", "Formaldehyde", "ghi chú A")
        self._rule("CAM_NHAP", "cas", "50-00-0", "ghi chú A")
        with self.conn.cursor() as cur:
            cur.execute("""INSERT INTO products(name,code,cas,brand,size,ship,price)
                           VALUES ('Dung dịch Formaldehyde chuẩn','CODE-X','50-00-0','Brand A','1g','1','100'),
                                  ('Không CAS','CODE-X',NULL,'Brand A','1g','1','100'),
                                  ('Không khớp','FREE',NULL,'Brand A','1g','1','100')""")
        data = self._admin_client().get("/search", query_string={"query": ""}).get_json()["results"]
        by_code = {row["Name"]: row for row in data}
        self.assertEqual(by_code["Dung dịch Formaldehyde chuẩn"]["Compliance_Status"], "CẤM NHẬP")
        self.assertEqual(by_code["Dung dịch Formaldehyde chuẩn"]["Compliance_Note"], "ghi chú A\nghi chú Z")
        self.assertEqual(by_code["Dung dịch Formaldehyde chuẩn"]["Compliance_Export_Policy"], "BLOCK")
        self.assertEqual(by_code["Không CAS"]["Compliance_Status"], "CẤM NHẬP")
        self.assertEqual(by_code["Không khớp"]["Compliance_Status"], "")

    def test_manual_allow_and_rename_keep_stable_policy(self):
        self._rule("CAM_NHAP", "cas", "50-00-0", "blocked")
        with self.conn.cursor() as cur:
            cur.execute("INSERT INTO brand_compliance_settings VALUES ('BRAND A',true,now())")
            cur.execute("""INSERT INTO products(name,code,cas,brand,size,ship,price,manual_compliance)
                           VALUES ('Manual','M-1','50-00-0','Brand A','1g','1','100','Được bán') RETURNING manual_compliance_status_id""")
            manual_status_id = cur.fetchone()[0]
            cur.execute("SELECT id FROM regulatory_statuses WHERE stable_key='CAM_NHAP'")
            blocked_status_id = cur.fetchone()[0]
            cur.execute("UPDATE regulatory_statuses SET label='Không được nhập khẩu',updated_at=now() WHERE id=%s", (blocked_status_id,))
            cur.execute("UPDATE regulatory_rules SET rule_label='Không được nhập khẩu' WHERE status_id=%s", (blocked_status_id,))
        row = self._admin_client().get("/search", query_string={"query": "M-1"}).get_json()["results"][0]
        self.assertEqual(row["Compliance_Status"], "Được bán")
        self.assertEqual(row["Compliance_Export_Policy"], "ALLOW")
        self.assertEqual(row["compliance_source"], "manual")
        self.assertEqual(row["Compliance_Status"], self._status("DUOC_BAN")[1])
        self.assertEqual(manual_status_id, self._status("DUOC_BAN")[0])
        license_result = self._admin_client().get(
            "/check_cas", query_string={"cas": "50-00-0"}
        ).get_json()
        self.assertEqual(license_result["warning_type"], "Không được nhập khẩu")
        self.assertEqual(license_result["export_policy"], "BLOCK")

    def test_check_license_is_independent_from_catalog_and_team_scope(self):
        self._rule("CAM_NHAP", "cas", "50-00-0", "độc lập catalog")
        response = self._admin_client().get("/check_cas_batch", query_string={"cas": "50-00-0\n64-17-5"})
        self.assertEqual(response.status_code, 200)
        rows = response.get_json()["results"]
        self.assertEqual(rows[0]["Compliance_Status"], "CẤM NHẬP")
        self.assertEqual(rows[0]["Compliance_Note"], "độc lập catalog")
        self.assertEqual(rows[1]["Compliance_Status"], "")
        self.assertNotIn("Brand", rows[0])

    def test_check_license_note_permission_uses_only_regulatory_note_alias(self):
        self._rule("CAM_NHAP", "cas", "50-00-0", "ghi chú quản lý riêng")
        base={"CHECK_LICENSE","SEARCH_BY_CAS","VIEW_COMPLIANCE"}
        product_note_only=self._team_client(base|{"VIEW_NOTE"})
        for path in ("/check_cas", "/check_cas_batch"):
            response=product_note_only.get(path,query_string={"cas":"50-00-0"})
            self.assertEqual(response.status_code,200,response.get_data(as_text=True))
            payload=response.get_json()
            self.assertNotIn("Compliance_Note",payload if path=="/check_cas" else payload["results"][0])
            self.assertNotIn("compliance_note",payload if path=="/check_cas" else payload["results"][0])
        single=product_note_only.get('/check_cas',query_string={'cas':'50-00-0'}).get_json()
        self.assertEqual(single['warning_type'],'CẤM NHẬP')
        self.assertEqual(single['export_policy'],'BLOCK')
        self.assertIn('message',single)

        regulatory_note_only=self._team_client(base|{"VIEW_COMPLIANCE_NOTE"})
        single=regulatory_note_only.get('/check_cas',query_string={'cas':'50-00-0'}).get_json()
        self.assertEqual(single['Compliance_Note'],'ghi chú quản lý riêng')
        self.assertNotIn('note',single)
        batch=regulatory_note_only.get('/check_cas_batch',query_string={'cas':'50-00-0'}).get_json()['results'][0]
        self.assertEqual(batch['Compliance_Note'],'ghi chú quản lý riêng')

    def test_long_aggregated_notes_are_preserved_search_quick_quote_and_export(self):
        note_a='A'*2500
        note_b='B'*2500
        expected=note_a+'\n'+note_b
        self._rule('PHU_LUC_II','cas','50-00-0',note_a)
        self._rule('PHU_LUC_II','code','LONG-NOTE',note_b)
        with self.conn.cursor() as cur:
            cur.execute("""INSERT INTO products(name,code,cas,brand,size,ship,price)
                           VALUES ('Long notes','LONG-NOTE','50-00-0','Brand A','1g','1','100') RETURNING id""")
            product_id=cur.fetchone()[0]
        client=self._admin_client()
        search_row=client.get('/search',query_string={'query':'LONG-NOTE'}).get_json()['results'][0]
        self.assertEqual(search_row['Compliance_Note'],expected)
        quote=client.post('/api/quote-assistant/match',json={
            'rows':[{'code':'LONG-NOTE'}],'selection_strategy':'LOWEST_UNIT_PRICE',
        }).get_json()['results'][0]
        self.assertEqual(quote['candidates'][0]['Compliance_Note'],expected)
        with search.app.test_request_context('/'):
            with self.conn.cursor() as _cur:
                export_rows=search._quote_export_products(
                    self.conn,[{'ord':1,'product_id':product_id}],
                    {'is_admin':True,'team_id':None,
                     'grants':frozenset({'VIEW_COMPLIANCE','VIEW_COMPLIANCE_NOTE'})},
                )
        self.assertEqual(export_rows[0]['Compliance_Note'],expected)

    def test_import_validation_duplicate_scoped_replace_and_stale_preview(self):
        self._rule("CAM_NHAP", "code", "OLD-CODE", "old")
        self._rule("CAM_NHAP", "cas", "64-17-5", "keep cas")
        bad_path = Path(os.environ["IMPORT_UPLOAD_DIR"]) / "bad.xlsx"
        bad_path.write_bytes(_workbook_bytes(
            ["CAS", "Tình trạng quản lý", "Ghi chú quản lý"],
            [["123-45-6", "CẤM NHẬP", "bad"]],
        ).read())
        with self.assertRaisesRegex(ImportProblem, "checksum"):
            regulatory_import_jobs._parse_workbook(bad_path)

        duplicate_path = Path(os.environ["IMPORT_UPLOAD_DIR"]) / "duplicate.xlsx"
        duplicate_path.write_bytes(_workbook_bytes(
            ["Code", "Tình trạng quản lý", "Ghi chú quản lý"],
            [["X", "CẤM NHẬP", "a"], ["x", "CẤM NHẬP", "b"]],
        ).read())
        rows = regulatory_import_jobs._parse_workbook(duplicate_path)
        with self.conn.cursor() as cur:
            with self.assertRaisesRegex(ImportProblem, "trùng Code"):
                regulatory_import_jobs.build_plan(cur, rows, "upsert")

        replace_path = Path(os.environ["IMPORT_UPLOAD_DIR"]) / "replace.xlsx"
        replace_path.write_bytes(_workbook_bytes(
            ["Code", "Tình trạng quản lý", "Ghi chú quản lý"],
            [["NEW-CODE", "CẤM NHẬP", "new"]],
        ).read())
        rows = regulatory_import_jobs._parse_workbook(replace_path)
        with self.conn:
            with self.conn.cursor() as cur:
                cur.execute("SELECT pg_advisory_xact_lock(%s)", (regulatory_import_jobs.REGULATORY_LOCK_KEY,))
                plan = regulatory_import_jobs.build_plan(cur, rows, "replace_scoped")
                self.assertEqual(plan["deleted"], 1)
                regulatory_import_jobs.apply_plan(cur, rows, "replace_scoped", plan)
        with self.conn.cursor() as cur:
            cur.execute("SELECT match_field,match_value FROM regulatory_rules ORDER BY match_field,match_value")
            self.assertEqual(cur.fetchall(), [("cas", "64-17-5"), ("code", "NEW-CODE")])
            stale = regulatory_import_jobs.build_plan(cur, rows, "upsert")
            cur.execute("UPDATE regulatory_statuses SET updated_at=now(),label='CẤM NHẬP mới' WHERE stable_key='CAM_NHAP'")
            with self.assertRaisesRegex(ImportProblem, "đổi từ lúc xem trước"):
                regulatory_import_jobs.apply_plan(cur, rows, "upsert", stale)

    def test_new_status_is_allow_and_last(self):
        rows = [{"row_number": 2, "match_field": "code", "match_value": "NEW",
                 "status_label": "Theo dõi nội bộ", "note": ""}]
        with self.conn:
            with self.conn.cursor() as cur:
                plan = regulatory_import_jobs.build_plan(cur, rows, "upsert")
                regulatory_import_jobs.apply_plan(cur, rows, "upsert", plan)
        with self.conn.cursor() as cur:
            cur.execute("SELECT export_policy,priority=(SELECT max(priority) FROM regulatory_statuses),color_key FROM regulatory_statuses WHERE label='Theo dõi nội bộ'")
            self.assertEqual(cur.fetchone(), ("ALLOW", True, "gray"))

    def test_admin_ui_status_create_rename_and_templates(self):
        client = self._admin_client()
        page = client.get("/admin/regulatory")
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn("Quy tắc quản lý", html)
        self.assertIn("Thay đúng nhóm trường + tình trạng", html)
        self.assertIn("Nhập dữ liệu an toàn", html)
        self.assertIn("Đổi màu", html)
        self.assertIn("Bảng màu cho", html)
        for raw in ("Phase 6D1", "Policy cho/chặn", "Background import", "format/checksum", "giữ rule"):
            self.assertNotIn(raw, html)
        self.assertIn('meta name="viewport"', html)
        response = client.post("/admin/regulatory/statuses", data={
            "csrf_token": "phase6d1-csrf", "action": "add", "label": "Theo dõi UI",
        })
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("err=", response.location, response.location)
        with self.conn.cursor() as cur:
            cur.execute("SELECT id,stable_key,export_policy,color_key,updated_at FROM regulatory_statuses WHERE label='Theo dõi UI'")
            status_id, stable_key, policy, color_key, revision = cur.fetchone()
        self.assertEqual(policy, "ALLOW")
        self.assertEqual(color_key, "gray")
        response = client.post("/admin/regulatory/statuses", data={
            "csrf_token": "phase6d1-csrf", "action": "set_color", "status_id": status_id,
            "revision": str(revision), "color_key": "purple",
        })
        self.assertEqual(response.status_code, 302)
        with self.conn.cursor() as cur:
            cur.execute("SELECT color_key,updated_at FROM regulatory_statuses WHERE id=%s", (status_id,))
            self.assertEqual((color_key := cur.fetchone())[0], "purple")
            revision = color_key[1]
        response = client.post("/admin/regulatory/statuses", data={
            "csrf_token": "phase6d1-csrf", "action": "rename", "status_id": status_id,
            "revision": str(revision), "label": "Theo dõi UI mới",
        })
        self.assertEqual(response.status_code, 302)
        with self.conn.cursor() as cur:
            cur.execute("SELECT stable_key,label,export_policy,color_key FROM regulatory_statuses WHERE id=%s", (status_id,))
            self.assertEqual(cur.fetchone(), (stable_key, "Theo dõi UI mới", "ALLOW", "purple"))
            cur.execute("SELECT event FROM regulatory_status_events WHERE status_id=%s ORDER BY id", (status_id,))
            self.assertEqual([row[0] for row in cur.fetchall()], ["created", "color_changed", "renamed"])
        template = client.get("/admin/regulatory/templates/cas")
        self.assertEqual(template.status_code, 200)
        wb = load_workbook(BytesIO(template.data), read_only=True)
        try:
            self.assertEqual([cell.value for cell in next(wb.active.iter_rows())],
                             ["CAS", "Tình trạng quản lý", "Ghi chú quản lý"])
        finally:
            wb.close()

    def test_color_validation_stale_guard_policy_independence_and_no_product_write(self):
        client = self._admin_client()
        self._rule("CAM_NHAP", "code", "COLOR-BLOCK", "Vẫn chặn theo chính sách")
        self._rule("CAM_NHAP", "cas", "50-00-0", "Màu chung cho tra cứu CAS")
        self._rule("DUOC_BAN", "code", "COLOR-ALLOW", "Vẫn cho xuất theo chính sách")
        with self.conn.cursor() as cur:
            cur.execute("""INSERT INTO products(name,code,brand,size,ship,price)
                           VALUES ('Blocked green','COLOR-BLOCK','Brand A','1g','1','100'),
                                  ('Allowed red','COLOR-ALLOW','Brand A','1g','1','100')""")
            cur.execute("INSERT INTO brand_compliance_settings VALUES ('BRAND A',true,now())")
            cur.execute("""INSERT INTO products(name,code,cas,brand,size,ship,price,manual_compliance)
                           VALUES ('Manual final colour','COLOR-MANUAL','50-00-0','Brand A','1g','1','100','Được bán')""")
            cur.execute("SELECT xmin::text FROM products ORDER BY code")
            product_versions = cur.fetchall()
            cur.execute("SELECT id,updated_at FROM regulatory_statuses WHERE stable_key='CAM_NHAP'")
            blocked_id, blocked_revision = cur.fetchone()
            cur.execute("SELECT id,updated_at FROM regulatory_statuses WHERE stable_key='DUOC_BAN'")
            allowed_id, allowed_revision = cur.fetchone()

        invalid = client.post("/admin/regulatory/statuses", data={
            "csrf_token": "phase6d1-csrf", "action": "set_color", "status_id": blocked_id,
            "revision": str(blocked_revision), "color_key": "url(javascript:bad)",
        })
        self.assertIn("err=", invalid.location)
        changed = client.post("/admin/regulatory/statuses", data={
            "csrf_token": "phase6d1-csrf", "action": "set_color", "status_id": blocked_id,
            "revision": str(blocked_revision), "color_key": "green",
        })
        self.assertNotIn("err=", changed.location)
        stale = client.post("/admin/regulatory/statuses", data={
            "csrf_token": "phase6d1-csrf", "action": "set_color", "status_id": blocked_id,
            "revision": str(blocked_revision), "color_key": "purple",
        })
        self.assertIn("err=", stale.location)
        allowed = client.post("/admin/regulatory/statuses", data={
            "csrf_token": "phase6d1-csrf", "action": "set_color", "status_id": allowed_id,
            "revision": str(allowed_revision), "color_key": "red",
        })
        self.assertNotIn("err=", allowed.location)

        rows = client.get("/search", query_string={"query": "COLOR-"}).get_json()["results"]
        by_code = {row["Code"]: row for row in rows}
        self.assertEqual(by_code["COLOR-BLOCK"]["Compliance_Css"], "regulatory-color-green")
        self.assertEqual(by_code["COLOR-BLOCK"]["Compliance_Export_Policy"], "BLOCK")
        self.assertEqual(by_code["COLOR-ALLOW"]["Compliance_Css"], "regulatory-color-red")
        self.assertEqual(by_code["COLOR-ALLOW"]["Compliance_Export_Policy"], "ALLOW")
        manual = client.get("/search", query_string={"query": "COLOR-MANUAL"}).get_json()["results"][0]
        self.assertEqual((manual["Compliance_Status"], manual["Compliance_Css"],
                          manual["Compliance_Export_Policy"]),
                         ("Được bán", "regulatory-color-red", "ALLOW"))
        found = client.post("/find_code_batch", data={"codes": "COLOR-BLOCK"}).get_json()["results"][0]
        self.assertEqual(found["Compliance_Css"], "regulatory-color-green")
        checked = client.post("/check_cas_batch", data={"cas": "50-00-0"}).get_json()["results"][0]
        self.assertEqual(checked["Compliance_Css"], "regulatory-color-green")
        quote = client.post("/api/quote-assistant/match", json={
            "rows": [{"code": "COLOR-BLOCK"}, {"code": "COLOR-ALLOW"}],
            "selection_strategy": "LOWEST_UNIT_PRICE",
        }).get_json()["results"]
        self.assertEqual(quote[0]["candidates"][0]["ineligible_reason"], "COMPLIANCE_BLOCKED")
        self.assertNotEqual(quote[1]["candidates"][0]["ineligible_reason"], "COMPLIANCE_BLOCKED")
        with self.conn.cursor() as cur:
            cur.execute("SELECT xmin::text FROM products ORDER BY code")
            self.assertEqual(cur.fetchall(), product_versions)

    def test_set_color_revalidates_and_locks_actor_after_regulatory_wait(self):
        cases = {
            "auth_version": "UPDATE app_users SET auth_version=2 WHERE id=%s",
            "demoted": "UPDATE app_users SET is_admin=false WHERE id=%s",
            "suspended": "UPDATE app_users SET account_status='SUSPENDED' WHERE id=%s",
        }
        for label, revoke_sql in cases.items():
            with self.subTest(label=label):
                with self.conn.cursor() as cur:
                    cur.execute("""UPDATE app_users
                                   SET is_admin=true,account_status='ACTIVE',auth_version=1
                                   WHERE id=%s""", (self.admin_id,))
                    cur.execute("""UPDATE regulatory_statuses
                                   SET color_key='red',updated_at=clock_timestamp()
                                   WHERE stable_key='CAM_NHAP' RETURNING id,updated_at""")
                    status_id, revision = cur.fetchone()
                    cur.execute("DELETE FROM regulatory_status_events WHERE status_id=%s", (status_id,))

                locker = psycopg2.connect(self.dsn)
                locker.autocommit = False
                with locker.cursor() as cur:
                    cur.execute("SELECT pg_advisory_xact_lock(%s)",
                                (regulatory_import_jobs.REGULATORY_LOCK_KEY,))
                reached_lock = threading.Event()
                real_acquire = admin_regulatory.acquire_regulatory_lock
                result = {}

                def marked_acquire(cur):
                    reached_lock.set()
                    return real_acquire(cur)

                def request_color_change():
                    result["response"] = self._real_admin_client().post(
                        "/admin/regulatory/statuses",
                        data={"csrf_token": "phase6d1-csrf", "action": "set_color",
                              "status_id": status_id, "revision": str(revision),
                              "color_key": "purple"},
                    )

                thread = threading.Thread(target=request_color_change)
                try:
                    with mock.patch.object(admin_regulatory, "acquire_regulatory_lock",
                                           side_effect=marked_acquire):
                        thread.start()
                        self.assertTrue(reached_lock.wait(5), "request never reached regulatory lock")
                        self.assertTrue(thread.is_alive(), "request did not block on regulatory lock")
                        updater = psycopg2.connect(self.dsn)
                        try:
                            with updater, updater.cursor() as cur:
                                cur.execute(revoke_sql, (self.admin_id,))
                        finally:
                            updater.close()
                        locker.rollback()
                        thread.join(10)
                finally:
                    locker.close()
                    if thread.is_alive():
                        thread.join(5)

                self.assertFalse(thread.is_alive(), "request hung after regulatory lock release")
                self.assertEqual(result["response"].status_code, 403)
                with self.conn.cursor() as cur:
                    cur.execute("SELECT color_key FROM regulatory_statuses WHERE id=%s", (status_id,))
                    self.assertEqual(cur.fetchone()[0], "red")
                    cur.execute("SELECT count(*) FROM regulatory_status_events WHERE status_id=%s",
                                (status_id,))
                    self.assertEqual(cur.fetchone()[0], 0)

        with self.conn.cursor() as cur:
            cur.execute("""UPDATE app_users SET is_admin=true,account_status='ACTIVE',auth_version=1
                           WHERE id=%s""", (self.admin_id,))
            cur.execute("SELECT id,updated_at FROM regulatory_statuses WHERE stable_key='CAM_NHAP'")
            status_id, revision = cur.fetchone()
        response = self._real_admin_client().post("/admin/regulatory/statuses", data={
            "csrf_token": "phase6d1-csrf", "action": "set_color", "status_id": status_id,
            "revision": str(revision), "color_key": "purple",
        })
        self.assertEqual(response.status_code, 302)
        with self.conn.cursor() as cur:
            cur.execute("SELECT color_key FROM regulatory_statuses WHERE id=%s", (status_id,))
            self.assertEqual(cur.fetchone()[0], "purple")
            cur.execute("""SELECT count(*) FROM regulatory_status_events
                           WHERE status_id=%s AND event='color_changed'""", (status_id,))
            self.assertEqual(cur.fetchone()[0], 1)

    def test_set_color_requires_csrf_and_admin_session(self):
        with self.conn.cursor() as cur:
            cur.execute("SELECT id,color_key,updated_at FROM regulatory_statuses WHERE stable_key='CAM_NHAP'")
            status_id, original_color, revision = cur.fetchone()
        admin = self._real_admin_client()
        missing_csrf = admin.post("/admin/regulatory/statuses", data={
            "action": "set_color", "status_id": status_id,
            "revision": str(revision), "color_key": "purple",
        })
        self.assertEqual(missing_csrf.status_code, 400)
        marker = secrets.token_hex(4)
        with self.conn.cursor() as cur:
            cur.execute("INSERT INTO teams(name,lifecycle_status) VALUES (%s,'ACTIVE') RETURNING id",
                        (f"Color guard {marker}",))
            team_id = cur.fetchone()[0]
            cur.execute("""INSERT INTO app_users
                               (username,password_hash,team_id,is_admin,account_status,auth_version)
                           VALUES (%s,'x',%s,false,'ACTIVE',1) RETURNING id""",
                        (f"color_guard_{marker}", team_id))
            staff_id = cur.fetchone()[0]
        staff = search.app.test_client()
        with staff.session_transaction() as sess:
            sess.update(authenticated=True,user_id=staff_id,auth_version=1,is_admin=False,
                        team_id=team_id,username=f"color_guard_{marker}",csrf_token="staff-csrf")
        nonadmin = staff.post("/admin/regulatory/statuses", data={
            "csrf_token": "staff-csrf", "action": "set_color", "status_id": status_id,
            "revision": str(revision), "color_key": "purple",
        })
        self.assertEqual(nonadmin.status_code, 403)
        with self.conn.cursor() as cur:
            cur.execute("SELECT color_key FROM regulatory_statuses WHERE id=%s", (status_id,))
            self.assertEqual(cur.fetchone()[0], original_color)

    def test_color_metadata_is_redacted_without_compliance_permission(self):
        self._rule("PHU_LUC_II", "code", "HIDDEN-COLOR", "regulatory note")
        with self.conn.cursor() as cur:
            cur.execute("""INSERT INTO products(name,code,brand,size,ship,price,note)
                           VALUES ('Visible product','HIDDEN-COLOR','Brand A','1g','1','100','Visible product note')""")
        grants = {"SEARCH", "FIND_CODE", "QUICK_QUOTE", "VIEW_NAME", "VIEW_CODE",
                  "VIEW_PRICE", "VIEW_NOTE"}
        client = self._team_client(grants)
        with client.session_transaction() as sess:
            team_id = sess["team_id"]
        with self.conn.cursor() as cur:
            cur.execute("INSERT INTO team_brands(team_id,brand) VALUES (%s,'Brand A')", (team_id,))
        responses = [
            client.get("/search", query_string={"query": "HIDDEN-COLOR"}),
            client.post("/find_code_batch", data={"codes": "HIDDEN-COLOR"}),
            client.post("/api/quote-assistant/match", json={"rows": [{"code": "HIDDEN-COLOR"}]}),
        ]
        for response in responses:
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
            raw = response.get_data(as_text=True)
            self.assertIn("Visible product note", raw)
            for forbidden in ("Compliance_Color", "compliance_color", "Compliance_Css",
                              "compliance_css", "regulatory-color-amber", "PHU_LUC_II"):
                self.assertNotIn(forbidden, raw)
        self.assertEqual(client.get("/check_cas", query_string={"cas": "50-00-0"}).status_code, 403)

    def test_background_preview_confirm_apply_and_audit(self):
        upload = _workbook_bytes(
            ["CAS", "Tình trạng quản lý", "Ghi chú quản lý"],
            [["50-00-0", "CẤM NHẬP", "background"]],
        )
        job_id = regulatory_import_jobs.submit(
            FileStorage(stream=upload, filename="rules.xlsx"), "upsert", "phase6d1_admin",
            self.admin_id, 1, str(__import__("uuid").uuid4()),
        )
        self.assertTrue(regulatory_import_jobs.run_once(job_id))
        with self.conn.cursor(cursor_factory=RealDictCursor) as cur:
            preview_job = regulatory_import_jobs.fetch_job(cur, job_id)
        self.assertEqual(preview_job["status"], "completed", preview_job["errors"])
        self.assertTrue(preview_job["preview_ready"])
        self.assertEqual(preview_job["preview"]["inserted"], 1)
        detail=self._admin_client().get(f"/admin/regulatory/jobs/{job_id}").get_data(as_text=True)
        for translated in ("Hoàn tất","Xem trước","Đã tiếp nhận tệp","Xem trước hoàn tất"):
            self.assertIn(translated,detail)
        for raw_cell in (">completed<",">preview<",">uploaded<",">preview_completed<"):
            self.assertNotIn(raw_cell,detail)
        regulatory_import_jobs.control(
            job_id, "apply", "phase6d1_admin", preview_job["preview"]["fingerprint"], ""
        )
        self.assertTrue(regulatory_import_jobs.run_once(job_id))
        with self.conn.cursor() as cur:
            cur.execute("SELECT note FROM regulatory_rules WHERE match_field='cas' AND match_value='50-00-0'")
            self.assertEqual(cur.fetchone()[0], "background")
            cur.execute("SELECT event FROM regulatory_import_events WHERE job_id=%s ORDER BY id", (job_id,))
            self.assertEqual([row[0] for row in cur.fetchall()],
                             ["uploaded", "preview_started", "preview_completed", "apply", "apply_started", "apply_completed"])

    def test_worker_lock_serializes_and_partial_failure_rolls_back(self):
        upload = _workbook_bytes(
            ["Code", "Tình trạng quản lý", "Ghi chú quản lý"],
            [["LOCKED", "Tình trạng tạm", "race"]],
        )
        job_id = regulatory_import_jobs.submit(
            FileStorage(stream=upload, filename="race.xlsx"), "upsert", "phase6d1_admin",
            self.admin_id, 1, str(__import__("uuid").uuid4()),
        )
        lock_conn = psycopg2.connect(self.dsn)
        lock_conn.autocommit = False
        with lock_conn.cursor() as cur:
            cur.execute("SELECT pg_advisory_xact_lock(%s)", (regulatory_import_jobs.REGULATORY_LOCK_KEY,))
        result = {}
        thread = threading.Thread(target=lambda: result.setdefault("worked", regulatory_import_jobs.run_once(job_id)))
        thread.start()
        thread.join(timeout=0.5)
        self.assertTrue(thread.is_alive(), "worker did not contend on the regulatory advisory lock")
        lock_conn.commit()
        lock_conn.close()
        thread.join(timeout=10)
        self.assertFalse(thread.is_alive())
        self.assertTrue(result.get("worked"))

        with self.conn.cursor(cursor_factory=RealDictCursor) as cur:
            preview_job = regulatory_import_jobs.fetch_job(cur, job_id)
        self.assertEqual(preview_job["status"], "completed", preview_job["errors"])
        regulatory_import_jobs.control(job_id, "apply", "phase6d1_admin",
                                       preview_job["preview"]["fingerprint"], "")

        real_apply = regulatory_import_jobs.apply_plan
        def fail_after_write(cur, rows, mode, expected):
            cur.execute("INSERT INTO regulatory_statuses(stable_key,label,priority,export_policy) VALUES ('ROLLBACK_TEST','Rollback test',9999,'ALLOW')")
            raise ImportProblem("forced rollback")
        try:
            regulatory_import_jobs.apply_plan = fail_after_write
            self.assertTrue(regulatory_import_jobs.run_once(job_id))
        finally:
            regulatory_import_jobs.apply_plan = real_apply
        with self.conn.cursor() as cur:
            cur.execute("SELECT 1 FROM regulatory_statuses WHERE stable_key='ROLLBACK_TEST'")
            self.assertIsNone(cur.fetchone())
            cur.execute("SELECT status,errors->>0 FROM regulatory_import_jobs WHERE id=%s", (job_id,))
            self.assertEqual(cur.fetchone(), ("failed", "forced rollback"))

    def test_cleanup_respects_live_job_lock_and_purges_only_terminal_jobs(self):
        completed=regulatory_import_jobs.submit(
            FileStorage(stream=_workbook_bytes(
                ['Code','Tình trạng quản lý','Ghi chú quản lý'],[['DONE','CẤM NHẬP','done']]),
                filename='done.xlsx'), 'upsert','phase6d1_admin',self.admin_id,1,
            str(__import__('uuid').uuid4()),
        )
        self.assertTrue(regulatory_import_jobs.run_once(completed))
        failed=regulatory_import_jobs.submit(
            FileStorage(stream=_workbook_bytes(
                ['CAS','Tình trạng quản lý','Ghi chú quản lý'],[['123-45-6','CẤM NHẬP','bad']]),
                filename='failed.xlsx'), 'upsert','phase6d1_admin',self.admin_id,1,
            str(__import__('uuid').uuid4()),
        )
        self.assertTrue(regulatory_import_jobs.run_once(failed))
        with self.conn.cursor() as cur:
            cur.execute("UPDATE regulatory_import_jobs SET expires_at=now()-interval '1 second' WHERE id=ANY(%s::uuid[])",([completed,failed],))
        terminal_paths=[regulatory_import_jobs.upload_path(item) for item in (completed,failed)]
        regulatory_import_jobs.cleanup()
        with self.conn.cursor() as cur:
            cur.execute("SELECT status,purged_at IS NOT NULL FROM regulatory_import_jobs WHERE id=ANY(%s::uuid[]) ORDER BY status",([completed,failed],))
            self.assertEqual(cur.fetchall(),[('completed',True),('failed',True)])
            cur.execute("SELECT count(*) FROM regulatory_import_events WHERE job_id=ANY(%s::uuid[])",([completed,failed],))
            self.assertGreater(cur.fetchone()[0],0)
        self.assertTrue(all(not path.exists() for path in terminal_paths))

        live=regulatory_import_jobs.submit(
            FileStorage(stream=_workbook_bytes(
                ['Code','Tình trạng quản lý','Ghi chú quản lý'],[['LIVE','CẤM NHẬP','live']]),
                filename='live.xlsx'), 'upsert','phase6d1_admin',self.admin_id,1,
            str(__import__('uuid').uuid4()),
        )
        self.assertTrue(regulatory_import_jobs.run_once(live))
        with self.conn.cursor(cursor_factory=RealDictCursor) as cur:
            live_preview=regulatory_import_jobs.fetch_job(cur,live)['preview']
        regulatory_import_jobs.control(live,'apply','phase6d1_admin',live_preview['fingerprint'],'')
        live_path=regulatory_import_jobs.upload_path(live)
        lock=psycopg2.connect(self.dsn)
        with lock.cursor() as cur:
            cur.execute('SELECT pg_advisory_xact_lock(%s)',(regulatory_import_jobs.REGULATORY_LOCK_KEY,))
        thread=threading.Thread(target=lambda:regulatory_import_jobs.run_once(live))
        thread.start()
        deadline=time.monotonic()+5
        while time.monotonic()<deadline:
            with self.conn.cursor() as cur:
                cur.execute('SELECT status FROM regulatory_import_jobs WHERE id=%s',(live,))
                if cur.fetchone()[0]=='running':break
            time.sleep(.02)
        with self.conn.cursor() as cur:
            cur.execute("UPDATE regulatory_import_jobs SET expires_at=now()-interval '1 second' WHERE id=%s",(live,))
        regulatory_import_jobs.cleanup()
        with self.conn.cursor() as cur:
            cur.execute('SELECT cancel_requested,purged_at FROM regulatory_import_jobs WHERE id=%s',(live,))
            cancel_requested,purged_at=cur.fetchone()
        self.assertTrue(cancel_requested)
        self.assertIsNone(purged_at)
        self.assertTrue(live_path.exists())
        lock.rollback();lock.close();thread.join(10)
        self.assertFalse(thread.is_alive())
        with self.conn.cursor() as cur:
            cur.execute('SELECT status FROM regulatory_import_jobs WHERE id=%s',(live,))
            self.assertEqual(cur.fetchone()[0],'cancelled')
            cur.execute("SELECT count(*) FROM regulatory_rules WHERE match_field='code' AND match_value='LIVE'")
            self.assertEqual(cur.fetchone()[0],0)
        regulatory_import_jobs.cleanup()
        self.assertFalse(live_path.exists())
        with self.conn.cursor() as cur:
            cur.execute('SELECT purged_at IS NOT NULL FROM regulatory_import_jobs WHERE id=%s',(live,))
            self.assertTrue(cur.fetchone()[0])


class Phase6D1PureTests(unittest.TestCase):
    def test_cas_format_and_checksum(self):
        self.assertEqual(normalize_cas(" 50-00-0 "), "50-00-0")
        for value in ("123-45-6", "50 00 0", "50-00-1", ""):
            with self.assertRaises(ValueError, msg=value):
                normalize_cas(value)

    def test_name_sql_is_substring_and_no_field_order(self):
        sql = product_resolver_lateral("p", "bcs")
        self.assertIn("strpos(upper(btrim(p.name)), upper(btrim(r.match_value))) > 0", sql)
        self.assertIn("ORDER BY priority, status_id", sql)
        self.assertNotIn("CASE r.match_field", sql)

    def test_text_identity_preserves_accents(self):
        self.assertNotEqual(normalized_identity("Hóa chất"), normalized_identity("Hoa chat"))

    def test_frontend_uses_api_palette_and_workbook_mapping_excludes_color(self):
        search_js = (ROOT / "static" / "script.js").read_text(encoding="utf-8")
        quote_js = (ROOT / "static" / "quick_quote.js").read_text(encoding="utf-8")
        css = (ROOT / "static" / "styles.css").read_text(encoding="utf-8")
        self.assertNotIn("COMPLIANCE_CLASS", search_js)
        self.assertNotIn("qqComplianceClass(compLabel)", quote_js)
        self.assertIn("td:not(.qq-cell-lifecycle)", css)
        mapping = search._quote_template_mapping_snapshot()["mapping"]
        self.assertNotIn("Compliance_Color", mapping)
        self.assertNotIn("Compliance_Css", mapping)


class Phase6D1LegacyMigrationRehearsal(unittest.TestCase):
    @unittest.skipUnless(pg_temp_db.probe_postgres_reachable(), "isolated local PostgreSQL required")
    def test_pre026_transition_removes_only_ton_kho_and_is_rerunnable(self):
        db_name = pg_temp_db._TEST_DB_PREFIX + "phase6d1_migration_" + secrets.token_hex(4)
        dsn = pg_temp_db.dsn_for(db_name)
        maint = psycopg2.connect(pg_temp_db.maintenance_dsn())
        maint.autocommit = True
        with maint.cursor() as cur:
            cur.execute(f'CREATE DATABASE "{db_name}"')
        maint.close()
        try:
            conn = psycopg2.connect(dsn)
            with conn:
                with conn.cursor() as cur:
                    cur.execute(pg_temp_db._MINIMAL_BASE_SCHEMA_SQL)
                    for name in pg_temp_db._FULL_SCHEMA_SQL_FILES:
                        if name in {"migration_026_regulatory_management.sql",
                                  "migration_027_regulatory_status_colors.sql"}:
                            continue
                        cur.execute(pg_temp_db._read_sql(name))
                    cur.execute("INSERT INTO regulatory_rules(rule_type,rule_label,match_field,match_value,priority,is_active,note) VALUES ('CAM_NHAP','old label','cas','50-00-0',999,true,'keep'),('TON_KHO','TỒN KHO','code','STOCK',1,true,'remove')")
                    cur.execute("INSERT INTO products(name,code,brand,manual_compliance) VALUES ('P','P1','B','Được bán')")
                    cur.execute(MIGRATION_026)
                    cur.execute(MIGRATION_026)
                    cur.execute("SELECT stable_key,label,priority,export_policy FROM regulatory_statuses WHERE stable_key IN ('CAM_NHAP','PHU_LUC_II','PHU_LUC_III') ORDER BY priority")
                    self.assertEqual(cur.fetchall(), [
                        ("CAM_NHAP","CẤM NHẬP",10,"BLOCK"),
                        ("PHU_LUC_II","Phụ lục II",20,"ALLOW"),
                        ("PHU_LUC_III","Phụ lục III",30,"ALLOW"),
                    ])
                    cur.execute("SELECT rule_type FROM regulatory_rules ORDER BY id")
                    self.assertEqual(cur.fetchall(), [("CAM_NHAP",)])
                    cur.execute("SELECT manual_compliance_status_id IS NOT NULL FROM products WHERE code='P1'")
                    self.assertTrue(cur.fetchone()[0])
            conn.close()
        finally:
            pg_temp_db.drop_temp_db(db_name)

    @unittest.skipUnless(pg_temp_db.probe_postgres_reachable(), "isolated local PostgreSQL required")
    def test_027_runs_twice_from_exact_026_schema_without_product_writes(self):
        db_name = pg_temp_db._TEST_DB_PREFIX + "phase6d11_migration_" + secrets.token_hex(4)
        dsn = pg_temp_db.dsn_for(db_name)
        maint = psycopg2.connect(pg_temp_db.maintenance_dsn())
        maint.autocommit = True
        with maint.cursor() as cur:
            cur.execute(f'CREATE DATABASE "{db_name}"')
        maint.close()
        try:
            conn = psycopg2.connect(dsn)
            with conn:
                with conn.cursor() as cur:
                    cur.execute(pg_temp_db._MINIMAL_BASE_SCHEMA_SQL)
                    for name in pg_temp_db._FULL_SCHEMA_SQL_FILES:
                        if name in {"migration_026_regulatory_management.sql",
                                  "migration_027_regulatory_status_colors.sql"}:
                            continue
                        cur.execute(pg_temp_db._read_sql(name))
                    cur.execute("INSERT INTO products(name,code,brand) VALUES ('P','NO-WRITE','B')")
                    cur.execute(MIGRATION_026)
                    cur.execute("SELECT xmin::text FROM products WHERE code='NO-WRITE'")
                    product_version = cur.fetchone()[0]
                    cur.execute("SELECT id,label,priority FROM regulatory_statuses WHERE stable_key='CAM_NHAP'")
                    status_id, label, priority = cur.fetchone()
                    cur.execute("""INSERT INTO regulatory_rules
                                       (rule_type,rule_label,match_field,match_value,priority,is_active,status_id)
                                   VALUES ('CAM_NHAP',%s,'code','NO-WRITE',%s,true,%s)""",
                                (label, priority, status_id))
                    cur.execute(f"""SELECT rr.rule_label,rr.export_policy,rr.source,
                                            rr.stable_key,rr.status_id,rr.color_key
                                     FROM products p {product_resolver_lateral('p', manual_enabled_expr='FALSE')}
                                     WHERE p.code='NO-WRITE'""")
                    pre_migration_result = cur.fetchone()
                    self.assertIsNone(pre_migration_result[5])
                    resolved = search.resolve_compliance_precedence(
                        brand_manual_enabled=False, manual_compliance=None,
                        manual_compliance_note=None, legacy_compliance=pre_migration_result[0],
                        legacy_compliance_note=None, cas=None,
                        export_policy=pre_migration_result[1], source=pre_migration_result[2],
                        stable_key=pre_migration_result[3], status_id=pre_migration_result[4],
                        color_key=pre_migration_result[5],
                    )
                    self.assertEqual(resolved["compliance_css"], "regulatory-color-red")
                    cur.execute("""INSERT INTO app_users
                                       (username,password_hash,is_admin,account_status,auth_version)
                                   VALUES ('pre027-admin','x',true,'ACTIVE',1) RETURNING id""")
                    pre027_admin_id = cur.fetchone()[0]
                    cur.execute("SELECT updated_at FROM regulatory_statuses WHERE id=%s", (status_id,))
                    pre027_revision = cur.fetchone()[0]
                conn.commit()
                with mock.patch.dict(os.environ, {"DATABASE_URL": dsn, "DISABLE_IP_ALLOWLIST": "1"}):
                    search.app.testing = True
                    client = search.app.test_client()
                    with client.session_transaction() as sess:
                        sess.update(authenticated=True, user_id=pre027_admin_id, auth_version=1,
                                    is_admin=True, username="pre027-admin", csrf_token="pre027-csrf")
                    response = client.post("/admin/regulatory/statuses", data={
                        "csrf_token": "pre027-csrf", "action": "set_color",
                        "status_id": status_id, "revision": str(pre027_revision), "color_key": "blue",
                    })
                    self.assertIn("migration+027", response.location)
                with conn.cursor() as cur:
                    cur.execute(MIGRATION_027)
                    cur.execute("""SELECT stable_key,color_key FROM regulatory_statuses
                                   WHERE stable_key IN ('CAM_NHAP','PHU_LUC_II','PHU_LUC_III','DUOC_BAN')
                                   ORDER BY stable_key""")
                    self.assertEqual(cur.fetchall(), [
                        ("CAM_NHAP", "red"), ("DUOC_BAN", "green"),
                        ("PHU_LUC_II", "amber"), ("PHU_LUC_III", "teal"),
                    ])
                    cur.execute("UPDATE regulatory_statuses SET color_key='purple' WHERE stable_key='CAM_NHAP'")
                    cur.execute(MIGRATION_027)
                    cur.execute("SELECT color_key FROM regulatory_statuses WHERE stable_key='CAM_NHAP'")
                    self.assertEqual(cur.fetchone()[0], "purple")
                    cur.execute("SELECT xmin::text FROM products WHERE code='NO-WRITE'")
                    self.assertEqual(cur.fetchone()[0], product_version)
                    with self.assertRaises(psycopg2.errors.CheckViolation):
                        with conn.cursor() as invalid_cur:
                            invalid_cur.execute("UPDATE regulatory_statuses SET color_key='raw-css' WHERE stable_key='CAM_NHAP'")
            conn.close()
        finally:
            pg_temp_db.drop_temp_db(db_name)


if __name__ == "__main__":
    unittest.main()
