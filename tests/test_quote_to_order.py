"""Menu "Báo giá → Đơn hàng": đọc báo giá, ghép cột, tính giá, tạo file đơn hàng, route/quyền/CSRF.

Không chạm Postgres thật: phiên đăng nhập dùng `auth_test_helpers`, IP allowlist tắt như các test route khác.
LibreOffice không có trên máy dev, nên file "có giá trị tính sẵn" được dựng bằng cách chèn <v> vào XML
của ô công thức (đúng dạng Excel lưu).
"""
import io
import re
import unittest
import zipfile
from decimal import Decimal
from unittest import mock

from openpyxl import Workbook, load_workbook

import search  # noqa: E402
import quote_to_order as q  # noqa: E402
import team_permissions  # noqa: E402
from auth_test_helpers import start_auth_db_patch  # noqa: E402

RENAMES = {"Tên hàng": "Mặt hàng", "Code": "Mã số"}
FORM_HEADERS = [
    "TT", "Tên hàng", "Code", "Cas", "Hãng", "Đơn vị tính", "Số lượng", "Đơn giá có VAT (VNĐ)",
    "Đơn giá chưa VAT (VNĐ)", "Thành tiền gồm VAT (VNĐ)", "Thời gian đặt hàng", "Đơn giá trước giảm (Nếu có)",
    "Ghi chú hàng hóa", "Ghi chú khác", "Ghi chú nội bộ", "Phụ Lục", "Giá bán chưa VAT", "%LN",
    "Giá nhập chưa VAT", "Loại Hàng", "Kho yêu cầu",
]


def make_quote(items, *, header_row=10, formulas=True, cached=None, footer=("Tổng giá gồm VAT", "Thuế VAT 8%"),
               headers=None, sheet="BG", after_footer=True, renames=None):
    """Dựng form báo giá giống ảnh PO gửi. `items`: dict các trường; công thức H/I/Q như form thật."""
    headers = headers or FORM_HEADERS
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(row=1, column=2, value="CÔNG TY TEST")
    for c, title in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=(renames or {}).get(title, title))
    pos = {title: i for i, title in enumerate(headers, start=1)}
    row = header_row + 1
    for n, item in enumerate(items, start=1):
        ws.cell(row=row, column=1, value=n)
        mapping = {"name": "Tên hàng", "code": "Code", "cas": "Cas", "brand": "Hãng", "unit": "Đơn vị tính",
                   "qty": "Số lượng", "note_goods": "Ghi chú hàng hóa", "note_other": "Ghi chú khác",
                   "internal": "Ghi chú nội bộ", "cost": "Giá nhập chưa VAT", "margin": "%LN", "type": "Loại Hàng",
                   "warehouse": "Kho yêu cầu",
                   "phu_luc": "Phụ Lục"}
        for key, title in mapping.items():
            if key in item and title in pos:
                ws.cell(row=row, column=pos[title], value=item[key])
        if formulas and "Giá nhập chưa VAT" in pos and item.get("cost") is not None:
            S, R, Q = (get(pos[t], row) for t in ("Giá nhập chưa VAT", "%LN", "Giá bán chưa VAT"))
            ws.cell(row=row, column=pos["Giá bán chưa VAT"], value=f"=ROUNDUP({S}/(1-{R}),-4)")
            ws.cell(row=row, column=pos["Đơn giá chưa VAT (VNĐ)"], value=f"={Q}")
            I = get(pos["Đơn giá chưa VAT (VNĐ)"], row)
            ws.cell(row=row, column=pos["Đơn giá có VAT (VNĐ)"], value=f"=ROUND({I}*1.08,0)")
        row += 1
    if footer:
        ws.cell(row=row, column=1, value=footer[0])
        ws.cell(row=row + 1, column=1, value=footer[1])
        if after_footer:
            ws.cell(row=row + 2, column=1, value="Ghi chú")
            ws.cell(row=row + 2, column=2, value="Hàng sau phần tổng không được đọc")
            ws.cell(row=row + 2, column=3, value="ZZ-AFTER")
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()
    return add_cached_values(raw, cached) if cached else raw


def get(col, row):
    from openpyxl.utils import get_column_letter
    return f"{get_column_letter(col)}{row}"


def add_cached_values(raw, values):
    """Chèn giá trị tính sẵn vào ô công thức: {'H11': 10800000}."""
    src, out = zipfile.ZipFile(io.BytesIO(raw)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                for ref, value in values.items():
                    pattern = re.compile(r'(<c r="%s"[^>]*>\s*<f>[^<]*</f>)\s*(?:<v\s*/>|<v></v>)' % ref)
                    text, n = pattern.subn(r"\g<1><v>%s</v>" % value, text)
                    assert n == 1, f"không chèn được giá trị tính sẵn cho {ref}"
                data = text.encode("utf-8")
            dst.writestr(info, data)
    return out.getvalue()


ITEM = dict(name="Benzene", code="M-502-01N", cas="71-43-2", brand="AccuStandard", unit="1g", qty=3,
            cost=8500000, margin=0.15, type="Nhập khẩu", warehouse="Hà Nội", note_goods="HSD 12 tháng", note_other="Có phép",
            internal="BÍ MẬT NỘI BỘ", phu_luc="PL-1")


class RoundUpTests(unittest.TestCase):
    def test_excel_roundup_boundary_cases(self):
        cases = [(120000, "0.25", 160000), (300000, "0.2", 380000),
                 (8500000, "0.15", 10000000), (1000000, "0.3", 1430000)]
        for cost, margin, expected in cases:
            with self.subTest(cost=cost, margin=margin):
                self.assertEqual(q.calc_price_ex_vat(Decimal(cost), Decimal(margin)), expected)

    def test_roundup_rounds_away_from_zero(self):
        self.assertEqual(q.excel_roundup(Decimal("10000.01"), -4), Decimal(20000))
        self.assertEqual(q.excel_roundup(Decimal("10000"), -4), Decimal(10000))
        self.assertEqual(q.excel_roundup(Decimal("2.341"), 2), Decimal("2.35"))

    def test_invalid_margin_returns_none(self):
        self.assertIsNone(q.calc_price_ex_vat(Decimal(100), Decimal(1)))
        self.assertIsNone(q.calc_price_ex_vat(Decimal(100), Decimal("-0.1")))

    def test_number_text_parsing(self):
        for text, expected in [("1.250.000", 1250000), ("0,5", Decimal("0.5")), ("1.5", Decimal("1.5")),
                               ("1.250.000,5", Decimal("1250000.5")), ("12", 12), (" 3 ", 3)]:
            self.assertEqual(q.parse_number_text(text), Decimal(expected), text)
        for bad in ("abc", "", "1.2.3,4,5", "--1"):
            self.assertIsNone(q.parse_number_text(bad), bad)


class ParseQuoteTests(unittest.TestCase):
    def test_parse_without_cached_values_recomputes_price(self):
        result = q.parse_quote(make_quote([ITEM]))
        self.assertEqual(result["header_row"], 10)
        item = result["items"][0]
        self.assertEqual(item["price"], 10800000)        # ROUNDUP(8.5tr/0.85,-4)=10tr, x1.08
        self.assertEqual(item["cost"], 8500000)
        self.assertEqual(item["qty"], 3)

    def test_parse_with_cached_values_uses_stored_numbers(self):
        raw = make_quote([ITEM], cached={"H11": 11000000, "I11": 10185185, "Q11": 10185185})
        item = q.parse_quote(raw)["items"][0]
        self.assertEqual(item["price"], 11000000)         # lấy số Excel đã tính, không tính lại

    def test_cached_ex_vat_price_is_used_for_recompute(self):
        raw = make_quote([ITEM], cached={"I11": 10000000, "Q11": 10000000})
        item = q.parse_quote(raw)["items"][0]
        self.assertEqual(item["price"], 10800000)         # H chưa có số -> 10tr x 1.08

    def test_vat_rate_is_read_from_footer(self):
        raw = make_quote([ITEM], footer=("Tổng giá gồm VAT", "Thuế VAT 10%"))
        result = q.parse_quote(raw)
        self.assertEqual(result["vat_rate"], 0.1)
        self.assertEqual(result["items"][0]["price"], 11000000)

    def test_placeholder_rows_and_footer_are_not_read_as_items(self):
        items = [ITEM, dict(qty=None), dict(name="Hóa chất B", code="B-1", brand="X", unit="g", qty=1)]
        parsed = q.parse_quote(make_quote(items))["items"]
        self.assertEqual([i["code"] for i in parsed], ["M-502-01N", "B-1"])
        self.assertNotIn("ZZ-AFTER", [i["code"] for i in parsed])

    def test_header_is_found_by_content_at_any_row(self):
        for header_row in (3, 10, 16):
            with self.subTest(header_row=header_row):
                result = q.parse_quote(make_quote([ITEM], header_row=header_row))
                self.assertEqual(result["header_row"], header_row)
                self.assertEqual(len(result["items"]), 1)

    def test_default_header_row_is_16_and_other_rows_raise_a_warning(self):
        at16 = q.parse_quote(make_quote([ITEM], header_row=16))
        self.assertEqual((at16["header_row"], at16["header_mismatch"], at16["default_header_row"]), (16, False, 16))
        for row in (11, 13):
            with self.subTest(header_row=row):
                found = q.parse_quote(make_quote([ITEM], header_row=row))
                self.assertEqual((found["header_row"], found["header_mismatch"]), (row, True))
                self.assertEqual(len(found["items"]), 1)             # vẫn đọc được, chỉ cảnh báo

    def test_row_16_wins_over_an_earlier_lookalike_header(self):
        wb = load_workbook(io.BytesIO(make_quote([ITEM], header_row=16)))
        ws = wb["BG"]
        ws["B5"], ws["C5"] = "Tên hàng", "Code"                       # dòng giống tiêu đề ở phía trên
        buf = io.BytesIO()
        wb.save(buf)
        result = q.parse_quote(buf.getvalue())
        self.assertEqual(result["header_row"], 16)
        self.assertFalse(result["header_mismatch"])

    def test_user_chosen_header_row_is_not_flagged_again(self):
        raw = make_quote([ITEM], header_row=13)
        self.assertTrue(q.parse_quote(raw)["header_mismatch"])
        confirmed = q.parse_quote(raw, header_row=13)
        self.assertEqual(confirmed["header_row"], 13)
        self.assertFalse(confirmed["header_mismatch"])

    def test_each_file_keeps_its_own_header_row(self):
        a = q.parse_quote(make_quote([ITEM], header_row=11))
        b = q.parse_quote(make_quote([ITEM], header_row=13))
        self.assertEqual((a["header_row"], b["header_row"]), (11, 13))

    def test_form_without_cas_column_still_parses(self):
        headers = [h for h in FORM_HEADERS if h != "Cas"]
        result = q.parse_quote(make_quote([ITEM], headers=headers, header_row=16))
        self.assertEqual(result["items"][0]["cas"], "")
        self.assertEqual(result["items"][0]["code"], "M-502-01N")
        self.assertIsNone(result["mapping"]["cas"])

    def test_missing_header_asks_for_header_row_and_manual_row_works(self):
        raw = make_quote([ITEM], renames=RENAMES)
        result = q.parse_quote(raw)
        self.assertTrue(result["needs_header"])
        manual = q.parse_quote(raw, header_row=10)
        self.assertFalse(manual["needs_header"])
        self.assertEqual(manual["unmapped_required"][:2], ["name", "code"])

    def test_manual_mapping_overrides_auto_detection(self):
        raw = make_quote([ITEM], renames=RENAMES)
        first = q.parse_quote(raw, header_row=10)
        cols = {c["title"]: c["index"] for c in first["columns"]}
        mapping = {"name": cols["Mặt hàng"], "code": cols["Mã số"]}
        result = q.parse_quote(raw, header_row=10, mapping=mapping)
        self.assertEqual(result["items"][0]["name"], "Benzene")
        self.assertEqual(result["items"][0]["code"], "M-502-01N")
        self.assertEqual(result["unmapped_required"], [])

    def test_mapping_none_blanks_field_and_bad_index_is_ignored(self):
        result = q.parse_quote(make_quote([ITEM]), mapping={"brand": None, "unit": 999, "qty": "3"})
        item = result["items"][0]
        self.assertEqual(item["brand"], "")
        self.assertEqual(item["unit"], "")
        self.assertIsNone(item["qty"])

    def test_price_can_be_mapped_to_ex_vat_column(self):
        raw = make_quote([ITEM])
        cols = {c["title"]: c["index"] for c in q.parse_quote(raw)["columns"]}
        item = q.parse_quote(raw, mapping={"price": cols["Đơn giá chưa VAT (VNĐ)"]})["items"][0]
        self.assertEqual(item["price"], 10000000)

    def test_warehouse_column_is_detected_and_normalized(self):
        items = [dict(ITEM, warehouse="Hà Nội"), dict(ITEM, warehouse="hồ chí minh"), dict(ITEM, warehouse="HCM"),
                 dict(ITEM, warehouse="ha noi"), dict(ITEM, warehouse="Đà Nẵng"), dict(ITEM, warehouse=None)]
        result = q.parse_quote(make_quote(items))
        self.assertIsNotNone(result["auto_mapping"]["warehouse"])
        self.assertEqual([i["warehouse"] for i in result["items"]],
                         ["Hà Nội", "Hồ Chí Minh", "Hồ Chí Minh", "Hà Nội", "", ""])

    def test_quote_without_warehouse_column_leaves_it_for_the_user(self):
        headers = [h for h in FORM_HEADERS if h != "Kho yêu cầu"]
        result = q.parse_quote(make_quote([ITEM], headers=headers))
        self.assertIn("warehouse", result["unmapped_required"])
        self.assertEqual(result["items"][0]["warehouse"], "")

    def test_code_and_cas_stay_text_and_types_are_normalized(self):
        items = [dict(name="A", code="00123", cas="64-17-5", brand="b", unit="g", qty=1, type="nhập khẩu"),
                 dict(name="B", code=456.0, cas="1-2-3", brand="b", unit="g", qty=1, type="mua trong nước"),
                 dict(name="C", code="C", brand="b", unit="g", qty=1, type="khác")]
        parsed = q.parse_quote(make_quote(items))["items"]
        self.assertEqual([i["code"] for i in parsed], ["00123", "456", "C"])
        self.assertTrue(all(isinstance(i["code"], str) for i in parsed))
        self.assertEqual([i["type"] for i in parsed], ["Nhập khẩu", "Mua trong nước", ""])

    def test_multiple_sheets_prefers_bg_and_allows_choosing(self):
        wb = load_workbook(io.BytesIO(make_quote([ITEM])))
        other = wb.create_sheet("Khác", 0)
        other["A1"] = "x"
        buf = io.BytesIO()
        wb.save(buf)
        result = q.parse_quote(buf.getvalue())
        self.assertEqual(result["sheet"], "BG")
        self.assertEqual(sorted(result["sheets"]), ["BG", "Khác"])
        self.assertTrue(q.parse_quote(buf.getvalue(), sheet="Khác")["needs_header"])

    def test_invalid_files_are_rejected_with_vietnamese_errors(self):
        for raw in (b"", b"not a zip at all" * 20, b"PK" + b"\x00" * 100):
            with self.subTest(size=len(raw)), self.assertRaises(q.QuoteParseError):
                q.parse_quote(raw)

    def test_macro_workbook_is_rejected(self):
        raw = make_quote([ITEM])
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w") as dst:
            for info in src.infolist():
                dst.writestr(info, src.read(info.filename))
            dst.writestr("xl/vbaProject.bin", b"x")
        with self.assertRaises(q.QuoteParseError):
            q.parse_quote(out.getvalue())


def valid_row(**overrides):
    row = dict(name="Benzene", code="M-502-01N", cas="71-43-2", brand="AccuStandard", unit="1g", qty=2,
               price=10800000, cost=8500000, min_price=None, type="Nhập khẩu", warehouse="Hà Nội",
               note_goods="HSD 12 tháng", note_other="Có phép")
    row.update(overrides)
    return row


class ExportTests(unittest.TestCase):
    def build(self, rows):
        clean, errors = q.clean_order_rows(rows)
        self.assertEqual(errors, [])
        return load_workbook(io.BytesIO(q.build_order_file(clean)))

    def test_output_keeps_template_headers_and_both_sheets(self):
        template = load_workbook(q.ORDER_TEMPLATE)
        wb = self.build([valid_row()])
        self.assertEqual(wb.sheetnames, template.sheetnames)
        self.assertEqual([c.value for c in wb["Exported file"][1]], [c.value for c in template["Exported file"][1]])

    def test_rows_amount_and_types(self):
        wb = self.build([valid_row(qty=2), valid_row(code="X-9", qty="0,5", price="1.234.567")])
        ws = wb["Exported file"]
        header = {c.value: c.column for c in ws[1]}
        self.assertEqual(ws.max_row, 3)
        r2, r3 = 2, 3
        self.assertEqual(ws.cell(r2, header["Số lượng (*)"]).value, 2)
        self.assertEqual(ws.cell(r2, header["Thành tiền"]).value, 21600000)    # SL x Đơn giá, ghi dạng số
        self.assertEqual(ws.cell(r3, header["Thành tiền"]).value, 617283.5)
        for r in (r2, r3):
            self.assertEqual(ws.cell(r, header["Thành tiền"]).data_type, "n")   # giá trị, không phải công thức
            self.assertEqual(ws.cell(r, header["Số lượng (*)"]).data_type, "n")
            self.assertEqual(ws.cell(r, header["Đơn giá (*)"]).data_type, "n")
            self.assertEqual(ws.cell(r, header["Code (*)"]).data_type, "s")
            self.assertEqual(ws.cell(r, header["Cas"]).number_format, "@")
        self.assertEqual(ws.cell(r2, header["Đơn giá mua dự kiến (Sales)"]).value, 8500000)
        self.assertEqual(ws.cell(r2, header["Giá bán tối thiểu"]).value, 10800000)   # trống => bằng Đơn giá
        self.assertEqual(ws.cell(r2, header["Loại hàng (*)"]).value, "Nhập khẩu")

    def test_min_price_defaults_to_unit_price_but_keeps_explicit_value(self):
        wb = self.build([valid_row(min_price=None), valid_row(min_price="9.000.000"), valid_row(min_price=0)])
        ws = wb["Exported file"]
        header = {c.value: c.column for c in ws[1]}
        self.assertEqual([ws.cell(r, header["Giá bán tối thiểu"]).value for r in (2, 3, 4)],
                         [10800000, 9000000, 0])
        self.assertEqual(ws.cell(2, header["Đơn giá (*)"]).value, ws.cell(2, header["Giá bán tối thiểu"]).value)

    def test_empty_or_zero_min_price_in_quote_file_is_treated_as_blank(self):
        raw = make_quote([ITEM], headers=FORM_HEADERS + ["Giá bán tối thiểu"])
        cols = {c["title"]: c["index"] for c in q.parse_quote(raw)["columns"]}
        self.assertEqual(q.parse_quote(raw)["auto_mapping"]["min_price"], cols["Giá bán tối thiểu"])
        self.assertIsNone(q.parse_quote(raw)["items"][0]["min_price"])

    def test_warehouse_is_written_to_the_last_column(self):
        wb = self.build([valid_row(warehouse="Hồ Chí Minh"), valid_row(warehouse="Hà Nội")])
        ws = wb["Exported file"]
        header = {c.value: c.column for c in ws[1]}
        self.assertEqual(list(header)[-1], "Kho yêu cầu (*)")
        self.assertEqual([ws.cell(r, header["Kho yêu cầu (*)"]).value for r in (2, 3)], ["Hồ Chí Minh", "Hà Nội"])
        self.assertEqual(ws.cell(2, header["Kho yêu cầu (*)"]).data_type, "s")

    def test_text_starting_with_equals_is_not_a_formula(self):
        wb = self.build([valid_row(name="=1+1", note_other="@SUM(A1)")])
        cell = wb["Exported file"]["A2"]
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")

    def test_order_qty_overrides_quote_qty_end_to_end(self):
        parsed = q.parse_quote(make_quote([ITEM]))["items"][0]
        parsed["qty"] = 7
        clean, errors = q.clean_order_rows([parsed])
        self.assertEqual(errors, [])
        ws = load_workbook(io.BytesIO(q.build_order_file(clean)))["Exported file"]
        self.assertEqual(ws["F2"].value, 7)
        self.assertEqual(ws["H2"].value, 7 * 10800000)

    def test_internal_note_is_never_in_the_order_file(self):
        parsed = q.parse_quote(make_quote([ITEM]))["items"]
        self.assertNotIn("BÍ MẬT NỘI BỘ", repr(parsed))
        clean, _ = q.clean_order_rows(parsed)
        ws = load_workbook(io.BytesIO(q.build_order_file(clean)))["Exported file"]
        values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertNotIn("BÍ MẬT NỘI BỘ", values)

    def test_validation_flags_missing_and_invalid_fields(self):
        rows = [valid_row(), valid_row(name="", code=" "), valid_row(qty=0), valid_row(qty=None), valid_row(qty=-1),
                valid_row(price=None), valid_row(type=""), valid_row(type="Khác"), valid_row(qty="abc"),
                valid_row(unit="", brand=""), valid_row(warehouse=""), valid_row(warehouse="Đà Nẵng")]
        clean, errors = q.clean_order_rows(rows)
        self.assertEqual(len(clean), 1)
        by_index = {e["index"]: e for e in errors}
        self.assertEqual(sorted(by_index), list(range(1, 12)))
        self.assertIn("Kho yêu cầu", by_index[10]["missing"])
        self.assertIn("Kho yêu cầu", by_index[11]["invalid"])
        self.assertEqual(by_index[1]["missing"], ["Tên hàng", "Code"])
        self.assertIn("Số lượng", by_index[2]["invalid"])
        self.assertIn("Số lượng", by_index[3]["missing"])
        self.assertIn("Số lượng", by_index[4]["invalid"])
        self.assertIn("Đơn giá", by_index[5]["missing"])
        self.assertIn("Loại hàng", by_index[6]["missing"])
        self.assertIn("Loại hàng", by_index[7]["invalid"])
        self.assertIn("Số lượng", by_index[8]["invalid"])
        self.assertEqual(by_index[9]["missing"], ["Hãng", "ĐVT"])

    def test_validation_rejects_oversized_and_non_scalar_values(self):
        _, errors = q.clean_order_rows([valid_row(name="x" * 501), valid_row(code=["a"]), "not a dict",
                                        valid_row(qty=float("inf")), valid_row(price=10 ** 16)])
        self.assertEqual(len(errors), 5)

    def test_control_characters_are_stripped(self):
        clean, errors = q.clean_order_rows([valid_row(name="A\x00B\x07C")])
        self.assertEqual(errors, [])
        self.assertEqual(clean[0]["name"], "ABC")
        q.build_order_file(clean)  # không ném IllegalCharacterError


class RouteTests(unittest.TestCase):
    CSRF = "test-csrf-token"

    def setUp(self):
        search.app.testing = True
        self.client = search.app.test_client()
        start_auth_db_patch(self)
        patcher = mock.patch.dict("os.environ", {"DISABLE_IP_ALLOWLIST": "1"})
        patcher.start()
        self.addCleanup(patcher.stop)

    def login(self, admin=True):
        with self.client.session_transaction() as sess:
            sess.update(authenticated=True, user_id=1, auth_version=1, is_admin=admin, csrf_token=self.CSRF)
            if not admin:
                sess["team_id"] = 1

    def post_parse(self, raw=None, name="bao-gia.xlsx", csrf=True, **extra):
        data = {"quote": (io.BytesIO(raw if raw is not None else make_quote([ITEM])), name), **extra}
        headers = {"X-CSRF-Token": self.CSRF} if csrf else {}
        return self.client.post("/quote-to-order/parse", data=data, headers=headers, content_type="multipart/form-data")

    def test_page_requires_login(self):
        response = self.client.get("/quote-to-order/")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/login", response.headers["Location"])

    def test_api_requires_login(self):
        self.assertEqual(self.client.post("/quote-to-order/parse").status_code, 401)
        self.assertEqual(self.client.post("/quote-to-order/export", json={"rows": [valid_row()]}).status_code, 401)

    def test_team_without_quick_quote_is_denied(self):
        self.login(admin=False)
        with mock.patch.object(team_permissions, "current_permissions", return_value=frozenset({"SEARCH"})):
            self.assertEqual(self.client.get("/quote-to-order/").status_code, 403)
            self.assertEqual(self.post_parse().status_code, 403)
            response = self.client.post("/quote-to-order/export", json={"rows": [valid_row()]},
                                        headers={"X-CSRF-Token": self.CSRF})
            self.assertEqual(response.status_code, 403)

    def test_team_with_quick_quote_sees_page_and_menu_link(self):
        self.login(admin=False)
        grants = frozenset({"QUICK_QUOTE", "VIEW_PRICE", "SEARCH"})
        with mock.patch.object(team_permissions, "current_permissions", return_value=grants):
            response = self.client.get("/quote-to-order/")
            self.assertEqual(response.status_code, 200)
            body = response.get_data(as_text=True)
            self.assertIn("Báo giá → Đơn hàng", body)
            self.assertIn('href="/quote-to-order/"', body)
            self.assertIn("qtoConfig", body)
            self.assertIn("no-store", response.headers["Cache-Control"])

    def test_menu_link_hidden_without_quick_quote(self):
        self.login(admin=False)
        with mock.patch.object(team_permissions, "current_permissions", return_value=frozenset({"SEARCH"})):
            body = self.client.get("/").get_data(as_text=True)
        self.assertNotIn('href="/quote-to-order/"', body)

    def test_csrf_is_required_on_both_posts(self):
        self.login()
        self.assertEqual(self.post_parse(csrf=False).status_code, 400)
        response = self.client.post("/quote-to-order/export", json={"rows": [valid_row()]})
        self.assertEqual(response.status_code, 400)
        response = self.client.post("/quote-to-order/export", json={"rows": [valid_row()]},
                                    headers={"X-CSRF-Token": "sai"})
        self.assertEqual(response.status_code, 400)

    def test_parse_happy_path_returns_items_and_mapping(self):
        self.login()
        response = self.post_parse()
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["filename"], "bao-gia.xlsx")
        self.assertEqual(data["items"][0]["price"], 10800000)
        self.assertEqual(data["unmapped_required"], [])

    def test_parse_accepts_mapping_and_header_row_form_fields(self):
        self.login()
        raw = make_quote([ITEM], renames=RENAMES)
        first = self.post_parse(raw).get_json()
        self.assertTrue(first["needs_header"])
        second = self.post_parse(raw, header_row="10").get_json()
        cols = {c["title"]: c["index"] for c in second["columns"]}
        import json
        third = self.post_parse(raw, header_row="10",
                                mapping=json.dumps({"name": cols["Mặt hàng"], "code": cols["Mã số"]})).get_json()
        self.assertEqual(third["items"][0]["code"], "M-502-01N")

    def test_parse_rejects_wrong_extension_garbage_and_missing_file(self):
        self.login()
        for response in (self.post_parse(name="bao-gia.xls"), self.post_parse(name="bao-gia.xlsm"),
                         self.post_parse(raw=b"hello world" * 50),
                         self.client.post("/quote-to-order/parse", data={}, headers={"X-CSRF-Token": self.CSRF})):
            self.assertEqual(response.status_code, 400)
            self.assertFalse(response.get_json()["ok"])
            self.assertNotIn("Traceback", response.get_data(as_text=True))

    def test_parse_rejects_oversized_upload(self):
        self.login()
        big = b"PK" + b"\x00" * (q.MAX_XLSX_BYTES + 10)
        response = self.post_parse(raw=big)
        self.assertIn(response.status_code, (400, 413))

    def test_export_returns_workbook(self):
        self.login()
        response = self.client.post("/quote-to-order/export", json={"rows": [valid_row(), valid_row(code="B")]},
                                    headers={"X-CSRF-Token": self.CSRF})
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.headers["Content-Type"])
        self.assertRegex(response.headers["Content-Disposition"], r'filename=bang-hang-hoa_\d{8}_\d{4}\.xlsx')
        ws = load_workbook(io.BytesIO(response.data))["Exported file"]
        self.assertEqual(ws.max_row, 3)

    def test_export_with_nothing_selected_creates_no_file(self):
        self.login()
        for payload in ({"rows": []}, {}, {"rows": "x"}):
            response = self.client.post("/quote-to-order/export", json=payload, headers={"X-CSRF-Token": self.CSRF})
            self.assertEqual(response.status_code, 400)
            self.assertEqual(response.get_json()["error"], "Chưa chọn dòng hàng nào.")

    def test_export_blocks_incomplete_rows_and_reports_which(self):
        self.login()
        response = self.client.post("/quote-to-order/export",
                                    json={"rows": [valid_row(), valid_row(type=""), valid_row(qty=0)]},
                                    headers={"X-CSRF-Token": self.CSRF})
        self.assertEqual(response.status_code, 400)
        data = response.get_json()
        self.assertEqual([e["index"] for e in data["rows"]], [1, 2])

    def test_export_row_limit(self):
        self.login()
        rows = [valid_row()] * (q.MAX_ORDER_ROWS + 1)
        response = self.client.post("/quote-to-order/export", json={"rows": rows}, headers={"X-CSRF-Token": self.CSRF})
        self.assertEqual(response.status_code, 400)

    def test_export_failure_does_not_leak_internals(self):
        self.login()
        with mock.patch.object(q, "build_order_file", side_effect=RuntimeError("secret path /opt/x")):
            response = self.client.post("/quote-to-order/export", json={"rows": [valid_row()]},
                                        headers={"X-CSRF-Token": self.CSRF})
        self.assertEqual(response.status_code, 500)
        self.assertNotIn("secret path", response.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
