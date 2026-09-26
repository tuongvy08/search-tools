"""Inventory notes parsing, preview, visibility, and template tests without a DB."""
from datetime import date
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from flask import Flask
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook, load_workbook

import admin_stock
from import_engine import ImportProblem
import stock
import stock_import_jobs as jobs
import team_permissions


class StockNotesTests(unittest.TestCase):
    def parse(self, note="Kho A\nHàng mẫu", header="Ghi chú", legacy=False):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "notes.xlsx"
            workbook = Workbook()
            workbook.active.append(list(jobs.HEADERS[:-1]) + ([] if legacy else [header]))
            workbook.active.append(
                ["A", "C1", "50-00-0", "Brand A", "1g", None, 1, date(2027, 1, 1)]
                + ([] if legacy else [note])
            )
            workbook.save(path)
            workbook.close()
            return jobs.parse_workbook(path)[0]

    def test_notes_roundtrip_text_newlines_and_aliases(self):
        note = 'Kho A\nHàng mẫu <script>alert("x")</script>'
        for header in ("Ghi chú", "note", "notes"):
            with self.subTest(header=header):
                self.assertEqual(self.parse(note, header)["stock_note"], note)

    def test_legacy_and_empty_notes(self):
        self.assertEqual(self.parse(legacy=True)["stock_note"], "")
        self.assertEqual(self.parse(None)["stock_note"], "")
        self.assertEqual(self.parse("  ")["stock_note"], "")

    def test_note_limit_and_formula_rejection(self):
        self.assertEqual(len(self.parse("a" * 2000)["stock_note"]), 2000)
        for note in ("a" * 2001, '=HYPERLINK("https://example.com")'):
            with self.subTest(note=note[:20]), self.assertRaises(ImportProblem):
                self.parse(note)

    def test_legacy_header_does_not_import_unmapped_ninth_cell(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "legacy.xlsx"
            workbook = Workbook()
            workbook.active.append(list(jobs.HEADERS[:-1]))
            workbook.active.append(["A", "C1", None, "Brand A", "1g", None, 1, None, "Unmapped"])
            workbook.save(path)
            workbook.close()
            self.assertEqual(jobs.parse_workbook(path)[0]["stock_note"], "")

    def test_unknown_ninth_header_rejected(self):
        with self.assertRaises(ImportProblem):
            self.parse(header="Unknown")

    def test_note_edit_is_changed_row_and_changes_confirmation_digest(self):
        original = self.parse("Kho A")
        revised = dict(original, stock_note="Kho B — hàng mẫu")
        cur = mock.Mock()
        cur.fetchall.return_value = [tuple(original[key] for key in (
            "name", "code", "cas", "brand", "size", "stock_price_vnd", "quantity", "expiry_date", "stock_note"
        ))]
        with mock.patch.object(jobs, "_prepare_rows", side_effect=lambda cur, rows, **kw: (rows, [])), \
             mock.patch.object(jobs, "active_snapshot", return_value={"id": "snapshot", "revision": 1}), \
             mock.patch.object(jobs, "snapshot_fingerprint", return_value="fingerprint"):
            unchanged = jobs.build_plan(cur, [original])
            changed = jobs.build_plan(cur, [revised])
        self.assertEqual(unchanged["unchanged"], 1)
        self.assertEqual(changed["changed"], 1)
        self.assertEqual(changed["added"], 0)
        self.assertNotEqual(unchanged["plan_digest"], changed["plan_digest"])
        self.assertEqual(changed["sample"][0]["stock_note"], revised["stock_note"])

    def test_stock_note_visibility_uses_view_note_permission(self):
        cur = mock.Mock()
        cur.fetchall.return_value = [(1, "A", "C1", "50-00-0", "Brand A", "1g", None,
                                     1, None, "c1", "50-00-0", "Kho A", None)]
        for grants in ([], ["VIEW_NOTE"]):
            with self.subTest(grants=grants):
                result = stock.fetch_stock_options(cur, codes=["C1"], is_admin=False,
                                                   team_id=1, grants=grants, allow_same_cas=False)
                option = result["all"][0]
                self.assertEqual("Stock_Note" in option, "VIEW_NOTE" in grants)
                if grants:
                    self.assertEqual(option["Stock_Note"], "Kho A")
        redacted = team_permissions.redact({"Stock_Options": [{"Stock_Note": "Private"}]}, grants=[])
        self.assertEqual(redacted["Stock_Options"], [{}])

    def test_downloaded_template_has_ninth_column_and_parses(self):
        app = Flask(__name__)
        app.secret_key = "notes-test-only"
        admin_stock.register(app, lambda: None, lambda: "test")
        with mock.patch.object(jobs, "connection") as connection:
            connection.return_value.__enter__.return_value.cursor.return_value.__enter__.return_value.fetchone.return_value = (1,)
            response = app.test_client().get("/admin/stock/template")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data))
        self.assertEqual(workbook.active.max_column, 9)
        self.assertEqual(workbook.active.cell(1, 9).value, "Ghi chú")
        self.assertEqual(workbook.active.cell(2, 9).value, "Kho A — hàng mẫu")
        workbook.close()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "template.xlsx"
            path.write_bytes(response.data)
            self.assertEqual(jobs.parse_workbook(path)[0]["stock_note"], "Kho A — hàng mẫu")

    def test_preview_escapes_notes_and_handles_pre_upgrade_preview(self):
        env = Environment(loader=FileSystemLoader(Path(__file__).resolve().parents[1] / "templates"),
                          autoescape=select_autoescape())
        env.globals.update(url_for=lambda *a, **kw: "/test", csrf_token=lambda: "test")
        template = env.get_template("admin_stock.html")
        sample = dict(code="C1", name="A", brand="Brand A", size="1g", cas="", quantity=1,
                      stock_price="", expiry="")
        for notes in ({}, {"stock_note": "Kho A\n<script>alert(1)</script>"}):
            with self.subTest(notes=notes):
                job = dict(id="test", status="completed", phase="preview", preview_ready=True,
                           preview=dict(sample=[dict(sample, **notes)], brands=[], new_brands=[]))
                html = template.render(session={}, job=job, status_labels={}, phase_labels={}, events=[])
                self.assertIn("<th>Ghi chú</th>", html)
                self.assertNotIn("<script>alert(1)</script>", html)
                if notes:
                    self.assertIn("Kho A\n&lt;script&gt;alert(1)&lt;/script&gt;", html)


if __name__ == "__main__":
    unittest.main()
