"""Phase 6C0.1 boundary/security gates against isolated PostgreSQL."""
import io
import json
import os
import unittest
from unittest import mock
from urllib.parse import parse_qs, urlparse

import psycopg2
from openpyxl import load_workbook
import search
import team_permissions as permissions
from pg_temp_db import (create_full_schema_temp_db, drop_temp_db, probe_postgres_reachable,
                        apply_brand_master_and_currency_migrations, apply_dynamic_brand_currency_migration)
from test_quote_workbook_export import make_workbook


class RegistryTests(unittest.TestCase):
    def test_explicit_empty_and_unknown(self):
        self.assertEqual(permissions.validate_permissions([]), [])
        for value in (None, 'SEARCH', ['SEARCH', 'TYPO'], [False]):
            with self.assertRaises(ValueError):
                permissions.validate_permissions(value)

    def test_recursive_aliases_and_combined_fields(self):
        sample = {'Name': 'name', 'Cas': 'cas', 'Unit_Price': '999', 'selected': {
            'price': 999, 'ship': 88, 'unit_price_value': 999, 'Compliance_Combined': 'secret',
            'Compliance_Note': 'secret', 'requested_cas': 'cas', 'warnings': ['CẤM NHẬP']}}
        value = permissions.redact(sample, {'VIEW_NAME'})
        self.assertEqual(value, {'Name': 'name', 'selected': {}})


@unittest.skipUnless(probe_postgres_reachable(), 'local PostgreSQL required')
class TeamCapabilitiesPgTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dbname, cls.dsn = create_full_schema_temp_db()
        cls.env = mock.patch.dict(os.environ, DATABASE_URL=cls.dsn, DISABLE_IP_ALLOWLIST='1')
        cls.env.start()
        try:
            conn = psycopg2.connect(cls.dsn)
            with conn, conn.cursor() as cur:
                apply_brand_master_and_currency_migrations(cur)
                apply_dynamic_brand_currency_migration(cur)
                cur.execute("INSERT INTO teams(name) VALUES ('Biology') RETURNING id")
                cls.team = cur.fetchone()[0]
                cur.execute("INSERT INTO app_users(username,password_hash,is_admin) VALUES ('admin','x',true),('admin2','x',true) RETURNING id")
                cls.admin, cls.admin2 = [row[0] for row in cur.fetchall()]
                cur.execute("INSERT INTO app_users(username,password_hash,team_id,auth_provider,account_status,email,google_sub) VALUES "
                            "('local','x',%s,'LOCAL','ACTIVE',NULL,NULL),"
                            "('google',NULL,%s,'GOOGLE','ACTIVE','staff@example.test','synthetic-google') RETURNING id",
                            (cls.team, cls.team))
                cls.local, cls.google = [row[0] for row in cur.fetchall()]
                cur.execute("INSERT INTO team_brands(team_id,brand) VALUES (%s,'TRC')", (cls.team,))
                cur.execute("INSERT INTO products(name,code,cas,brand,source_brand,size,ship,price,note,manual_compliance,manual_compliance_note) "
                            "VALUES ('HiddenName','ITEM-X','123-45-6','TRC','TRC','100mg','1','100','SecretNote','Được bán','SecretCompliance') RETURNING id")
                cls.product = cur.fetchone()[0]
                cur.execute("INSERT INTO products(name,code,cas,brand,source_brand,size,ship,price,note,manual_compliance) "
                            "VALUES ('NoPrice','ITEM-NO-PRICE','223-45-6','TRC','TRC','100mg','1',NULL,'','Được bán') RETURNING id")
                cls.no_price_product = cur.fetchone()[0]
                cur.execute("INSERT INTO brand_compliance_settings(brand_norm,manual_compliance_priority) VALUES ('TRC',true)")
                template_raw = make_workbook()
                cur.execute("""INSERT INTO quote_templates(filename,content,content_sha256,content_size,profile_version,mapping_json,is_active)
                    VALUES ('team-test.xlsx',%s,%s,%s,'BG_V1',%s::jsonb,true)""",
                    (psycopg2.Binary(template_raw), __import__('hashlib').sha256(template_raw).hexdigest(),
                     len(template_raw), json.dumps(search._quote_template_mapping_snapshot())))
            conn.close()
        except Exception:
            cls.env.stop()
            drop_temp_db(cls.dbname)
            raise

    @classmethod
    def tearDownClass(cls):
        cls.env.stop()
        drop_temp_db(cls.dbname)

    def setUp(self):
        self.set_grants(permissions.LEGACY_PERMISSIONS)
        self.client = self.login(self.local)

    def sql(self, query, args=(), fetch=False):
        conn = psycopg2.connect(self.dsn)
        try:
            with conn, conn.cursor() as cur:
                cur.execute(query, args)
                return cur.fetchall() if fetch else None
        finally:
            conn.close()

    def set_grants(self, grants):
        self.sql("UPDATE teams SET permission_keys=%s, updated_at=clock_timestamp() WHERE id=%s", (list(grants),self.team))

    def login(self, user):
        client = search.app.test_client()
        with client.session_transaction() as sess:
            sess.update(authenticated=True, user_id=user, auth_version=1,
                        is_admin=user in {self.admin,self.admin2}, team_id=self.team,
                        auth_provider='GOOGLE' if user == self.google else 'LOCAL', csrf_token='csrf-test')
        return client

    def match(self, client=None, **extra):
        return (client or self.client).post('/api/quote-assistant/match', json={'rows':[{'code':'ITEM-X'}], **extra})

    def test_legacy_defaults_admin_bypass_and_provider_equivalence(self):
        local = self.client.get('/search?query=ITEM-X')
        google = self.login(self.google).get('/search?query=ITEM-X')
        self.assertEqual(local.status_code,200)
        self.assertEqual(local.json,google.json)
        self.assertIn('Unit_Price',local.json['results'][0])
        self.set_grants([])
        admin = self.login(self.admin).get('/search?query=ITEM-X')
        self.assertEqual(admin.status_code,200)
        self.assertEqual(admin.json,local.json)

    def test_live_redaction_in_all_search_responses_and_transfer(self):
        grants = set(permissions.LEGACY_PERMISSIONS) - {'VIEW_PRICE','VIEW_CAS','VIEW_COMPLIANCE','VIEW_COMPLIANCE_NOTE'}
        self.set_grants(grants)
        for client in (self.client,self.login(self.google)):
            responses = [client.get('/search?query=ITEM-X'),client.post('/find_code_batch',data={'codes':'ITEM-X'}),
                         client.post('/advanced_search',data={'cas':'123-45-6'})]
            for response in responses:
                self.assertEqual(response.status_code,200,response.data)
                self.assertTrue(response.json['results'])
                raw = response.get_data(as_text=True)
                for forbidden in ('Unit_Price','Currency_Rate','123-45-6','SecretCompliance','compliance'):
                    self.assertNotIn(forbidden,raw)
            self.assertEqual(self.match(client).status_code,403)
        payload = {'source':'SEARCH','rows':[{
            'Name':'Name', 'Cas':'123-45-6','Unit_Price':'987654321','Compliance':'CẤM NHẬP'}]}
        copied = self.client.post('/api/results/copy', json=payload,
                                  headers={'X-CSRF-Token':'csrf-test'})
        self.assertEqual(copied.status_code,200)
        self.assertNotIn('987654321',copied.get_data(as_text=True))
        self.assertNotIn('123-45-6',copied.get_data(as_text=True))
        exported = self.client.post('/api/results/export', json=payload,
                                    headers={'X-CSRF-Token':'csrf-test'})
        self.assertEqual(exported.status_code,403)
        self.assertNotIn('987654321',exported.get_data(as_text=True))
        self.assertNotIn('123-45-6',exported.get_data(as_text=True))

    def test_feature_routes_fail_closed_and_source_cannot_bypass(self):
        self.set_grants([])
        for path, method in [('/search','get'),('/check_cas','get'),('/check_cas_batch','post'),
                             ('/find_code_batch','post'),('/advanced_search/options','post'),('/advanced_search','post'),
                             ('/quote-assistant/quick','get'),('/api/quote-assistant/preflight','post'),
                             ('/api/quote-assistant/match','post'),('/api/quote-assistant/workbook/export','post'),
                             ('/api/quote-assistant/workbook/template','get'),('/api/results/copy','post'),('/api/results/export','post')]:
            self.assertEqual(getattr(self.client,method)(path).status_code,403,path)
        self.assertEqual(self.client.get('/').status_code,200)
        self.set_grants(['COPY'])
        self.assertEqual(self.client.post('/api/results/copy',json={'source':'SEARCH','rows':[{}]},headers={'X-CSRF-Token':'csrf-test'}).status_code,403)
        self.assertEqual(search.app.test_client().get('/search').status_code,401)

    def test_search_by_cas_is_independent_and_all_quote_aliases_guarded(self):
        self.set_grants(set(permissions.LEGACY_PERMISSIONS)-{'SEARCH_BY_CAS'})
        self.assertEqual(self.client.get('/search?query=123-45-6').json['results'],[])
        self.assertEqual(self.client.get('/search?query=ITEM-X').json['results'][0]['Cas'],'123-45-6')
        for endpoint in ('match','preflight'):
            for payload in ({'rows':[{'cas':'123-45-6'}]},
                            {'rows':[{'code':'ITEM-X','equivalent_override':True}]},
                            {'rows':[{'code':'ITEM-X'}],'equivalent_search_default':True}):
                r=self.client.post('/api/quote-assistant/'+endpoint,json=payload)
                self.assertEqual(r.status_code,403,(endpoint,payload,r.data))
        self.assertEqual(self.match().status_code,200)

    def test_quote_nested_redaction_copy_and_workbook(self):
        self.set_grants(set(permissions.LEGACY_PERMISSIONS)-{'VIEW_CAS','VIEW_COMPLIANCE','VIEW_COMPLIANCE_NOTE','VIEW_NOTE','VIEW_NAME'})
        response=self.match()
        self.assertEqual(response.status_code,200,response.data)
        self.assertTrue(response.json['results'][0]['candidates'])
        for forbidden in ('HiddenName','123-45-6','SecretCompliance','SecretNote'):
            self.assertNotIn(forbidden,response.get_data(as_text=True))
        data={'source':'QUICK_QUOTE','rows':[{'product_id':self.product}]}
        response=self.client.post('/api/results/copy',json=data,headers={'X-CSRF-Token':'csrf-test'})
        self.assertEqual(response.status_code,200,response.data)
        self.assertNotIn('123-45-6',response.get_data(as_text=True))
        response=self.client.post('/api/quote-assistant/workbook/export',data={
            'selections':json.dumps([{'product_id':self.product}])},headers={'X-CSRF-Token':'csrf-test'})
        # Deny before product lookup: success/failure must not reveal hidden compliance.
        self.assertEqual(response.status_code,403,response.data[:200])
        values=response.get_data(as_text=True)
        for forbidden in ('HiddenName','123-45-6','SecretCompliance','SecretNote'):
            self.assertNotIn(forbidden,values)

    def test_search_quote_export_without_view_price_uses_clean_allowed_columns(self):
        self.set_grants(set(permissions.LEGACY_PERMISSIONS)-{'VIEW_PRICE','QUICK_QUOTE'})
        response=self.client.post('/api/results/quote-export',data={'source':'SEARCH',
            'selections':json.dumps([{'product_id':self.product}])},headers={'X-CSRF-Token':'csrf-test'})
        self.assertEqual(response.status_code,200,response.data[:200])
        check=load_workbook(io.BytesIO(response.data),data_only=False)
        self.assertEqual(check.sheetnames,['Báo giá'])
        headers=[cell.value for cell in check['Báo giá'][1]]
        self.assertIn('Code',headers)
        self.assertNotIn('Unit Price',headers)
        self.assertEqual(check['Báo giá'].cell(2,headers.index('Code')+1).value,'ITEM-X')
        check.close()

        # Price validity is not an oracle: missing-price and valid-price rows
        # both export successfully without a price column.
        missing=self.client.post('/api/results/quote-export',data={'source':'SEARCH',
            'selections':json.dumps([{'product_id':self.no_price_product}])},headers={'X-CSRF-Token':'csrf-test'})
        self.assertEqual(missing.status_code,200,missing.data[:200])
        missing_check=load_workbook(io.BytesIO(missing.data),data_only=False)
        missing_headers=[cell.value for cell in missing_check['Báo giá'][1]]
        self.assertNotIn('Unit Price',missing_headers)
        self.assertEqual(missing_check['Báo giá'].cell(2,missing_headers.index('Code')+1).value,'ITEM-NO-PRICE')
        missing_check.close()

    def test_hidden_compliance_reasons_and_fallback_counts_are_generic(self):
        self.set_grants(set(permissions.LEGACY_PERMISSIONS)-{'VIEW_COMPLIANCE','VIEW_COMPLIANCE_NOTE'})
        self.sql("UPDATE products SET manual_compliance='CẤM NHẬP' WHERE id=%s",(self.product,))
        try:
            response=self.match()
            self.assertEqual(response.status_code,200)
            row=response.json['results'][0]
            self.assertEqual(row['reason_code'],'NOT_ELIGIBLE')
            self.assertFalse(row['candidates'][0]['eligible'])
            self.assertEqual(row['candidates'][0]['ineligible_reason'],'NOT_ELIGIBLE')
            self.assertNotIn('COMPLIANCE',response.get_data(as_text=True))
            self.assertEqual(self.client.post('/api/results/copy',json={
                'source':'QUICK_QUOTE','rows':[{'product_id':self.product}]},
                headers={'X-CSRF-Token':'csrf-test'}).status_code,400)
            response=self.client.post('/api/quote-assistant/match',json={
                'rows':[{'cas':'123-45-6'}], 'global_brand_policy':{
                    'mode':'PRIORITY_FALLBACK','priority_tiers':[{'brands':['TRC']}]} })
            self.assertEqual(response.status_code,200,response.data)
            self.assertNotIn('COMPLIANCE',response.get_data(as_text=True))
        finally:
            self.sql("UPDATE products SET manual_compliance='Được bán' WHERE id=%s",(self.product,))

    def preview(self, admin, grants):
        r=admin.post('/admin/teams/preview',data={'csrf_token':'csrf-test','team_id':str(self.team),
            'ip_policy':'INHERIT','brands':['TRC'],'permissions_present':'1','permissions':list(grants)})
        self.assertIn('preview=',r.location,r.location)
        return parse_qs(urlparse(r.location).query)['preview'][0]

    def test_preview_confirm_live_change_history_and_staleness(self):
        admin=self.login(self.admin)
        grants=set(permissions.LEGACY_PERMISSIONS)-{'VIEW_PRICE'}
        token=self.preview(admin,grants)
        self.assertIn('Unit_Price',self.client.get('/search?query=ITEM-X').json['results'][0])
        wrong=self.login(self.admin2).post('/admin/teams/confirm',data={'csrf_token':'csrf-test','preview_token':token})
        self.assertIn('err=',wrong.location)
        result=admin.post('/admin/teams/confirm',data={'csrf_token':'csrf-test','preview_token':token,'permissions':list(permissions.REGISTRY)})
        self.assertIn('msg=',result.location)
        self.assertNotIn('Unit_Price',self.client.get('/search?query=ITEM-X').json['results'][0])
        self.assertTrue(self.sql('SELECT id FROM team_capability_history WHERE team_id=%s',(self.team,),True))
        self.assertTrue(self.sql("SELECT id FROM login_audit_events WHERE reason_code='TEAM_CAPABILITIES_UPDATED' AND target_team_id=%s",(self.team,),True))
        token=self.preview(admin,[])
        self.set_grants(permissions.LEGACY_PERMISSIONS)
        stale=admin.post('/admin/teams/confirm',data={'csrf_token':'csrf-test','preview_token':token})
        self.assertIn('err=',stale.location)

    def test_csrf_unknown_policy_and_unavailable_schema(self):
        admin=self.login(self.admin)
        self.assertEqual(admin.post('/admin/teams/preview',data={'team_id':self.team}).status_code,400)
        self.assertEqual(self.client.post('/api/results/copy',json={}).status_code,400)
        r=admin.post('/admin/teams/preview',data={'csrf_token':'csrf-test','team_id':str(self.team),
            'permissions_present':'1','permissions':['TYPO']})
        self.assertIn('err=',r.location)
        self.sql('UPDATE teams SET permission_keys=ARRAY[\'TYPO\'] WHERE id=%s',(self.team,))
        self.assertEqual(self.client.get('/search').status_code,503)


if __name__ == '__main__':
    unittest.main()
